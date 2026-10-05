import json
import hmac
import hashlib
import logging
import csv
import io
import os
import re
import time
from threading import Lock
import uuid
import base64
import requests
from io import BytesIO
from django.core.files.base import ContentFile
from django.http import HttpResponse, HttpResponseForbidden, JsonResponse
from django.views.decorators.http import require_GET, require_POST, require_http_methods
from django.views.decorators.csrf import csrf_exempt
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import datetime, timedelta
from reportlab.lib.units import mm
from reportlab.graphics.barcode import createBarcodeDrawing
from reportlab.graphics import renderSVG
from django.db.models.deletion import ProtectedError
from django.utils.html import escape
from openpyxl import load_workbook
import qrcode
from django.conf import settings
from django.core import signing
from django.core.signing import BadSignature, SignatureExpired
from django.shortcuts import get_object_or_404, redirect, render
from django.contrib import messages
from django.contrib.messages import get_messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.models import User
from django.db import transaction
from django.db.models import (
    Count,
    DecimalField,
    ExpressionWrapper,
    F,
    OuterRef,
    Q,
    Subquery,
    Sum,
    Value,
)
from django.db.models.functions import Coalesce
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone

from .models import (
    ActivityLog,
    Bidhaa,
    BUSINESS_TYPE_CHOICES,
    BusinessSubscription,
    Duka,
    Kategoria,
    LoginHistory,
    Mauzo,
    Matumizi,
    Mteja,
    MzigoUlioingia,
    Notification,
    Profile,
    PurchaseRequest,
    PurchaseRequestItem,
    SaleReturn,
    VochaYaDuka,
    CommunicationLog,
    DailyReportSettings,
    PaymentTransaction,
    ProductStock,
    StockLocation,
    StockMovement,
    StockTransfer,
)

from .inventory_services import (
    ensure_inventory_locations,
    current_quantity,
    add_stock,
    perform_transfer,
    perform_adjustments,
    perform_stock_take,
    get_last_verified,
)

from .tenant import reset_current_duka, set_current_duka_id


logger = logging.getLogger(__name__)

# AI image provider guard: keep automatic product-image generation gentle
# enough for normal API rate limits. This is process-local and additive; it
# does not affect sales, stock, checkout, or other POS requests.
_ai_image_rate_lock = Lock()
_ai_image_next_request_at = 0.0
_AI_IMAGE_MIN_INTERVAL = max(1.0, float(os.environ.get("TRADECORE_AI_IMAGE_MIN_INTERVAL", "8")))


# TENANT ISOLATION TOUCH 2026-10-03: view-layer defense-in-depth only.
# Models/managers remain the primary scope; request tenant is established by middleware.

# =========================================================
# HELPERS
# =========================================================

def parse_decimal(value, default=Decimal("0.00")):
    if value in (None, ""):
        return default

    try:
        return Decimal(str(value))
    except (InvalidOperation, TypeError, ValueError):
        return default


def parse_positive_int(value, default=0):
    try:
        number = int(value)
        return number if number > 0 else default
    except (TypeError, ValueError):
        return default


def parse_nonnegative_int(value, default=0):
    try:
        number = int(value)
        return number if number >= 0 else default
    except (TypeError, ValueError):
        return default


def create_activity(user, action, kind="kawaida"):
    profile = getattr(user, "profile", None) if getattr(user, "is_authenticated", False) else None
    duka = getattr(user, "duka", None) or (getattr(profile, "duka", None) if profile else None)
    ActivityLog.objects.create(
        duka=duka,
        mhusika=user
        if getattr(user, "is_authenticated", False)
        else None,
        kitendo=action,
        aina=kind,
    )
    

def get_profile_role(user):
    if user.is_superuser:
        return "admin"

    profile = getattr(user, "profile", None)
    if not profile:
        return None

    if getattr(profile, "ni_admin", False):
        return "admin"

    # SURGERY: Tunatumia Dynamic Permissions kuamua mfanyakazi aanzie page gani (Landing Page)
    if getattr(profile, "ruhusa_stoo", False) and not getattr(profile, "ruhusa_mauzo", False):
        return "stoo"
    if getattr(profile, "ruhusa_mauzo", False):
        return "cashier"

    # Legacy fallback kwa wale wa zamani
    if getattr(profile, "ni_stoo", False):
        return "stoo"
    if getattr(profile, "ni_cashier", False):
        return "cashier"

    role = (getattr(profile, "jukumu", None) or "").strip().lower()
    return role if role in {"admin", "cashier", "stoo"} else None


def get_client_ip(request):
    forwarded = request.META.get("HTTP_X_FORWARDED_FOR")

    if forwarded:
        return forwarded.split(",")[0].strip()

    return request.META.get("REMOTE_ADDR")


def redirect_by_role(user):
    if user.is_superuser:
        return redirect("dashboard")

    role = get_profile_role(user)

    if role == "cashier":
        return redirect("mauzo")

    if role == "stoo":
        return redirect("stoo_bidhaa")

    if role == "admin":
        return redirect("dashboard")

    return redirect("dashboard")


def user_can(request, *roles):
    if request.user.is_superuser:
        return True
        
    profile = getattr(request.user, "profile", None)
    if not profile:
        return False

    if getattr(profile, "ni_admin", False):
        return True

    # SURGERY: Tunaruhusu access kwa kuangalia URL na Checkbox husika
    path = request.path
    if "/wateja" in path and getattr(profile, "ruhusa_wateja", False): return True
    if "/matumizi" in path and getattr(profile, "ruhusa_matumizi", False): return True
    if "/ripoti" in path and getattr(profile, "ruhusa_ripoti", False): return True
    if "/ujumbe" in path and getattr(profile, "ruhusa_mawasiliano", False): return True
    if "/stoo" in path and getattr(profile, "ruhusa_stoo", False): return True
    if "/mauzo" in path and getattr(profile, "ruhusa_mauzo", False): return True

    return get_profile_role(request.user) in roles


def require_roles(request, *roles):
    if user_can(request, *roles):
        return None
    messages.error(request, "Huna ruhusa ya kufungua ukurasa huu.")
    return redirect_by_role(request.user)


def make_invoice_number():
    return f"TC-{timezone.now():%Y%m%d}-{uuid.uuid4().hex[:8].upper()}"


# =========================================================
# 1. USIMAMIZI WA STOO NA INVENTORY
# =========================================================

def _ean13_checksum(body12):
    digits = [int(ch) for ch in str(body12)]
    total = sum((d if i % 2 == 0 else d * 3) for i, d in enumerate(digits))
    return (10 - (total % 10)) % 10


def generate_internal_barcode():
    used = set(Bidhaa.objects.exclude(barcode__isnull=True).exclude(barcode="").values_list("barcode", flat=True))
    next_id = (Bidhaa.objects.order_by("-id").values_list("id", flat=True).first() or 0) + 1
    body = f"200{next_id:09d}"[:12]
    code = body + str(_ean13_checksum(body))
    counter = next_id
    while code in used:
        counter += 1
        body = f"200{counter:09d}"[:12]
        code = body + str(_ean13_checksum(body))
    return code


BUSINESS_CATEGORY_PACKS = {
    "FASHION": ["Men", "Women", "Kids", "Shirts", "T-Shirts", "Trousers", "Jeans", "Dresses", "Skirts", "Jackets", "Sweaters", "Shoes", "Sandals", "Bags", "Belts", "Sportswear", "School Uniforms", "Kanga & Vitenge", "Fabrics", "Accessories", "Mitumba", "New Stock", "Clearance"],
    "COMPUTERS": ["Laptops", "Desktop Computers", "Monitors", "Keyboards", "Mouse", "Printers", "Scanners", "Hard Drives", "SSD", "RAM", "Flash Disks", "Routers", "Switches", "Network Devices", "Cables", "Chargers", "Batteries", "CCTV", "Software", "Computer Accessories", "Gaming"],
    "SUPERMARKET": ["Beverages", "Water", "Juices", "Soft Drinks", "Milk & Dairy", "Bread & Bakery", "Biscuits", "Rice", "Sugar", "Flour", "Cooking Oil", "Spices", "Canned Foods", "Snacks", "Personal Care", "Cleaning Products", "Baby Products", "Household", "Stationery", "Toiletries", "Fresh Foods"],
    "COSMETICS": ["Skincare", "Face Care", "Body Care", "Hair Care", "Shampoo", "Conditioner", "Lotions", "Perfumes", "Makeup", "Lip Products", "Foundation", "Powder", "Hair Extensions", "Wigs", "Nail Care", "Beauty Tools", "Men's Grooming", "Baby Care", "Accessories"],
    "PHARMACY": ["Pain Relief", "Cold & Flu", "Digestive Care", "First Aid", "Vitamins", "Personal Care", "Baby Care", "Medical Devices", "Bandages", "Antiseptics", "Oral Care", "Skin Care", "Women's Health", "Men's Health", "Supplements", "Other Health Products"],
    "HARDWARE": ["Cement", "Sand", "Aggregate", "Steel", "Iron Sheets", "Timber", "Paint", "Plumbing", "Electrical", "Tools", "Nails", "Screws", "Bolts & Nuts", "Locks", "Door Hardware", "Building Chemicals", "Tiles", "Pipes", "Fittings", "Safety Equipment", "Adhesives"],
    "RESTAURANT": ["Breakfast", "Lunch", "Dinner", "Main Dishes", "Snacks", "Fast Food", "Rice Dishes", "Ugali Dishes", "Meat", "Chicken", "Fish", "Vegetarian", "Salads", "Soups", "Desserts", "Juices", "Soft Drinks", "Hot Drinks", "Combos", "Takeaway"],
    "STATIONERY": ["Exercise Books", "Notebooks", "Pens", "Pencils", "Markers", "Erasers", "Rulers", "Files & Folders", "Paper", "Printing Paper", "Envelopes", "Staplers", "Staples", "Calculators", "School Supplies", "Office Supplies", "Art Supplies", "Books", "Textbooks", "Printing Accessories"],
    "PHONES": ["Smartphones", "Feature Phones", "Phone Cases", "Screen Protectors", "Chargers", "Power Banks", "USB Cables", "Earphones", "Headphones", "Memory Cards", "Phone Batteries", "Adapters", "Smart Watches", "Speakers", "Phone Accessories", "SIM Accessories", "Networking Accessories"],
    "AUTOPARTS": ["Engine Parts", "Brake System", "Suspension", "Steering", "Electrical", "Filters", "Belts", "Bearings", "Clutch", "Cooling System", "Fuel System", "Body Parts", "Lights", "Wipers", "Tyres", "Batteries", "Lubricants", "Auto Accessories", "Tools", "Motorcycle Parts"],
    "FURNITURE": ["Sofas", "Chairs", "Tables", "Beds", "Wardrobes", "Cabinets", "Desks", "Office Furniture", "Dining Sets", "TV Stands", "Shelves", "Mattresses", "Outdoor Furniture", "Home Decor", "Lighting", "Mirrors", "Curtains", "Carpets", "Kitchen Furniture"],
    "GENERAL": ["General Goods", "Household", "Personal Care", "Food", "Beverages", "Electronics", "Clothing", "Accessories", "Stationery", "Cleaning", "Hardware", "Other"],
}

SMART_CATEGORY_HINTS = {
    "COMPUTERS": {"elitebook":"Laptops","thinkpad":"Laptops","latitude":"Laptops","macbook":"Laptops","laptop":"Laptops","notebook":"Laptops","desktop":"Desktop Computers","monitor":"Monitors","printer":"Printers","laserjet":"Printers","keyboard":"Keyboards","mouse":"Mouse","ssd":"SSD","hdd":"Hard Drives","ram":"RAM","flash":"Flash Disks","router":"Routers","switch":"Switches","cctv":"CCTV","hdmi":"Cables","charger":"Chargers"},
    "FASHION": {"air max":"Shoes","sneaker":"Shoes","shoe":"Shoes","shirt":"Shirts","t-shirt":"T-Shirts","tshirt":"T-Shirts","jeans":"Jeans","trouser":"Trousers","dress":"Dresses","skirt":"Skirts","jacket":"Jackets","sweater":"Sweaters","sandal":"Sandals","bag":"Bags","belt":"Belts","uniform":"School Uniforms","kanga":"Kanga & Vitenge","vitenge":"Kanga & Vitenge","mitumba":"Mitumba"},
    "SUPERMARKET": {"coca cola":"Soft Drinks","fanta":"Soft Drinks","sprite":"Soft Drinks","soda":"Soft Drinks","water":"Water","juice":"Juices","milk":"Milk & Dairy","bread":"Bread & Bakery","biscuit":"Biscuits","rice":"Rice","sugar":"Sugar","flour":"Flour","oil":"Cooking Oil","spice":"Spices","snack":"Snacks","detergent":"Cleaning Products"},
    "COSMETICS": {"shampoo":"Shampoo","conditioner":"Conditioner","lotion":"Lotions","perfume":"Perfumes","foundation":"Foundation","lipstick":"Lip Products","wig":"Wigs","extension":"Hair Extensions","nail":"Nail Care","serum":"Face Care"},
    "PHONES": {"iphone":"Smartphones","samsung":"Smartphones","tecno":"Smartphones","infinix":"Smartphones","xiaomi":"Smartphones","nokia":"Feature Phones","screen protector":"Screen Protectors","case":"Phone Cases","cover":"Phone Cases","charger":"Chargers","power bank":"Power Banks","earphone":"Earphones","headphone":"Headphones","memory card":"Memory Cards"},
    "STATIONERY": {"exercise":"Exercise Books","notebook":"Notebooks","pen":"Pens","pencil":"Pencils","marker":"Markers","eraser":"Erasers","ruler":"Rulers","paper":"Paper","calculator":"Calculators","textbook":"Textbooks","book":"Books"},
    "HARDWARE": {"cement":"Cement","paint":"Paint","nail":"Nails","screw":"Screws","bolt":"Bolts & Nuts","nut":"Bolts & Nuts","pipe":"Pipes","plumbing":"Plumbing","electrical":"Electrical","tile":"Tiles","lock":"Locks","hammer":"Tools","drill":"Tools","adhesive":"Adhesives"},
    "AUTOPARTS": {"brake":"Brake System","clutch":"Clutch","filter":"Filters","bearing":"Bearings","belt":"Belts","shock":"Suspension","battery":"Batteries","tyre":"Tyres","tire":"Tyres","wiper":"Wipers","headlight":"Lights","engine":"Engine Parts","oil":"Lubricants"},
    "FURNITURE": {"sofa":"Sofas","chair":"Chairs","table":"Tables","bed":"Beds","wardrobe":"Wardrobes","cabinet":"Cabinets","desk":"Desks","mattress":"Mattresses","shelf":"Shelves","carpet":"Carpets","curtain":"Curtains","mirror":"Mirrors"},
    "RESTAURANT": {"pizza":"Fast Food","burger":"Fast Food","chips":"Fast Food","fries":"Fast Food","chicken":"Chicken","fish":"Fish","beef":"Meat","nyama":"Meat","juice":"Juices","soda":"Soft Drinks","coffee":"Hot Drinks","tea":"Hot Drinks","salad":"Salads","soup":"Soups","dessert":"Desserts"},
    "PHARMACY": {"paracetamol":"Pain Relief","ibuprofen":"Pain Relief","pain":"Pain Relief","cold":"Cold & Flu","flu":"Cold & Flu","vitamin":"Vitamins","bandage":"Bandages","antiseptic":"Antiseptics","thermometer":"Medical Devices","gloves":"First Aid","sanitizer":"Antiseptics","toothpaste":"Oral Care"},
}

def get_store_profile(request, create=False):
    if not getattr(request.user, "is_authenticated", False):
        return None

    # Admin/store-owner: keep using the existing direct Duka relation.
    own_duka = getattr(request.user, "duka", None)
    if own_duka:
        set_current_duka_id(own_duka.pk)
        return own_duka

    # Staff: use a Duka linked on Profile when the model provides that field.
    profile = getattr(request.user, "profile", None)
    staff_duka = getattr(profile, "duka", None) if profile else None
    if staff_duka:
        set_current_duka_id(staff_duka.pk)
        return staff_duka

    # Single-store installations can safely auto-link legacy staff to the
    # only existing Duka. With multiple stores we do not guess.
    if profile is not None and hasattr(profile, "duka_id"):
        existing_duka_count = Duka.objects.count()
        if existing_duka_count == 1:
            staff_duka = Duka.objects.first()
            profile.duka = staff_duka
            profile.save(update_fields=["duka"])
            set_current_duka_id(staff_duka.pk)
            return staff_duka

    # Do not create a separate store for Cashier/Storekeeper accounts.
    role = get_profile_role(request.user)
    if role in {"cashier", "stoo"}:
        return None

    if create:
        duka, _ = Duka.objects.get_or_create(
            mwenye_duka=request.user,
            defaults={"business_type": "GENERAL"},
        )
        set_current_duka_id(duka.pk)
        return duka

    return None

def ensure_business_categories(business_type):
    business_type = business_type if business_type in BUSINESS_CATEGORY_PACKS else "GENERAL"
    for name in BUSINESS_CATEGORY_PACKS[business_type]:
        Kategoria.objects.get_or_create(jina=name, business_type=business_type, defaults={"is_custom":False})

def categories_for_store(duka=None):
    business_type = getattr(duka, "business_type", "GENERAL") if duka else "GENERAL"
    ensure_business_categories(business_type)
    return Kategoria.objects.filter(business_type__in=[business_type,"GENERAL"]).order_by("jina")

def smart_category_suggestion(name, business_type="GENERAL"):
    lowered=(name or "").strip().lower()
    hints=SMART_CATEGORY_HINTS.get(business_type,{})
    for keyword in sorted(hints,key=len,reverse=True):
        if keyword in lowered:
            category=Kategoria.objects.filter(business_type=business_type,jina=hints[keyword]).first()
            if category:
                return {"id":category.id,"name":category.jina,"confidence":"high"}
    return None

def _is_ajax(request):
    return request.headers.get("X-Requested-With")=="XMLHttpRequest" or "application/json" in request.headers.get("Accept","")


# =========================================================
# MACHINE-READABLE BARCODE ENGINE
# Real EAN-13 for valid 13-digit internal/factory codes; Code128 for all
# other supported values. No fake-bar fallback and no stretching/distortion.
# =========================================================

def _is_valid_ean13(value):
    value = str(value or '').strip()
    if len(value) != 13 or not value.isdigit():
        return False
    digits = [int(ch) for ch in value]
    total = sum(digits[i] if i % 2 == 0 else digits[i] * 3 for i in range(12))
    check = (10 - (total % 10)) % 10
    return check == digits[12]


def _barcode_svg(value, max_width_mm=46, height_mm=15):
    """Return a real, proportion-preserving SVG barcode."""
    value = str(value or '').strip()
    if not value:
        raise ValueError('Barcode value ni tupu.')

    target_h = float(height_mm) * mm

    if _is_valid_ean13(value):
        drawing = createBarcodeDrawing(
            'EAN13',
            value=value,
            barHeight=target_h,
            barWidth=0.31 * mm,
            humanReadable=False,
        )
    else:
        # Code128-B supports the printable ASCII product codes we use.
        if any(ord(ch) < 32 or ord(ch) > 127 for ch in value):
            raise ValueError('Barcode ina characters zisizoungwa mkono na Code 128.')

        probe = createBarcodeDrawing(
            'Code128',
            value=value,
            barHeight=target_h,
            barWidth=0.33 * mm,
            humanReadable=False,
        )
        width_mm = float(probe.width) / mm
        if width_mm > max_width_mm:
            adaptive_bar_width = (0.33 * mm) * (45.5 / width_mm)
            adaptive_bar_width = max(0.20 * mm, adaptive_bar_width)
        else:
            adaptive_bar_width = 0.33 * mm
        drawing = createBarcodeDrawing(
            'Code128',
            value=value,
            barHeight=target_h,
            barWidth=adaptive_bar_width,
            humanReadable=False,
        )

    raw = renderSVG.drawToString(drawing)
    # Remove XML/doctype wrappers; the SVG fragment is embedded directly in HTML.
    raw = re.sub(r'^<\?xml[^>]*>\s*', '', raw, flags=re.I)
    raw = re.sub(r'^<!DOCTYPE[^>]*>\s*', '', raw, flags=re.I)
    # Give the SVG an opaque white canvas before the barcode group is drawn.
    raw = re.sub(r'(<svg\b[^>]*>)', r'\1<rect width="100%" height="100%" fill="white"/>', raw, count=1, flags=re.I)
    # Never use preserveAspectRatio="none": that stretches bars and can break scanning.
    raw = raw.replace('preserveAspectRatio="xMinYMin meet"', 'preserveAspectRatio="xMidYMid meet"')
    raw = raw.replace('preserveAspectRatio="xMinYMin meet"', 'preserveAspectRatio="xMidYMid meet"')
    return raw


def barcode_svg_data_uri(barcode_value):
    try:
        raw = _barcode_svg(str(barcode_value or '')).encode('utf-8')
        return 'data:image/svg+xml;base64,' + base64.b64encode(raw).decode('ascii')
    except Exception:
        logger.exception('Machine-readable barcode rendering failed')
        return ''

def _get_import_queue(request):
    return request.session.get("tradecore_import_completion_queue",[]) or []

def _save_import_queue(request,queue):
    request.session["tradecore_import_completion_queue"]=queue; request.session.modified=True

def _remove_import_queue_item(request,product_id):
    queue=[x for x in _get_import_queue(request) if int(x.get("id"))!=int(product_id)]
    _save_import_queue(request,queue); return queue

def _category_allowed_for_store(category_id,duka):
    if not category_id: return None
    business_type=getattr(duka,"business_type","GENERAL") if duka else "GENERAL"
    return Kategoria.objects.filter(pk=category_id,business_type__in=[business_type,"GENERAL"]).first()


@login_required(login_url="login")
def inventory_view(request):
    denied = require_roles(request, 'admin', 'stoo')
    if denied:
        return denied

    duka = get_store_profile(request)
    store_location, shop_location = ensure_inventory_locations(duka)

    if request.method == "GET" and request.GET.get("inventory_alerts") == "1":
        if duka and getattr(duka, "inventory_location_mode", True) and store_location:
            rows = []
            product_ids = list(Bidhaa.objects.filter(duka=duka).values_list("id", flat=True))
            store_map = dict(ProductStock.objects.filter(location=store_location, bidhaa_id__in=product_ids).values_list("bidhaa_id", "quantity"))
            shop_map = dict(ProductStock.objects.filter(location=shop_location, bidhaa_id__in=product_ids).values_list("bidhaa_id", "quantity"))
            transit_map = {}
            pending_transfers = StockTransfer.objects.filter(duka=duka, completed_at__isnull=True).prefetch_related("items")
            for transfer in pending_transfers:
                for item in transfer.items.all():
                    transit_map[item.bidhaa_id] = transit_map.get(item.bidhaa_id, 0) + int(item.quantity or 0)
            products = Bidhaa.objects.select_related("kategoria").filter(id__in=set(store_map) | set(shop_map) | set(transit_map)).order_by("jina_la_bidhaa")[:1000]
            for p in products:
                store_qty = int(store_map.get(p.id, 0) or 0)
                shop_qty = int(shop_map.get(p.id, 0) or 0)
                transit_qty = int(transit_map.get(p.id, 0) or 0)
                total_qty = store_qty + shop_qty + transit_qty
                rows.append({"id": p.id, "name": p.jina_la_bidhaa, "stock": total_qty, "store_stock": store_qty, "shop_stock": shop_qty, "transit_stock": transit_qty, "category": p.kategoria.jina if p.kategoria else "Kawaida"})
            low = [x for x in rows if 0 < x["stock"] <= 5]
            out = [x for x in rows if x["stock"] == 0]
        else:
            base_qs = Bidhaa.objects.select_related("kategoria").all()
            low = [{"id": p.id, "name": p.jina_la_bidhaa, "stock": int(p.idadi_stoo or 0), "category": p.kategoria.jina if p.kategoria else "Kawaida"} for p in base_qs.filter(idadi_stoo__gt=0, idadi_stoo__lte=5)[:100]]
            out = [{"id": p.id, "name": p.jina_la_bidhaa, "stock": 0, "category": p.kategoria.jina if p.kategoria else "Kawaida"} for p in base_qs.filter(idadi_stoo=0)[:100]]
        return JsonResponse({"ok": True, "low_stock": low, "out_of_stock": out, "generated_at": timezone.now().isoformat()}, json_dumps_params={"ensure_ascii": False})

    action = (request.POST.get("action") or "").strip()

    if request.method == "POST" and action == "set_business_type":
        value = (request.POST.get("business_type") or "GENERAL").strip().upper()
        valid = {key for key, _ in BUSINESS_TYPE_CHOICES}
        if value not in valid:
            value = "GENERAL"
        duka = get_store_profile(request, create=True)
        duka.business_type = value
        duka.save(update_fields=["business_type"])
        ensure_business_categories(value)
        return JsonResponse({"ok": True, "business_type": value}) if _is_ajax(request) else redirect("stoo_bidhaa")

    if request.method == "POST" and action == "add_category":
        name = (request.POST.get("name") or "").strip()
        business_type = getattr(duka, "business_type", "GENERAL") if duka else "GENERAL"
        if not name:
            return JsonResponse({"ok": False, "message": "Andika jina la category."}, status=400)
        category, created = Kategoria.objects.get_or_create(jina=name, business_type=business_type, defaults={"is_custom": True})
        return JsonResponse({"ok": True, "id": category.id, "name": category.jina, "created": created}) if _is_ajax(request) else redirect("stoo_bidhaa")

    if request.method == "GET" and request.GET.get("smart_category"):
        business_type = getattr(duka, "business_type", "GENERAL") if duka else "GENERAL"
        ensure_business_categories(business_type)
        return JsonResponse({"ok": True, "suggestion": smart_category_suggestion(request.GET.get("smart_category"), business_type)})

    bidhaa_zote = Bidhaa.objects.select_related("kategoria").filter(duka=duka).order_by("jina_la_bidhaa")
    if duka and getattr(duka, "inventory_location_mode", True) and store_location:
        product_ids = [p.id for p in bidhaa_zote]
        store_map = dict(ProductStock.objects.filter(location=store_location, bidhaa_id__in=product_ids).values_list("bidhaa_id", "quantity"))
        shop_map = dict(ProductStock.objects.filter(location=shop_location, bidhaa_id__in=product_ids).values_list("bidhaa_id", "quantity"))
        transit_map = {}
        pending_transfers = StockTransfer.objects.filter(duka=duka, completed_at__isnull=True).prefetch_related("items")
        for transfer in pending_transfers:
            for item in transfer.items.all():
                transit_map[item.bidhaa_id] = transit_map.get(item.bidhaa_id, 0) + int(item.quantity or 0)
        for p in bidhaa_zote:
            p.inventory_store_stock = int(store_map.get(p.id, 0) or 0)
            p.inventory_shop_stock = int(shop_map.get(p.id, 0) or 0)
            p.inventory_transit_stock = int(transit_map.get(p.id, 0) or 0)
            p.inventory_total_stock = p.inventory_store_stock + p.inventory_shop_stock + p.inventory_transit_stock
            p.inventory_stock_state = "out" if p.inventory_total_stock == 0 else ("low" if p.inventory_total_stock <= 5 else "in")
    else:
        for p in bidhaa_zote:
            p.inventory_store_stock = int(p.idadi_stoo or 0)
            p.inventory_stock_state = "out" if p.inventory_store_stock == 0 else ("low" if p.inventory_store_stock <= 5 else "in")

    # Existing import/completion actions remain available.
    if request.method == "POST" and action == "complete_imported_product":
        product_id = request.POST.get("product_id")
        bidhaa = get_object_or_404(Bidhaa, pk=product_id)
        queue_item = next((x for x in _get_import_queue(request) if int(x.get("id")) == int(product_id)), None)
        missing = queue_item.get("fields", []) if queue_item else []
        errors, updated = [], []
        if "category" in missing:
            category = _category_allowed_for_store(request.POST.get("kategoria_id"), duka)
            if not category: errors.append("Chagua category.")
            else: bidhaa.kategoria = category; updated.append("kategoria")
        if "buying" in missing:
            value = parse_decimal((request.POST.get("bei_ya_kununulia") or "").strip(), None)
            if value is None or value < 0: errors.append("Weka bei ya kununulia.")
            else: bidhaa.bei_ya_kununulia = value; updated.append("bei_ya_kununulia")
        if "selling" in missing:
            value = parse_decimal((request.POST.get("bei_ya_kuuzia") or "").strip(), None)
            if value is None or value < 0: errors.append("Weka bei ya kuuzia.")
            else: bidhaa.bei_ya_kuuzia = value; updated.append("bei_ya_kuuzia")
        if "stock" in missing:
            try:
                value = int(request.POST.get("idadi"))
            except (ValueError, TypeError):
                value = -1
            if value < 0:
                errors.append("Weka stock sahihi.")
            elif duka and getattr(duka, "inventory_location_mode", False):
                store_location, _ = ensure_inventory_locations(duka)
                current = current_quantity(bidhaa, store_location) if store_location else int(bidhaa.idadi_stoo or 0)
                if value != current and store_location:
                    direction = "ADD" if value > current else "REMOVE"
                    perform_adjustments(
                        duka=duka,
                        location=store_location,
                        items=[{"id": bidhaa.id, "qty": abs(value - current), "direction": direction, "reason": "Excel import / completion"}],
                        actor=request.user,
                        notes="Stock imekamilishwa kupitia Import Excel",
                    )
            else:
                bidhaa.idadi_stoo = value
                updated.append("idadi_stoo")
        if errors: return JsonResponse({"ok": False, "errors": errors}, status=400)
        if updated: bidhaa.save(update_fields=sorted(set(updated)))
        queue = _remove_import_queue_item(request, bidhaa.id)
        return JsonResponse({"ok": True, "remaining": len(queue), "product_name": bidhaa.jina_la_bidhaa})

    if request.method == "POST" and action == "clear_import_completion":
        _save_import_queue(request, [])
        return JsonResponse({"ok": True}) if _is_ajax(request) else redirect("stoo_bidhaa")

    jumla_bidhaa = bidhaa_zote.count()
    if duka and getattr(duka, "inventory_location_mode", True) and store_location:
        stock_values = ProductStock.objects.filter(location=store_location)
        store_total_units = sum(int(getattr(p, "inventory_store_stock", 0) or 0) for p in bidhaa_zote)
        low_stock = sum(1 for p in bidhaa_zote if int(getattr(p, "inventory_total_stock", 0) or 0) > 0 and int(getattr(p, "inventory_total_stock", 0) or 0) <= 5)
        out_of_stock = sum(1 for p in bidhaa_zote if int(getattr(p, "inventory_total_stock", 0) or 0) == 0)
        thamani_stoo = Decimal("0.00")
        for p in bidhaa_zote:
            # SURGERY FIX: Mahesabu yanasoma mzigo uliopo 'Store' pekee (bila kujumlisha wa 'Duka')
            thamani_stoo += (p.bei_ya_kuuzia or Decimal("0")) * int(getattr(p, "inventory_store_stock", 0) or 0)
    else:
        store_total_units = bidhaa_zote.aggregate(v=Sum("idadi_stoo"))["v"] or 0
        thamani_stoo = (bidhaa_zote.aggregate(total=Sum(ExpressionWrapper(F("bei_ya_kuuzia")*F("idadi_stoo"), output_field=DecimalField(max_digits=14, decimal_places=2))))['total'] or Decimal("0.00"))
        low_stock = bidhaa_zote.filter(idadi_stoo__lte=5, idadi_stoo__gt=0).count()
        out_of_stock = bidhaa_zote.filter(idadi_stoo=0).count()

    current_business_type = getattr(duka, "business_type", "GENERAL") if duka else "GENERAL"
    ensure_business_categories(current_business_type)
    import_queue = _get_import_queue(request)
    context = {
        "title": "Stoo na Bidhaa",
        "bidhaa_zote": bidhaa_zote,
        "makategoria": categories_for_store(duka),
        "jumla_bidhaa": jumla_bidhaa,
        "thamani_stoo": thamani_stoo,
        "store_total_units": int(store_total_units or 0),
        "low_stock": low_stock,
        "out_of_stock": out_of_stock,
        "business_type_choices": BUSINESS_TYPE_CHOICES,
        "current_business_type": current_business_type,
        "current_business_type_label": dict(BUSINESS_TYPE_CHOICES).get(current_business_type, "General Retail"),
        "import_completion_json": json.dumps(import_queue),
        "inventory_location_mode": bool(duka and getattr(duka, "inventory_location_mode", False)),
        "store_location_id": store_location.id if store_location else "",
        "shop_location_id": shop_location.id if shop_location else "",
    }
    return render(request, "inventory.html", context)


@login_required(login_url="login")
def inventory_scan_api(request):
    denied = require_roles(request, 'admin', 'stoo')
    if denied: return denied
    duka = get_store_profile(request)
    store, shop = ensure_inventory_locations(duka)
    q = (request.GET.get("q") or "").strip()
    if not q: return JsonResponse({"ok": False, "message": "Barcode haijawekwa."}, status=400)
    qs = Bidhaa.objects.select_related("kategoria").filter(duka=duka)
    # Barcode/QR/name must win over numeric product-ID matches. A short
    # numeric barcode can otherwise collide with another product's integer ID.
    bidhaa = (
        qs.filter(
            Q(barcode__iexact=q)
            | Q(qr_code__iexact=q)
            | Q(jina_la_bidhaa__iexact=q)
        )
        .order_by("id")
        .first()
    )
    if bidhaa is None and q.isdigit():
        bidhaa = qs.filter(id=int(q)).first()
    if not bidhaa:
        return JsonResponse({"ok": False, "message": "Bidhaa haijapatikana.", "barcode": q}, status=404)
    store_stock = current_quantity(bidhaa, store) if store else int(bidhaa.idadi_stoo or 0)
    shop_stock = current_quantity(bidhaa, shop) if shop else 0
    return JsonResponse({"ok": True, "product": {"id": bidhaa.id, "name": bidhaa.jina_la_bidhaa, "barcode": bidhaa.barcode or "", "category": bidhaa.kategoria.jina if bidhaa.kategoria else "Kawaida", "buying": str(bidhaa.bei_ya_kununulia), "selling": str(bidhaa.bei_ya_kuuzia), "store_stock": store_stock, "shop_stock": shop_stock, "total": store_stock+shop_stock}})


@login_required(login_url="login")
def inventory_product_detail(request, product_id):
    denied = require_roles(request, 'admin', 'stoo')
    if denied: return denied
    duka = get_store_profile(request)
    store, shop = ensure_inventory_locations(duka)
    bidhaa = get_object_or_404(Bidhaa.objects.select_related("kategoria"), pk=product_id, duka=duka)
    locations = []
    if store: locations.append(store)
    if shop: locations.append(shop)
    for loc in StockLocation.objects.filter(duka=duka, is_active=True).exclude(id__in=[x.id for x in locations]).order_by("jina") if duka else []: locations.append(loc)
    stock = [{"id": l.id, "name": l.jina, "quantity": current_quantity(bidhaa, l)} for l in locations]
    verified = []
    for l in locations:
        lv = get_last_verified(bidhaa, l)
        verified.append({"location": l.jina, "quantity": lv["quantity"] if lv else None, "date": lv["date"].isoformat() if lv else None, "reference": lv["reference"] if lv else None, "verified_by": lv["verified_by"] if lv else None})
    movements = []
    qs = StockMovement.objects.filter(duka=duka, bidhaa=bidhaa).select_related("actor", "from_location", "to_location", "location").order_by("-created_at")[:80]
    for m in qs:
        local = timezone.localtime(m.created_at)
        if m.movement_type == "TRANSFER":
            label = f"Transfer {m.from_location.jina if m.from_location else '-'} → {m.to_location.jina if m.to_location else '-'}"
            detail = f"{m.quantity} · {m.reason or ''}"
        else:
            label = dict(StockMovement.TYPE_CHOICES).get(m.movement_type, m.movement_type)
            detail = f"{m.quantity} · {m.location.jina if m.location else ''}"
        movements.append({"type": m.movement_type, "label": label, "detail": detail, "reference": m.reference, "actor": m.actor.username if m.actor else "Mfumo", "time": local.strftime("%d %b %Y · %H:%M"), "reason": m.reason or ""})
    return JsonResponse({"ok": True, "product": {"id": bidhaa.id, "name": bidhaa.jina_la_bidhaa, "barcode": bidhaa.barcode or "", "category": bidhaa.kategoria.jina if bidhaa.kategoria else "Kawaida", "buying": str(bidhaa.bei_ya_kununulia), "selling": str(bidhaa.bei_ya_kuuzia), "stock": stock, "verified": verified, "movements": movements}})


@login_required(login_url="login")
def inventory_stock_in(request):
    denied = require_roles(request, 'admin', 'stoo')
    if denied: return denied
    if request.method != "POST": return JsonResponse({"ok": False, "message": "POST only."}, status=405)
    duka = get_store_profile(request)
    store, _ = ensure_inventory_locations(duka)
    if not duka or not store: return JsonResponse({"ok": False, "message": "Inventory locations hazijapatikana."}, status=400)
    try:
        items = json.loads(request.POST.get("items") or "[]")
        supplier = (request.POST.get("supplier") or "").strip()
        notes = (request.POST.get("notes") or "").strip()
        ref = (request.POST.get("reference") or "").strip() or f"GRN-{timezone.now():%Y%m%d%H%M%S}-{uuid.uuid4().hex[:6].upper()}"
        if not supplier: raise ValueError("Weka supplier.")
        if not isinstance(items, list) or not items: raise ValueError("Hakuna bidhaa za Stock In.")
        with transaction.atomic():
            results=[]
            for item in items:
                bidhaa=get_object_or_404(Bidhaa,pk=int(item.get("id")), duka=duka)
                qty=int(item.get("qty") or 0)
                movement, after, _=add_stock(duka=duka,bidhaa=bidhaa,location=store,quantity=qty,actor=request.user,reference=ref,reason="Stock In",notes=notes,movement_type="STOCK_IN",supplier=supplier,source_ref=ref)
                results.append({"id":bidhaa.id,"name":bidhaa.jina_la_bidhaa,"quantity":after})
        try:
            create_activity(request.user, f"Stock In {ref} — {len(items)} products", "stoo")
        except Exception:
            logger.exception("Stock In activity log failed for reference=%s", ref)
        return JsonResponse({"ok":True,"reference":ref,"items":results})
    except Exception as exc:
        logger.exception("Stock In failed for user=%s",request.user.username)
        return JsonResponse({"ok":False,"message":str(exc) or "Stock In imeshindikana."},status=400)


@login_required(login_url="login")
def inventory_transfer(request):
    denied = require_roles(request, 'admin', 'stoo', 'cashier')
    if denied: return denied
    duka=get_store_profile(request); store,shop=ensure_inventory_locations(duka)

    # 1. VUTA MIZIGO INAYOSUBIRI KUPOKELEWA NA BARCODE ZAKE KWA AJILI YA POS
    if request.method == "GET" and request.GET.get("pending") == "1":
        pending = (
            StockTransfer.objects
            .filter(
                duka=duka,
                completed_at__isnull=True,
                destination_location__is_sales_location=True,
            )
            .select_related("performed_by", "source_location", "destination_location")
            .prefetch_related("items__bidhaa")
            .order_by("-created_at")
        )
        data = []
        for t in pending:
            items = [{"id": i.bidhaa.id, "name": i.bidhaa.jina_la_bidhaa, "qty": i.quantity, "barcode": i.bidhaa.barcode or ""} for i in t.items.all()]
            sender = t.performed_by.get_full_name() or t.performed_by.username if t.performed_by else "Admin"
            data.append({"id": t.id, "reference": t.reference, "from": t.source_location.jina if t.source_location else "", "to": t.destination_location.jina if t.destination_location else "", "sender": sender, "date": timezone.localtime(t.created_at).strftime("%d %b %Y %H:%M"), "items": items})
        return JsonResponse({"ok": True, "transfers": data})

    if request.method != "POST": return JsonResponse({"ok": False, "message": "POST only."}, status=405)

    action = request.POST.get("action")
    
    # 2. KUPOKEA MZIGO KUTOKA POS
    if action == "receive":
        try:
            transfer_id = request.POST.get("transfer_id")
            transfer_guard = StockTransfer.objects.filter(
                pk=transfer_id, duka=duka, completed_at__isnull=True
            ).first()
            if not transfer_guard:
                raise ValueError("Mzigo huu si wa biashara hii au tayari umeshashughulikiwa.")
            from .inventory_services import receive_transfer
            transfer = receive_transfer(transfer_id, request.user)
            received_items = []
            received_units = 0
            stock_verified = True
            for item in transfer.items.select_related("bidhaa").all():
                qty = int(item.quantity or 0)
                before = int(item.destination_before or 0)
                after = int(item.destination_after or 0)
                received_units += qty
                stock_verified = stock_verified and (after == before + qty)
                received_items.append({
                    "id": item.bidhaa_id,
                    "name": item.bidhaa.jina_la_bidhaa,
                    "quantity": qty,
                    "shop_before": before,
                    "shop_after": after,
                })
            create_activity(request.user, f"Amepokea mzigo dukani: {transfer.reference}", "stoo")
            return JsonResponse({
                "ok": True,
                "message": "Mzigo umepokelewa kikamilifu!",
                "reference": transfer.reference,
                "received_units": received_units,
                "items": received_items,
                "stock_verified": stock_verified,
                "destination": transfer.destination_location.jina,
            })
        except Exception as e:
            return JsonResponse({"ok": False, "message": str(e)}, status=400)

    if action == "reject":
        try:
            transfer_id = request.POST.get("transfer_id")
            reason = (request.POST.get("reason") or "").strip()
            transfer_guard = StockTransfer.objects.filter(
                pk=transfer_id, duka=duka, completed_at__isnull=True
            ).first()
            if not transfer_guard:
                raise ValueError("Mzigo huu si wa biashara hii au tayari umeshashughulikiwa.")
            from .inventory_services import reject_transfer
            transfer, restored = reject_transfer(transfer_id, request.user, reason)
            create_activity(request.user, f"Amekataa mzigo: {transfer.reference} — {reason}", "stoo")
            return JsonResponse({"ok": True, "message": "Mzigo umekataliwa na stock imerudishwa Store.", "reference": transfer.reference, "restored": restored})
        except Exception as e:
            return JsonResponse({"ok": False, "message": str(e)}, status=400)

    # 3. KUTUMA MZIGO (TRANSFER YA KAWAIDA)
    try:
        to_id=int(request.POST.get("to_location_id")); from_id=int(request.POST.get("from_location_id") or store.id)
        source=StockLocation.objects.get(pk=from_id,duka=duka,is_active=True); destination=StockLocation.objects.get(pk=to_id,duka=duka,is_active=True)
        items=json.loads(request.POST.get("items") or "[]")
        reason_code=(request.POST.get("reason_code") or "OTHER").strip()
        if reason_code in {"RETURN_TO_STORE", "DAMAGED", "EXPIRED", "CUSTOMER_RETURN"} and source.aina == "STORE" and destination.aina == "SHOP":
            source, destination = destination, source
        reason_text=(request.POST.get("reason_text") or "").strip()
        notes=(request.POST.get("notes") or "").strip()
        from .inventory_services import perform_transfer
        transfer=perform_transfer(duka=duka,source=source,destination=destination,items=items,actor=request.user,reason_code=reason_code,reason_text=reason_text,notes=notes)
        total=sum(int(i.get("qty") or 0) for i in items)
        try:
            create_activity(request.user, f"Transfer {transfer.reference}: {source.jina} → {destination.jina} (Inasubiri Kupokelewa)", "stoo")
        except Exception:
            logger.exception("Transfer activity log failed for reference=%s", transfer.reference)
        return JsonResponse({"ok":True,"reference":transfer.reference,"from":source.jina,"to":destination.jina,"items":len(items),"units":total})
    except Exception as exc:
        logger.exception("Inventory transfer failed")
        return JsonResponse({"ok":False,"message":str(exc) or "Transfer imeshindikana."},status=400)


@login_required(login_url="login")
def inventory_adjustment(request):
    denied=require_roles(request,'admin','stoo')
    if denied:return denied
    if request.method!="POST":return JsonResponse({"ok":False,"message":"POST only."},status=405)
    duka=get_store_profile(request); store,_=ensure_inventory_locations(duka)
    try:
        items=json.loads(request.POST.get("items") or "[]")
        ref,changed=perform_adjustments(duka=duka,location=store,items=items,actor=request.user,notes=(request.POST.get("notes") or "").strip())
        try:
            create_activity(request.user, f"Stock Adjustment {ref} — {len(items)} products", "stoo")
        except Exception:
            logger.exception("Stock Adjustment activity log failed for reference=%s", ref)
        return JsonResponse({"ok":True,"reference":ref,"items":len(changed)})
    except Exception as exc:
        logger.exception("Inventory adjustment failed")
        return JsonResponse({"ok":False,"message":str(exc) or "Adjustment imeshindikana."},status=400)


@login_required(login_url="login")
def inventory_stock_take(request):
    denied=require_roles(request,'admin','stoo')
    if denied:return denied
    if request.method!="POST":return JsonResponse({"ok":False,"message":"POST only."},status=405)
    duka=get_store_profile(request); store,_=ensure_inventory_locations(duka)
    try:
        items=json.loads(request.POST.get("items") or "[]")
        notes=(request.POST.get("notes") or "").strip()
        take=perform_stock_take(duka=duka,location=store,items=items,actor=request.user,notes=notes)
        try:
            create_activity(request.user, f"Stock Take {take.reference} imekamilika — {len(items)} products", "stoo")
        except Exception:
            logger.exception("Stock Take activity log failed for reference=%s", take.reference)
        return JsonResponse({"ok":True,"reference":take.reference,"items":len(items)})
    except Exception as exc:
        logger.exception("Stock take failed")
        return JsonResponse({"ok":False,"message":str(exc) or "Stock Take imeshindikana."},status=400)


@login_required(login_url="login")
def product_search_api(request):
    denied = require_roles(request, 'admin', 'stoo', 'cashier')
    if denied:
        return denied
    q = (request.GET.get("q") or "").strip()
    category = (request.GET.get("category") or "").strip()
    stock = (request.GET.get("stock") or "ALL").upper()
    location_key = (request.GET.get("location") or "").strip().lower()
    min_price_raw = request.GET.get("min_price")
    max_price_raw = request.GET.get("max_price")
    min_price = parse_decimal(min_price_raw, None) if min_price_raw not in (None, "") else None
    max_price = parse_decimal(max_price_raw, None) if max_price_raw not in (None, "") else None

    qs = Bidhaa.objects.select_related("kategoria").all()
    if q:
        query_filter = Q(jina_la_bidhaa__icontains=q) | Q(barcode__icontains=q) | Q(qr_code__icontains=q)
        if q.isdigit():
            query_filter |= Q(id=int(q))
        qs = qs.filter(query_filter)
    if category and category != "ALL":
        qs = qs.filter(kategoria_id=category)
    if min_price is not None:
        qs = qs.filter(bei_ya_kuuzia__gte=min_price)
    if max_price is not None:
        qs = qs.filter(bei_ya_kuuzia__lte=max_price)

    duka = get_store_profile(request)
    requested_location = None
    if duka and getattr(duka, "inventory_location_mode", False) and location_key in {"store", "shop", "duka", "pos"}:
        store_location, shop_location = ensure_inventory_locations(duka)
        requested_location = shop_location if location_key in {"shop", "duka", "pos"} else store_location

    if requested_location:
        stock_subquery = ProductStock.objects.filter(
            bidhaa_id=OuterRef("pk"),
            location=requested_location,
        ).values("quantity")[:1]
        qs = qs.annotate(_location_stock=Coalesce(Subquery(stock_subquery), Value(0)))
        if stock == "IN":
            qs = qs.filter(_location_stock__gt=5)
        elif stock == "LOW":
            qs = qs.filter(_location_stock__gt=0, _location_stock__lte=5)
        elif stock == "OUT":
            qs = qs.filter(_location_stock=0)
    else:
        if stock == "IN":
            qs = qs.filter(idadi_stoo__gt=5)
        elif stock == "LOW":
            qs = qs.filter(idadi_stoo__gt=0, idadi_stoo__lte=5)
        elif stock == "OUT":
            qs = qs.filter(idadi_stoo=0)

    values_fields = [
        "id", "jina_la_bidhaa", "barcode", "kategoria__jina",
        "idadi_stoo", "bei_ya_kununulia", "bei_ya_kuuzia",
    ]
    if requested_location:
        values_fields.append("_location_stock")
    results = []
    for item in qs.values(*values_fields).order_by("jina_la_bidhaa")[:100]:
        item_stock = item.get("_location_stock", item["idadi_stoo"])
        results.append({
            "id": item["id"],
            "name": item["jina_la_bidhaa"],
            "barcode": item["barcode"] or "",
            "barcode_type": item["barcode"],
            "category": item["kategoria__jina"] if item["kategoria__jina"] else "Kawaida",
            "stock": int(item_stock or 0),
            "buying_price": str(item["bei_ya_kununulia"]),
            "selling_price": str(item["bei_ya_kuuzia"]),
        })
    return JsonResponse({"results": results, "count": len(results)})


@login_required(login_url="login")
def sajili_bidhaa(request):
    denied = require_roles(request, 'admin', 'stoo')
    if denied:
        return denied
    if request.method != "POST":
        return redirect("stoo_bidhaa")
    jina = (request.POST.get("jina_la_bidhaa") or "").strip()
    factory_barcode = (request.POST.get("barcode") or "").strip()
    barcode = factory_barcode or generate_internal_barcode()
    bei_k = parse_decimal(request.POST.get("bei_ya_kununulia"))
    bei_u = parse_decimal(request.POST.get("bei_ya_kuuzia"))
    idadi = parse_positive_int(request.POST.get("idadi"), 0)
    kategoria_id = request.POST.get("kategoria_id") or ""
    maelezo = (request.POST.get("maelekezo_maalum") or "").strip()

    if not jina or bei_k < 0 or bei_u < 0:
        messages.error(request, "Jaza taarifa sahihi za bidhaa.")
        return redirect("stoo_bidhaa")
    if Bidhaa.objects.filter(barcode=barcode).exists():
        messages.error(request, f"Barcode {barcode} tayari imesajiliwa kwenye bidhaa nyingine.")
        return redirect("stoo_bidhaa")

    duka = get_store_profile(request)
    if duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("stoo_bidhaa")
    kategoria = _category_allowed_for_store(kategoria_id, duka) if kategoria_id else None
    bidhaa = Bidhaa.objects.create(
        duka=duka,
        jina_la_bidhaa=jina,
        kategoria=kategoria,
        barcode=barcode,
        bei_ya_kununulia=bei_k,
        bei_ya_kuuzia=bei_u,
        idadi_stoo=0 if duka and getattr(duka, "inventory_location_mode", True) else idadi,
        maelekezo_maalum=maelezo or None,
    )
    if duka and getattr(duka, "inventory_location_mode", True):
        store_location, _ = ensure_inventory_locations(duka)
        if idadi > 0 and store_location:
            add_stock(duka=duka, bidhaa=bidhaa, location=store_location, quantity=idadi, actor=request.user, reference=f"OPEN-{bidhaa.id}", reason="Opening stock ya bidhaa mpya", movement_type="OPENING")
    create_activity(request.user, f"Bidhaa mpya imesajiliwa: {bidhaa.jina_la_bidhaa} [{bidhaa.barcode}]", "stoo")
    messages.success(request, f"Bidhaa '{bidhaa.jina_la_bidhaa}' imesajiliwa. Barcode: {bidhaa.barcode}")
    return redirect("stoo_bidhaa")


@login_required(login_url="login")
def edit_product(request, product_id):
    denied = require_roles(request, 'admin', 'stoo')
    if denied:
        return denied
    bidhaa = get_object_or_404(Bidhaa, pk=product_id)
    if request.method != "POST":
        return redirect("stoo_bidhaa")
    jina = (request.POST.get("jina_la_bidhaa") or "").strip()
    barcode = (request.POST.get("barcode") or bidhaa.barcode or "").strip()
    if not jina or not barcode:
        messages.error(request, "Jina na barcode ni lazima.")
        return redirect("stoo_bidhaa")
    if Bidhaa.objects.filter(barcode=barcode).exclude(pk=bidhaa.pk).exists():
        messages.error(request, f"Barcode {barcode} tayari ipo kwenye bidhaa nyingine.")
        return redirect("stoo_bidhaa")
    duka = get_store_profile(request)
    category_id = request.POST.get("kategoria_id")
    selected_category = _category_allowed_for_store(category_id, duka) if category_id else None
    if category_id and not selected_category:
        messages.error(request, "Category uliyochagua si ya biashara hii.")
        return redirect("stoo_bidhaa")
    bidhaa.jina_la_bidhaa = jina
    bidhaa.barcode = barcode
    bidhaa.bei_ya_kununulia = parse_decimal(request.POST.get("bei_ya_kununulia"), bidhaa.bei_ya_kununulia)
    bidhaa.bei_ya_kuuzia = parse_decimal(request.POST.get("bei_ya_kuuzia"), bidhaa.bei_ya_kuuzia)
    requested_stock = parse_nonnegative_int(request.POST.get("idadi"), int(bidhaa.idadi_stoo or 0))
    if category_id:
        bidhaa.kategoria = selected_category
    bidhaa.maelekezo_maalum = (request.POST.get("maelekezo_maalum") or "").strip() or None
    if duka and getattr(duka, "inventory_location_mode", True):
        store_location, _ = ensure_inventory_locations(duka)
        current = current_quantity(bidhaa, store_location) if store_location else int(bidhaa.idadi_stoo or 0)
        if requested_stock != current and store_location:
            direction = "ADD" if requested_stock > current else "REMOVE"
            qty = abs(requested_stock-current)
            perform_adjustments(duka=duka, location=store_location, items=[{"id": bidhaa.id, "qty": qty, "direction": direction, "reason": "Edit ya stock"}], actor=request.user, notes="Stock imebadilishwa kupitia Edit Product")
    else:
        bidhaa.idadi_stoo = requested_stock
    bidhaa.save()
    create_activity(request.user, f"Bidhaa imehaririwa: {bidhaa.jina_la_bidhaa}", "stoo")
    messages.success(request, "Taarifa za bidhaa zimesasishwa.")
    return redirect("stoo_bidhaa")


@login_required(login_url="login")
def delete_product(request, product_id):
    denied = require_roles(request, 'admin')
    ajax = _is_ajax(request)
    if denied:
        if ajax:
            return JsonResponse({
                "ok": False,
                "code": "forbidden",
                "message": "Huna ruhusa ya kufanya operation hii.",
            }, status=403)
        return denied

    if request.method != "POST":
        if ajax:
            return JsonResponse({
                "ok": False,
                "code": "method_not_allowed",
                "message": "Tumia POST kufanya operation hii.",
            }, status=405)
        return redirect("stoo_bidhaa")

    bidhaa = get_object_or_404(Bidhaa, pk=product_id)
    jina = bidhaa.jina_la_bidhaa
    mode = (request.POST.get("mode") or "delete").strip().lower()

    if mode in {"deplete_stock", "stock_zero", "zero_stock"}:
        duka = get_store_profile(request)
        location_mode = bool(duka and getattr(duka, "inventory_location_mode", False))
        old_stock = 0
        try:
            if location_mode:
                with transaction.atomic():
                    store_location, shop_location = ensure_inventory_locations(duka)
                    active_locations = [x for x in (store_location, shop_location) if x and x.is_active]
                    stock_rows = ProductStock.objects.filter(
                        bidhaa=bidhaa,
                        location__in=active_locations,
                    ).select_related("location")
                    old_by_location = {
                        row.location_id: int(row.quantity or 0)
                        for row in stock_rows
                    }
                    old_stock = sum(old_by_location.values())
                    if old_stock > 0:
                        for location in active_locations:
                            current = old_by_location.get(location.id, 0)
                            if current > 0:
                                perform_adjustments(
                                    duka=duka,
                                    location=location,
                                    items=[{
                                        "id": bidhaa.id,
                                        "qty": current,
                                        "direction": "REMOVE",
                                        "reason": "Stock zero / bidhaa imeondolewa",
                                    }],
                                    actor=request.user,
                                    notes="Stock imewekwa 0 kupitia Product Management",
                                )
                try:
                    create_activity(
                        request.user,
                        f"Stock ya bidhaa imeondolewa: {jina} ({old_stock} -> 0)",
                        "stoo",
                    )
                except Exception:
                    logger.exception("Stock-zero activity log failed for product %s", product_id)
                message = (
                    f"Stock ya {jina} imewekwa 0."
                    if old_stock > 0
                    else f"{jina} tayari ilikuwa na stock 0."
                )
            else:
                old_stock = int(bidhaa.idadi_stoo or 0)
                if old_stock > 0:
                    bidhaa.idadi_stoo = 0
                    bidhaa.save(update_fields=["idadi_stoo"])
                    try:
                        create_activity(
                            request.user,
                            f"Stock ya bidhaa imeondolewa: {jina} ({old_stock} -> 0)",
                            "stoo",
                        )
                    except Exception:
                        logger.exception("Legacy stock-zero activity log failed for product %s", product_id)
                    message = f"Stock ya {jina} imewekwa 0."
                else:
                    message = f"{jina} tayari ilikuwa na stock 0."
        except Exception:
            logger.exception("Failed to zero stock for product %s", product_id)
            message = "Imeshindikana kubadilisha stock ya bidhaa."
            if ajax:
                return JsonResponse({
                    "ok": False,
                    "code": "stock_update_failed",
                    "id": bidhaa.id,
                    "product_name": jina,
                    "message": message,
                }, status=500)
            messages.error(request, message)
            return redirect("stoo_bidhaa")

    if mode != "delete":
        payload = {
            "ok": False,
            "code": "invalid_mode",
            "message": "Operation ya inventory haijatambuliwa.",
        }
        if ajax:
            return JsonResponse(payload, status=400)
        messages.error(request, payload["message"])
        return redirect("stoo_bidhaa")

    try:
        bidhaa.delete()
    except ProtectedError:
        reasons = []
        try:
            if Mauzo.objects.filter(bidhaa=bidhaa).exists():
                reasons.append("historia ya mauzo")
        except Exception:
            pass
        try:
            if PurchaseRequestItem.objects.filter(bidhaa=bidhaa).exists():
                reasons.append("historia ya purchase request")
        except Exception:
            pass
        reason_text = " na ".join(reasons) if reasons else "historia iliyolindwa kwenye database"
        payload = {
            "ok": False,
            "code": "protected",
            "id": bidhaa.id,
            "product_name": jina,
            "message": (
                f"{jina} haijafutwa kabisa kwa sababu ina {reason_text}. "
                "Tumia 'Weka Stock 0' ili kuiondoa kwenye stock bila kuharibu historia."
            ),
        }
        if ajax:
            return JsonResponse(payload, status=409)
        messages.error(request, payload["message"])
        return redirect("stoo_bidhaa")
    except Exception:
        logger.exception("Product deletion failed for product %s", product_id)
        payload = {
            "ok": False,
            "code": "delete_failed",
            "id": bidhaa.id,
            "product_name": jina,
            "message": "Bidhaa haijafutwa kutokana na kosa la mfumo.",
        }
        if ajax:
            return JsonResponse(payload, status=500)
        messages.error(request, payload["message"])
        return redirect("stoo_bidhaa")

    create_activity(request.user, f"Bidhaa imefutwa: {jina}", "stoo")
    payload = {
        "ok": True,
        "action": "deleted",
        "id": product_id,
        "product_name": jina,
        "message": f"Bidhaa '{jina}' imefutwa kabisa.",
    }
    if ajax:
        return JsonResponse(payload)
    messages.success(request, payload["message"])
    return redirect("stoo_bidhaa")


@login_required(login_url="login")
def regenerate_product_barcode(request, product_id):
    denied = require_roles(request, 'admin', 'stoo')
    if denied:
        return denied
    if request.method != "POST":
        return redirect("stoo_bidhaa")
    bidhaa = get_object_or_404(Bidhaa, pk=product_id)
    bidhaa.barcode = generate_internal_barcode()
    bidhaa.save(update_fields=["barcode"])
    if _is_ajax(request):
        return JsonResponse({
            "ok": True,
            "product_name": bidhaa.jina_la_bidhaa,
            "barcode_image": barcode_svg_data_uri(bidhaa.barcode),
        })
    messages.success(request, "Barcode mpya imetengenezwa.")
    return redirect("stoo_bidhaa")

@login_required(login_url="login")
def print_product_barcode(request, product_id):
    denied = require_roles(request, 'admin', 'stoo')
    if denied:
        return denied

    bidhaa = get_object_or_404(Bidhaa, pk=product_id)
    barcode_value = str(bidhaa.barcode or '').strip()

    if not barcode_value:
        return HttpResponse("Barcode haijapatikana kwa bidhaa hii.", status=400)

    try:
        barcode_svg = _barcode_svg(
            barcode_value,
            max_width_mm=72,
            height_mm=16,
        )
    except Exception:
        logger.exception(
            "Machine-readable barcode rendering failed for product %s",
            product_id,
        )
        return HttpResponse(
            "Barcode halisi imeshindikana kutengenezwa. Hakuna fake barcode itakayochapishwa.",
            status=500,
        )

    safe_name = escape(bidhaa.jina_la_bidhaa)
    category_name = escape(
        bidhaa.kategoria.jina if getattr(bidhaa, "kategoria", None) else ""
    )

    # Thermal-only Print Studio. Each sticker is a separate print page. This avoids
    # Chrome's long-page pagination (e.g. 100 labels becoming 2 A4-like pages).
    html = r'''<!doctype html>
<html lang="sw">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Barcode Print Studio — __PRODUCT_NAME_TITLE__</title>
<style>
*{box-sizing:border-box}
html,body{margin:0;padding:0}
body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",Arial,sans-serif;background:#eef0f4;color:#111827}
button,select,input{font:inherit}
button{border:0;border-radius:12px;padding:10px 15px;font-weight:800;cursor:pointer}
.topbar{position:sticky;top:0;z-index:50;background:rgba(255,255,255,.96);backdrop-filter:blur(14px);border-bottom:1px solid #e5e7eb;padding:14px 18px}
.topbar-inner{max-width:1220px;margin:auto;display:flex;align-items:center;justify-content:space-between;gap:14px}
.title{font-size:18px;font-weight:900}.subtitle{margin-top:3px;color:#6b7280;font-size:11px}
.actions{display:flex;gap:8px;flex-wrap:wrap}.btn-print{background:#4c1d95;color:#fbbf24}.btn-back{background:#f3f4f6;color:#374151}
.layout{max-width:1220px;margin:auto;padding:18px;display:grid;grid-template-columns:335px 1fr;gap:18px;align-items:start}
.panel{background:#fff;border:1px solid #e5e7eb;border-radius:22px;padding:18px;box-shadow:0 10px 30px rgba(15,23,42,.08)}
.section-title{font-size:10px;text-transform:uppercase;letter-spacing:.12em;font-weight:900;color:#6b7280;margin-bottom:8px}
.field{margin-bottom:14px}.field label{display:block;font-size:12px;font-weight:800;margin-bottom:6px}
.field select,.field input{width:100%;padding:11px 12px;border:1px solid #d1d5db;border-radius:12px;background:#fff;outline:none}
.field select:focus,.field input:focus{border-color:#7c3aed;box-shadow:0 0 0 3px rgba(124,58,237,.10)}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:8px}.custom-box{display:none;margin-top:8px;margin-bottom:14px;padding:12px;border-radius:16px;background:#faf5ff;border:1px solid #e9d5ff}
.summary{margin-top:10px;padding:12px;border-radius:16px;background:#f9fafb;border:1px solid #e5e7eb}.summary-row{display:flex;justify-content:space-between;gap:12px;padding:4px 0;font-size:11px}.summary-row strong{font-weight:900}
.fit-ok{color:#047857}.fit-warning{color:#b45309}.live-note{margin-top:12px;padding:11px 12px;border-radius:14px;background:#f5f3ff;border:1px solid #ddd6fe;color:#5b21b6;font-size:10px;line-height:1.5}
.help{color:#6b7280;font-size:10px;line-height:1.5;margin-top:12px}
.stage{min-height:calc(100vh - 100px);display:flex;justify-content:center;align-items:flex-start;padding:8px;overflow:auto}
.preview-roll{background:#d1d5db;padding:14px;border-radius:18px;box-shadow:inset 0 0 0 1px rgba(17,24,39,.08)}
.paper{position:relative;background:#fff;box-shadow:0 18px 50px rgba(15,23,42,.22);overflow:hidden}
.preview-label{position:absolute;left:7px;top:7px;z-index:5;padding:4px 8px;border-radius:999px;background:rgba(17,24,39,.86);color:#fff;font-size:8px;font-weight:900}
.label-stack{width:100%;display:flex;flex-direction:column;align-items:center}.print-page{width:100%;display:flex;justify-content:center;align-items:flex-start;flex:0 0 auto}.screen-gap{flex:0 0 auto}
.label{flex:0 0 auto;background:#fff;border:.25mm solid #111827;padding:2.7mm;display:flex;flex-direction:column;justify-content:space-between;overflow:hidden}
.product-name{font-size:9pt;font-weight:900;line-height:1.06;text-align:center;word-break:break-word;min-height:6mm;display:flex;align-items:center;justify-content:center}
.barcode-wrap{width:100%;display:flex;justify-content:center;align-items:center;min-height:10mm;overflow:hidden}.barcode-wrap svg{display:block;width:auto!important;height:auto!important;max-width:100%;max-height:14mm;shape-rendering:crispEdges}
.barcode-number,.price{display:none!important}.roll-note{margin-top:10px;font-size:10px;color:#6b7280;text-align:center}
@media(max-width:900px){.layout{grid-template-columns:1fr}}
@media print{
 html,body{margin:0!important;padding:0!important;background:#fff!important;width:__ROLL_WIDTH_DEFAULT_MM__mm!important;height:auto!important}
 .no-print{display:none!important}
 .layout{display:block!important;width:__ROLL_WIDTH_DEFAULT_MM__mm!important;max-width:none!important;margin:0!important;padding:0!important}
 .stage{display:block!important;width:__ROLL_WIDTH_DEFAULT_MM__mm!important;min-height:0!important;margin:0!important;padding:0!important;overflow:visible!important}
 .stage>div{width:__ROLL_WIDTH_DEFAULT_MM__mm!important;margin:0!important;padding:0!important}
 .preview-roll{width:__ROLL_WIDTH_DEFAULT_MM__mm!important;margin:0!important;padding:0!important;border-radius:0!important;background:#fff!important;box-shadow:none!important}
 .paper{width:__ROLL_WIDTH_DEFAULT_MM__mm!important;height:auto!important;margin:0!important;padding:0!important;background:#fff!important;box-shadow:none!important;overflow:visible!important}
 .preview-label{display:none!important}
 .label-stack{width:__ROLL_WIDTH_DEFAULT_MM__mm!important;display:block!important;margin:0!important;padding:0!important}
 .print-page{width:__ROLL_WIDTH_DEFAULT_MM__mm!important;height:__LABEL_HEIGHT_DEFAULT_MM__mm!important;min-height:__LABEL_HEIGHT_DEFAULT_MM__mm!important;display:flex!important;justify-content:center!important;align-items:flex-start!important;margin:0!important;padding:0!important;break-after:page;page-break-after:always;overflow:hidden!important}
 .print-page:last-child{break-after:auto!important;page-break-after:auto!important}
 .screen-gap{display:none!important}
 .label{width:__LABEL_WIDTH_DEFAULT_MM__mm!important;height:__LABEL_HEIGHT_DEFAULT_MM__mm!important;min-width:__LABEL_WIDTH_DEFAULT_MM__mm!important;min-height:__LABEL_HEIGHT_DEFAULT_MM__mm!important;margin:0!important;padding:2.7mm!important;box-shadow:none!important}
 .barcode-wrap svg{max-width:100%!important;height:auto!important;shape-rendering:crispEdges!important}
 @page{size:__ROLL_WIDTH_DEFAULT_MM__mm __LABEL_HEIGHT_DEFAULT_MM__mm;margin:0!important}
}
</style>
</head>
<body>
<div class="topbar no-print"><div class="topbar-inner"><div><div class="title">Barcode Print Studio</div><div class="subtitle">Thermal label roll only — hakuna A4. Bei na namba ya barcode hazichapishwi kwenye label.</div></div><div class="actions"><button class="btn-back" type="button" onclick="history.back()">← Rudi</button><button class="btn-print" type="button" onclick="window.print()">🖨 Print Labels</button></div></div></div>
<div class="layout">
<aside class="panel no-print">
<div class="section-title">Printer / Roll</div><div class="field"><label for="printerProfile">Upana wa Roll / Printer</label><select id="printerProfile"></select></div>
<div class="section-title">Label / Sticker</div><div class="field"><label for="labelProfile">Chagua Design / Size</label><select id="labelProfile"></select></div>
<div id="customBox" class="custom-box"><div class="section-title" style="color:#6d28d9">Custom Size</div><div class="grid2"><div class="field"><label for="customWidth">Label Width (mm)</label><input id="customWidth" type="number" min="10" step="0.1" value="50"></div><div class="field"><label for="customHeight">Label Height (mm)</label><input id="customHeight" type="number" min="10" step="0.1" value="30"></div></div><div class="field"><label for="customGap">Gap / Preview spacing (mm)</label><input id="customGap" type="number" min="0" step="0.1" value="2"></div><div class="field"><label for="customRollWidth">Roll / Media Width (mm)</label><input id="customRollWidth" type="number" min="20" step="0.1" value="58"></div></div>
<div class="section-title">Quantity</div><div class="field"><label for="quantity">Idadi ya Labels</label><input id="quantity" type="number" min="1" max="500" step="1" value="3"></div>
<div class="summary"><div class="summary-row"><span>Roll width</span><strong id="summaryPrinter">58 mm</strong></div><div class="summary-row"><span>Label size</span><strong id="summaryLabel">50 × 30 mm</strong></div><div class="summary-row"><span>Gap / preview</span><strong id="summaryGap">2 mm</strong></div><div class="summary-row"><span>Quantity</span><strong id="summaryQty">3</strong></div><div class="summary-row"><span>Print page</span><strong id="summaryPage">58 × 30 mm</strong></div><div class="summary-row"><span>Fit</span><strong id="summaryFit" class="fit-ok">OK</strong></div></div>
<div id="fitMessage" class="live-note">Kila label itatumwa kama print page yake kwenye thermal printer.</div>
<div class="help"><strong>__PRODUCT_NAME_HELP__</strong>__CATEGORY_HELP__<br>Barcode value imehifadhiwa database; hapa tunabadilisha presentation na print geometry tu.</div>
</aside>
<main class="stage"> <div><div class="preview-roll"><div id="paper" class="paper"><div class="preview-label">LIVE THERMAL ROLL PREVIEW</div><div id="labelStack" class="label-stack"></div></div></div><div class="roll-note no-print">Print itatoa label moja kwa moja baada ya nyingine kwenye thermal roll. Chrome itahesabu kila sticker kama page moja; hakuna A4.</div></div></main>
</div>
<style id="dynamicPrintStyle"></style>
<script>
const PRINTERS=[{id:'thermal58',name:'Thermal Roll — 58 mm',width:58},{id:'thermal80',name:'Thermal Roll — 80 mm',width:80}];
const LABEL_PRESETS=[{id:'40x20',name:'Small Barcode — 40 × 20 mm',width:40,height:20,gap:2,recommendedRoll:58},{id:'40x30',name:'Small Barcode — 40 × 30 mm',width:40,height:30,gap:2,recommendedRoll:58},{id:'50x25',name:'Barcode — 50 × 25 mm',width:50,height:25,gap:2,recommendedRoll:58},{id:'50x30',name:'Barcode — 50 × 30 mm',width:50,height:30,gap:2,recommendedRoll:58},{id:'60x30',name:'Barcode — 60 × 30 mm',width:60,height:30,gap:2,recommendedRoll:80},{id:'70x40',name:'Large Barcode — 70 × 40 mm',width:70,height:40,gap:2,recommendedRoll:80},{id:'80x50',name:'Large Label — 80 × 50 mm',width:80,height:50,gap:3,recommendedRoll:80},{id:'custom',name:'Custom — Weka vipimo',custom:true}];
const PRODUCT_NAME=__PRODUCT_NAME_JSON__;
const BARCODE_VALUE=__BARCODE_VALUE_JSON__;
const BARCODE_SVG=__BARCODE_SVG_JSON__;
const printerSelect=document.getElementById('printerProfile'),labelSelect=document.getElementById('labelProfile'),customBox=document.getElementById('customBox'),customWidth=document.getElementById('customWidth'),customHeight=document.getElementById('customHeight'),customGap=document.getElementById('customGap'),customRollWidth=document.getElementById('customRollWidth'),quantityInput=document.getElementById('quantity'),paper=document.getElementById('paper'),stack=document.getElementById('labelStack');
PRINTERS.forEach(p=>{const o=document.createElement('option');o.value=p.id;o.textContent=p.name;printerSelect.appendChild(o)});
LABEL_PRESETS.forEach(p=>{const o=document.createElement('option');o.value=p.id;o.textContent=p.name;labelSelect.appendChild(o)});
printerSelect.value='thermal58';labelSelect.value='50x30';
function getLabel(){const p=LABEL_PRESETS.find(x=>x.id===labelSelect.value);if(p&&!p.custom)return p;return {id:'custom',custom:true,width:Math.max(10,Number(customWidth.value)||50),height:Math.max(10,Number(customHeight.value)||30),gap:Math.max(0,Number(customGap.value)||0),recommendedRoll:Math.max(20,Number(customRollWidth.value)||58)}}
function getRollWidth(){return labelSelect.value==='custom'?Math.max(20,Number(customRollWidth.value)||58):(PRINTERS.find(p=>p.id===printerSelect.value)||PRINTERS[0]).width}
function getQuantity(){let q=Math.round(Number(quantityInput.value)||1);q=Math.max(1,Math.min(500,q));quantityInput.value=q;return q}
function updateDynamicPrintStyle(roll,label){document.getElementById('dynamicPrintStyle').textContent=`@media print{.paper,.label-stack{width:${roll}mm!important}.print-page{width:${roll}mm!important;height:${label.height}mm!important;min-height:${label.height}mm!important}.label{width:${label.width}mm!important;height:${label.height}mm!important}}@page{size:${roll}mm ${label.height}mm;margin:0}`}
function updateFit(roll,label){const fit=document.getElementById('summaryFit'),msg=document.getElementById('fitMessage');if(label.width<=roll){fit.textContent='OK';fit.className='fit-ok';msg.textContent='Kila label itatumwa kama print page yake kwenye thermal printer.'}else{fit.textContent='Too Wide';fit.className='fit-warning';msg.textContent='Label ni pana kuliko roll. Chagua roll pana zaidi au punguza Label Width.'}}
function renderPreview(){customBox.style.display=labelSelect.value==='custom'?'block':'none';const label=getLabel(),roll=getRollWidth(),qty=getQuantity();const totalHeight=(label.height*qty)+(label.gap*Math.max(0,qty-1));paper.style.width=roll+'mm';paper.style.height=totalHeight+'mm';stack.innerHTML='';for(let i=0;i<qty;i++){const page=document.createElement('div');page.className='print-page';page.style.width=roll+'mm';page.style.height=label.height+'mm';const el=document.createElement('div');el.className='label';el.style.width=label.width+'mm';el.style.height=label.height+'mm';const name=document.createElement('div');name.className='product-name';name.textContent=PRODUCT_NAME;const bw=document.createElement('div');bw.className='barcode-wrap';bw.innerHTML=BARCODE_SVG;el.appendChild(name);el.appendChild(bw);page.appendChild(el);stack.appendChild(page);if(i<qty-1&&label.gap>0){const gap=document.createElement('div');gap.className='screen-gap';gap.style.height=label.gap+'mm';stack.appendChild(gap)}}document.getElementById('summaryPrinter').textContent=roll.toFixed(1)+' mm';document.getElementById('summaryLabel').textContent=label.width.toFixed(1)+' × '+label.height.toFixed(1)+' mm';document.getElementById('summaryGap').textContent=label.gap.toFixed(1)+' mm';document.getElementById('summaryQty').textContent=qty;document.getElementById('summaryPage').textContent=roll.toFixed(1)+' × '+label.height.toFixed(1)+' mm';updateFit(roll,label);updateDynamicPrintStyle(roll,label)}
function syncPresetRoll(){const p=LABEL_PRESETS.find(x=>x.id===labelSelect.value);if(p&&p.recommendedRoll)printerSelect.value=p.recommendedRoll===80?'thermal80':'thermal58';renderPreview()}
printerSelect.addEventListener('change',renderPreview);labelSelect.addEventListener('change',syncPresetRoll);[customWidth,customHeight,customGap,customRollWidth,quantityInput].forEach(el=>el.addEventListener('input',renderPreview));renderPreview();
</script>
</body>
</html>'''

    replacements = {
        '__PRODUCT_NAME_TITLE__': str(bidhaa.jina_la_bidhaa).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;').replace('"','&quot;'),
        '__PRODUCT_NAME_HELP__': safe_name,
        '__CATEGORY_HELP__': (" · " + category_name) if category_name else "",
        '__PRODUCT_NAME_JSON__': json.dumps(str(bidhaa.jina_la_bidhaa), ensure_ascii=False),
        '__BARCODE_VALUE_JSON__': json.dumps(barcode_value, ensure_ascii=False),
        '__BARCODE_SVG_JSON__': json.dumps(barcode_svg, ensure_ascii=False),
        '__ROLL_WIDTH_DEFAULT_MM__': '58',
        '__LABEL_WIDTH_DEFAULT_MM__': '50',
        '__LABEL_HEIGHT_DEFAULT_MM__': '30',
    }
    for key, value in replacements.items():
        html = html.replace(key, value)

    return HttpResponse(html)


@login_required(login_url="login")
def import_products(request):
    denied = require_roles(request, 'admin', 'stoo')
    if denied:
        return denied
    if request.method != "POST":
        return redirect("stoo_bidhaa")
    upload = request.FILES.get("excel_file")
    if not upload:
        messages.error(request, "Chagua Excel file kwanza.")
        return redirect("stoo_bidhaa")
    try:
        # Current dependency set uses openpyxl, so XLSX/XLSM are the supported formats.
        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Excel haina data.")
        headers = [str(v or "").strip().lower() for v in rows[0]]
        # SURGERY FIX: Tumelegeza masharti ili Excel isome bila kikwazo
        aliases = {
            "name": ["jina_la_bidhaa", "jina", "product", "product name", "bidhaa", "item", "description", "jina la bidhaa"],
            "barcode": ["barcode", "factory barcode", "barcode ya kiwandani", "ean", "code"],
            "category": ["kategoria", "category", "aina"],
            "buying": ["bei_ya_kununulia", "buying price", "cost", "bei ya kununua", "bei"],
            "selling": ["bei_ya_kuuzia", "selling price", "price", "bei ya kuuza", "retail"],
            "qty": ["idadi", "qty", "quantity", "stock", "pcs", "mzigo", "kiasi"],
            "notes": ["maelekezo_maalum", "maelezo", "description", "notes"],
        }
        idx = {k: next((headers.index(n) for n in names if n in headers), None) for k, names in aliases.items()}
        
        # Kama hakuna headers, tunassumye Column A ni Majina na Column B ni Idadi (Fallback Mode)
        if idx["name"] is None:
            idx["name"] = 0
        if idx["qty"] is None:
            idx["qty"] = 1

        created = 0
        queue = []
        duka = get_store_profile(request, create=True)
        business_type = getattr(duka, "business_type", "GENERAL") or "GENERAL"
        ensure_business_categories(business_type)

        for row in rows[1:]:
            get = lambda key: row[idx[key]] if idx.get(key) is not None and idx[key] < len(row) else None
            name = str(get("name") or "").strip()
            if not name:
                continue

            factory = str(get("barcode") or "").strip()
            code = factory or generate_internal_barcode()
            if Bidhaa.objects.filter(barcode=code).exists():
                continue

            raw_category = str(get("category") or "").strip()
            cat = None
            if raw_category:
                cat, _ = Kategoria.objects.get_or_create(
                    jina=raw_category,
                    business_type=business_type,
                    defaults={"is_custom": True},
                )

            buying_raw = get("buying")
            selling_raw = get("selling")
            qty_raw = get("qty")
            notes_raw = get("notes")

            buying_missing = idx["buying"] is None or buying_raw in (None, "")
            selling_missing = idx["selling"] is None or selling_raw in (None, "")
            stock_missing = idx["qty"] is None or qty_raw in (None, "")
            category_missing = idx["category"] is None or raw_category == ""

            buying_value = Decimal("0.00")
            selling_value = Decimal("0.00")
            qty_value = 0
            if not buying_missing:
                buying_value = parse_decimal(buying_raw, None)
                if buying_value is None or buying_value < 0:
                    buying_missing = True
                    buying_value = Decimal("0.00")
            if not selling_missing:
                selling_value = parse_decimal(selling_raw, None)
                if selling_value is None or selling_value < 0:
                    selling_missing = True
                    selling_value = Decimal("0.00")
            if not stock_missing:
                try:
                    qty_value = int(qty_raw)
                    if qty_value < 0:
                        raise ValueError
                except (TypeError, ValueError):
                    stock_missing = True
                    qty_value = 0

            location_mode = bool(duka and getattr(duka, "inventory_location_mode", False))
            bidhaa = Bidhaa.objects.create(
                duka=duka,
                jina_la_bidhaa=name,
                barcode=code,
                kategoria=cat,
                bei_ya_kununulia=buying_value,
                bei_ya_kuuzia=selling_value,
                idadi_stoo=0 if location_mode else qty_value,
                maelekezo_maalum=str(notes_raw or "").strip() or None,
            )
            if location_mode and qty_value > 0:
                store_location, _ = ensure_inventory_locations(duka)
                add_stock(
                    duka=duka, bidhaa=bidhaa, location=store_location, quantity=qty_value,
                    actor=request.user, reference=f"IMPORT-{bidhaa.id}",
                    reason="Opening stock kupitia Import Excel", movement_type="OPENING",
                )
            created += 1

            missing_fields = []
            if category_missing:
                missing_fields.append("category")
            if buying_missing:
                missing_fields.append("buying")
            if selling_missing:
                missing_fields.append("selling")
            if stock_missing:
                missing_fields.append("stock")
            if missing_fields:
                queue.append({"id": bidhaa.id, "fields": missing_fields})

        _save_import_queue(request, queue)
        messages.success(request, f"Excel imeingizwa: bidhaa {created} zimesajiliwa.")
    except Exception as exc:
        logger.exception("Product import failed")
        messages.error(request, f"Import imeshindikana: {exc}")
    return redirect("stoo_bidhaa")


@login_required(login_url="login")
def import_stock_to_location(request):
    """Import a shipment into an existing Stock Location without creating a second product master.

    Existing products are matched by barcode first, then exact name. New products are created
    with an internal barcode when the spreadsheet does not provide one. The selected destination
    is Store by default, but Duka/POS is explicitly supported for direct-to-shop receiving.
    """
    denied = require_roles(request, 'admin', 'stoo')
    if denied:
        return denied
    ajax = _is_ajax(request)
    if request.method != "POST":
        return JsonResponse({"ok": False, "message": "POST only."}, status=405) if ajax else redirect("stoo_bidhaa")
    upload = request.FILES.get("excel_file")
    if not upload:
        messages.error(request, "Chagua Excel file kwanza.")
        return redirect("stoo_bidhaa")

    try:
        duka = get_store_profile(request)
        store, shop = ensure_inventory_locations(duka)
        if not duka or not store or not shop:
            raise ValueError("Inventory locations hazijapatikana.")
        destination_key = (request.POST.get("destination") or "store").strip().lower()
        destination = shop if destination_key in {"shop", "duka", "pos"} else store
        supplier = (request.POST.get("supplier") or "").strip()
        reference = (request.POST.get("reference") or "").strip() or f"GRN-{timezone.now():%Y%m%d%H%M%S}-{uuid.uuid4().hex[:6].upper()}"
        notes = (request.POST.get("notes") or "").strip()
        if not supplier:
            raise ValueError("Weka jina la supplier.")

        wb = load_workbook(upload, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError("Excel haina data.")
        headers = [str(v or "").strip().lower() for v in rows[0]]
        # SURGERY FIX: Tumelegeza masharti ili Excel isome bila kikwazo
        aliases = {
            "name": ["jina_la_bidhaa", "jina", "product", "product name", "bidhaa", "item", "description", "jina la bidhaa"],
            "barcode": ["barcode", "factory barcode", "barcode ya kiwandani", "ean", "code"],
            "category": ["kategoria", "category", "aina"],
            "buying": ["bei_ya_kununulia", "buying price", "cost", "bei ya kununua", "bei"],
            "selling": ["bei_ya_kuuzia", "selling price", "price", "bei ya kuuza", "retail"],
            "qty": ["idadi", "qty", "quantity", "stock", "pcs", "mzigo", "kiasi"],
            "notes": ["maelekezo_maalum", "maelezo", "description", "notes"],
        }
        idx = {k: next((headers.index(n) for n in names if n in headers), None) for k, names in aliases.items()}
        
        # Kama hakuna headers, tunassumye Column A ni Majina na Column B ni Idadi (Fallback Mode)
        if idx["name"] is None:
            idx["name"] = 0
        if idx["qty"] is None:
            idx["qty"] = 1

        created = received = skipped = 0
        issues = []
        with transaction.atomic():
            # Tunaanzia kusoma row ya 2 (ikiwa tunaamini row 1 ilikuwa headers hata kama haikutambuliwa vizuri)
            for row_no, row in enumerate(rows[1:], start=2):
                get = lambda key: row[idx[key]] if idx.get(key) is not None and idx[key] < len(row) else None
                name = str(get("name") or "").strip()
                if not name:
                    continue
                try:
                    qty = int(get("qty") or 0)
                except (TypeError, ValueError):
                    qty = 0
                if qty <= 0:
                    skipped += 1
                    issues.append(f"Row {row_no}: quantity si sahihi.")
                    continue

                barcode = str(get("barcode") or "").strip()
                raw_buying = get("buying")
                raw_selling = get("selling")
                raw_category = str(get("category") or "").strip()
                raw_notes = str(get("notes") or "").strip() or None
                product = None
                if barcode:
                    product = Bidhaa.objects.filter(duka=duka, barcode=barcode).first()
                if product is None:
                    product = Bidhaa.objects.filter(duka=duka, jina_la_bidhaa__iexact=name).first()

                if product is None:
                    buying = parse_decimal(raw_buying, Decimal("0.00")) if raw_buying not in (None, "") else Decimal("0.00")
                    selling = parse_decimal(raw_selling, Decimal("0.00")) if raw_selling not in (None, "") else Decimal("0.00")
                    category = None
                    if raw_category:
                        category, _ = Kategoria.objects.get_or_create(
                            jina=raw_category,
                            business_type=getattr(duka, "business_type", "GENERAL") or "GENERAL",
                            defaults={"is_custom": True},
                        )
                    # Location-aware stock is maintained exclusively through
                    # ProductStock/add_stock. Do not also mutate the legacy
                    # aggregate here or Store imports will be double-counted/stale.
                    product = Bidhaa.objects.create(
                        duka=duka,
                        jina_la_bidhaa=name,
                        barcode=barcode or generate_internal_barcode(),
                        kategoria=category,
                        bei_ya_kununulia=max(buying, Decimal("0.00")),
                        bei_ya_kuuzia=max(selling, Decimal("0.00")),
                        idadi_stoo=0,
                        maelekezo_maalum=raw_notes,
                    )
                    created += 1
                add_stock(
                    duka=duka, bidhaa=product, location=destination, quantity=qty, actor=request.user,
                    reference=reference, reason=("Import Mzigo → Duka" if destination.aina == "SHOP" else "Import Mzigo → Store"),
                    notes=notes, movement_type="STOCK_IN", supplier=supplier, source_ref=reference,
                )
                received += qty

        location_label = destination.jina
        message = f"Mzigo umeingizwa: {received} units → {location_label}. Bidhaa mpya: {created}."
        if skipped:
            message += f" Rows zilizorukwa: {skipped}."
        if issues:
            message += " " + " ".join(issues[:3])
        try:
            create_activity(request.user, f"Import Mzigo {reference} → {location_label} — {received} units", "stoo")
        except Exception:
            logger.exception("Stock import activity log failed for reference=%s", reference)
        messages.success(request, message)
        if ajax:
            return JsonResponse({"ok": True, "message": message, "destination": location_label, "reference": reference, "received": received, "created": created, "skipped": skipped, "issues": issues[:10]})
    except Exception as exc:
        logger.exception("Stock shipment import failed")
        message = f"Import Mzigo imeshindikana: {exc}"
        messages.error(request, message)
        if ajax:
            return JsonResponse({"ok": False, "message": message}, status=400)
    return redirect("stoo_bidhaa")


@login_required(login_url="login")
def tradecore_inventory_question(request):
    """Lightweight, data-backed TradeCore assistant for common inventory questions.

    It deliberately answers only from current database facts; unsupported questions receive
    example prompts instead of speculative answers.
    """
    denied = require_roles(request, 'admin', 'stoo', 'cashier')
    if denied:
        return denied
    q = (request.GET.get("q") or "").strip()
    if not q:
        return JsonResponse({"ok": True, "answer": "Uliza kuhusu stock, mauzo, transfers au variance."})
    duka = get_store_profile(request)
    if not duka:
        return JsonResponse({"ok": False, "message": "Biashara ya mtumiaji haijaunganishwa na Duka."}, status=400)
    store, shop = ensure_inventory_locations(duka)
    normalized = re.sub(r"\s+", " ", q.lower()).strip()

    # Find a specific product by barcode first, then by meaningful name tokens.
    product = None
    barcode_match = re.search(r"\b\d{6,}\b", q)
    if barcode_match:
        product = Bidhaa.objects.filter(barcode=barcode_match.group(0)).first()
    if product is None:
        stop = {"niko", "iko", "wapi", "zipi", "gani", "kwa", "nini", "hizi", "hivi", "stock", "mauzo", "bidhaa", "leo", "imepungua", "imebaki", "zimebaki", "zote", "transfer", "transfers", "to", "the"}
        tokens = [t for t in re.findall(r"[a-zA-Z0-9]+", normalized) if len(t) > 2 and t not in stop]
        qs = Bidhaa.objects.select_related("kategoria")
        if tokens:
            cond = Q()
            for token in tokens[:5]:
                cond |= Q(jina_la_bidhaa__icontains=token)
            product = qs.filter(cond).order_by("jina_la_bidhaa").first()

    if any(k in normalized for k in ["wapi", "iko wapi", "zimebaki", "imebaki", "stock ya"] ) and product:
        store_qty = current_quantity(product, store) if store else int(product.idadi_stoo or 0)
        shop_qty = current_quantity(product, shop) if shop else 0
        return JsonResponse({"ok": True, "intent": "stock", "answer": f"{product.jina_la_bidhaa}: Store {store_qty} pcs · Duka {shop_qty} pcs · Total {store_qty + shop_qty} pcs."})

    if any(k in normalized for k in ["mauzo", "zimeuzwa", "sold"]) and product:
        now = timezone.localtime(timezone.now())
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)
        qty = Mauzo.objects.filter(duka=duka, bidhaa=product, tarehe_ya_mauzo__gte=start).aggregate(v=Sum("idadi"))["v"] or 0
        return JsonResponse({"ok": True, "intent": "sales", "answer": f"{product.jina_la_bidhaa}: zimeuzwa {int(qty)} pcs leo."})

    if any(k in normalized for k in ["transfer", "hamisho", "zimehamishwa"]):
        now = timezone.localtime(timezone.now())
        start = now.replace(hour=0, minute=0, second=0, microsecond=0)
        transfers = StockTransfer.objects.filter(duka=duka, created_at__gte=start).select_related("source_location", "destination_location", "performed_by")
        if not transfers.exists():
            return JsonResponse({"ok": True, "intent": "transfers", "answer": "Hakuna transfer iliyorekodiwa leo."})
        parts = []
        for t in transfers[:12]:
            sender = t.performed_by.get_full_name() or t.performed_by.username if t.performed_by else "Mfumo"
            parts.append(f"{t.reference}: {t.source_location.jina} → {t.destination_location.jina} · {sender} · {timezone.localtime(t.created_at):%H:%M}")
        return JsonResponse({"ok": True, "intent": "transfers", "answer": "Transfer za leo:\n" + "\n".join(parts)})

    if any(k in normalized for k in ["low stock", "karibu kuisha", "zimekaribia", "stock ndogo"]):
        low = []
        if store:
            for ps in ProductStock.objects.filter(location=store, quantity__lte=5).select_related("bidhaa").order_by("quantity", "bidhaa__jina_la_bidhaa")[:20]:
                low.append(f"{ps.bidhaa.jina_la_bidhaa}: {ps.quantity} pcs")
        return JsonResponse({"ok": True, "intent": "low_stock", "answer": "Low stock:\n" + ("\n".join(low) if low else "Hakuna bidhaa yenye stock ≤ 5 pcs Store.")})

    if any(k in normalized for k in ["variance", "tofauti", "hesabu haifanani", "haifanani"]):
        take = StockTake.objects.filter(duka=duka).order_by("-created_at").first()
        if not take:
            return JsonResponse({"ok": True, "intent": "variance", "answer": "Bado hakuna Stock Take iliyorekodiwa."})
        items = list(take.items.select_related("bidhaa").filter(variance__isnull=False).order_by("-id")[:20])
        details = [f"{i.bidhaa.jina_la_bidhaa}: system {i.system_quantity}, physical {i.physical_quantity}, variance {i.variance:+d}" for i in items if i.variance != 0]
        return JsonResponse({"ok": True, "intent": "variance", "answer": f"Stock Take {take.reference}:\n" + ("\n".join(details) if details else "Hakuna variance kwenye Stock Take hiyo." )})

    if product and any(k in normalized for k in ["history", "historia", "movement", "kilichotokea", "kwa nini", "imepungua"]):
        movements = StockMovement.objects.filter(duka=duka, bidhaa=product).select_related("actor", "from_location", "to_location").order_by("-created_at")[:10]
        lines = []
        for m in movements:
            actor = m.actor.get_full_name() or m.actor.username if m.actor else "Mfumo"
            local = timezone.localtime(m.created_at)
            if m.movement_type == "TRANSFER":
                lines.append(f"{local:%d %b %H:%M} · {m.from_location.jina if m.from_location else '-'} → {m.to_location.jina if m.to_location else '-'} · {m.quantity} · {actor}")
            else:
                lines.append(f"{local:%d %b %H:%M} · {dict(StockMovement.TYPE_CHOICES).get(m.movement_type, m.movement_type)} · {m.quantity} · {actor}")
        return JsonResponse({"ok": True, "intent": "history", "answer": f"Historia ya {product.jina_la_bidhaa}:\n" + ("\n".join(lines) if lines else "Hakuna movement iliyorekodiwa.")})

    return JsonResponse({
        "ok": True,
        "intent": "help",
        "answer": "Ninaweza kujibu kutoka kwenye data ya TradeCore. Jaribu:\n• Pepsi iko wapi?\n• Ni Pepsi ngapi zimeuzwa leo?\n• Transfers za leo ni zipi?\n• Ni bidhaa gani zina stock ndogo?\n• Stock Take ina variance gani?\n• Nionyeshe historia ya Pepsi.",
    })


@login_required(login_url="login")
def stoo_bidhaa_view(request):
    return inventory_view(request)


# =========================================================
# 2. SEHEMU YA MAUZO (POS)
# =========================================================

@login_required(login_url="login")
def mauzo_view(request):
    denied = require_roles(request, 'admin', 'cashier')
    if denied:
        return denied

    sale_duka = get_store_profile(request)
    if sale_duka is None:
        if request.method == "POST":
            response = post_response(False, "Biashara ya mtumiaji haijapatikana.", status=400)
            if response is not None:
                return response
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("mauzo")
    sale_store_location = None
    sale_shop_location = None
    if sale_duka and getattr(sale_duka, "inventory_location_mode", True):
        sale_store_location, sale_shop_location = ensure_inventory_locations(sale_duka)

    ajax_post = (
        request.method == "POST"
        and request.headers.get("X-Requested-With") == "XMLHttpRequest"
    )

    def post_response(ok, message, *, receipt_url=None, status=200, duplicate=False):
        if ajax_post:
            payload = {"ok": bool(ok), "message": message, "duplicate": bool(duplicate)}
            if receipt_url:
                payload["receipt_url"] = receipt_url
            return JsonResponse(payload, status=status)
        return None

    # ------------------------------------------------------------------
    # AI PRODUCT IMAGE GENERATOR
    # Frontend calls this action asynchronously. It must be handled before
    # normal sale checkout so an image request can never fall through into
    # the cart_data validation branch (which previously caused POST /mauzo/
    # -> 400 on every POS visit).
    # ------------------------------------------------------------------
    if request.method == "POST" and request.POST.get("action") == "generate_product_ai_image":
        if not ajax_post:
            return JsonResponse({"ok": False, "message": "AI image request requires AJAX."}, status=400)

        api_key = (os.environ.get("OPENAI_API_KEY") or "").strip()
        if not api_key:
            return JsonResponse({
                "ok": False,
                "code": "ai_not_configured",
                "message": "AI product images are not configured. Set OPENAI_API_KEY.",
            }, status=503)

        product_id = (request.POST.get("product_id") or "").strip()
        if not product_id.isdigit():
            return JsonResponse({"ok": False, "code": "invalid_product", "message": "Invalid product."}, status=400)

        bidhaa = get_object_or_404(
            Bidhaa.objects.select_related("kategoria", "duka"),
            pk=int(product_id),
            duka=sale_duka,
        )

        # Cache-first: never call the image provider again when the product
        # already has a real image. This is the key to preventing repeat
        # generation and repeat page movement on subsequent POS visits.
        if bidhaa.picha:
            try:
                return JsonResponse({
                    "ok": True,
                    "cached": True,
                    "product_id": bidhaa.id,
                    "image_url": bidhaa.picha.url,
                })
            except Exception:
                pass

        product_name = (bidhaa.jina_la_bidhaa or "Product").strip()
        category_name = (bidhaa.kategoria.jina if bidhaa.kategoria else "General retail").strip()
        business_type = (getattr(sale_duka, "business_type", "GENERAL") or "GENERAL").strip()
        prompt = (
            "Create a clean commercial product photograph for a retail catalog. "
            f"Product: {product_name}. Category: {category_name}. Business type: {business_type}. "
            "Show one clearly identifiable product only, centered, front three-quarter view, "
            "professional studio lighting, soft neutral light background, realistic materials and proportions, "
            "subtle natural shadow, premium e-commerce photography, no people, no extra objects, "
            "no invented text, no promotional badges, no captions, square composition. "
            "Preserve recognizable physical characteristics suggested by the product name and do not add accessories."
        )

        model_name = (os.environ.get("TRADECORE_AI_IMAGE_MODEL") or "gpt-image-2.5-flare").strip()
        endpoint = (os.environ.get("TRADECORE_AI_IMAGE_URL") or "https://api.openai.com/v1/images/generations").strip()

        try:
            global _ai_image_next_request_at
            with _ai_image_rate_lock:
                now_mono = time.monotonic()
                wait_for = max(0.0, _ai_image_next_request_at - now_mono)
                if wait_for:
                    time.sleep(wait_for)
                _ai_image_next_request_at = time.monotonic() + _AI_IMAGE_MIN_INTERVAL

            provider_response = requests.post(
                endpoint,
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json",
                },
                json={
                    "model": model_name,
                    "prompt": prompt,
                    "size": "1024x1024",
                    "quality": "low",
                    "n": 1,
                    "output_format": "webp",
                    "background": "opaque",
                },
                timeout=90,
            )
        except requests.RequestException:
            logger.exception("AI product image provider request failed for product=%s", bidhaa.id)
            return JsonResponse({
                "ok": False,
                "code": "ai_provider_unreachable",
                "message": "AI image service is temporarily unavailable.",
            }, status=503)

        if not provider_response.ok:
            provider_status = provider_response.status_code
            try:
                provider_data = provider_response.json()
            except Exception:
                provider_data = {}
            provider_error = (provider_data.get("error") or {}) if isinstance(provider_data, dict) else {}
            provider_code = str(provider_error.get("code") or "").strip().lower()
            provider_type = str(provider_error.get("type") or "").strip().lower()
            provider_message = str(provider_error.get("message") or "").strip()

            logger.error(
                "AI product image provider returned HTTP %s for product=%s code=%s type=%s message=%s",
                provider_status, bidhaa.id, provider_code, provider_type, provider_message[:500],
            )

            if provider_status == 429:
                quota_like = any(token in f"{provider_code} {provider_type} {provider_message}".lower()
                                 for token in ("insufficient_quota", "quota", "billing", "credit"))
                if quota_like:
                    return JsonResponse({
                        "ok": False,
                        "code": "ai_quota_exceeded",
                        "message": "AI image quota/credits are unavailable. Add API credits and try again.",
                    }, status=503)
                retry_after = provider_response.headers.get("Retry-After")
                try:
                    retry_after_value = max(10, min(int(float(retry_after)), 300)) if retry_after else 60
                except (TypeError, ValueError):
                    retry_after_value = 60
                return JsonResponse({
                    "ok": False,
                    "code": "ai_rate_limited",
                    "message": "AI image generation is temporarily rate-limited.",
                    "retry_after": retry_after_value,
                }, status=429)

            return JsonResponse({
                "ok": False,
                "code": "ai_provider_error",
                "message": "AI image generation failed.",
            }, status=502)

        try:
            provider_data = provider_response.json()
            image_b64 = (provider_data.get("data") or [{}])[0].get("b64_json")
            if not image_b64:
                raise ValueError("Provider did not return b64_json.")
            image_bytes = base64.b64decode(image_b64, validate=True)
            if not image_bytes:
                raise ValueError("Generated image is empty.")
        except Exception:
            logger.exception("Invalid AI product image response for product=%s", bidhaa.id)
            return JsonResponse({
                "ok": False,
                "code": "ai_invalid_response",
                "message": "AI image response could not be read.",
            }, status=502)

        try:
            bidhaa.picha.save(
                f"ai_product_{bidhaa.id}_{uuid.uuid4().hex[:10]}.webp",
                ContentFile(image_bytes),
                save=True,
            )
            return JsonResponse({
                "ok": True,
                "cached": False,
                "product_id": bidhaa.id,
                "image_url": bidhaa.picha.url,
            })
        except Exception:
            logger.exception("Failed saving AI product image for product=%s", bidhaa.id)
            return JsonResponse({
                "ok": False,
                "code": "ai_save_failed",
                "message": "Generated image could not be saved.",
            }, status=500)

    # Lightweight live POS stock snapshot used after receiving a transfer.
    # The response is intentionally location-aware: Duka ProductStock is the
    # only source used by POS when inventory locations are enabled.
    if request.method == "GET" and request.GET.get("pos_stock_snapshot") == "1":
        products = []
        if sale_shop_location:
            stock_map = dict(
                ProductStock.objects
                .filter(location=sale_shop_location)
                .values_list("bidhaa_id", "quantity")
            )
            for pid, qty in stock_map.items():
                products.append({"id": pid, "stock": int(qty or 0)})
            return JsonResponse({
                "ok": True,
                "products": products,
                "reload_recommended": True,
                "location": sale_shop_location.jina,
            })
        for product in Bidhaa.objects.only("id", "idadi_stoo"):
            products.append({"id": product.id, "stock": int(product.idadi_stoo or 0)})
        return JsonResponse({"ok": True, "products": products, "reload_recommended": True, "location": "legacy"})

    if request.method == "POST":
        client_transaction_id = (
            request.POST.get("client_transaction_id") or ""
        ).strip()
        if not re.fullmatch(r"[A-Za-z0-9_-]{12,80}", client_transaction_id):
            client_transaction_id = ""

        # Idempotency guard for online retries and offline sync retries. The
        # marker is stored in the existing ActivityLog inside the same DB
        # transaction as the sale, so a lost HTTP response cannot create a
        # duplicate sale when the client retries the same transaction.
        if client_transaction_id:
            marker = f"[PWA:{client_transaction_id}]"
            existing_activity = (
                ActivityLog.objects
                .filter(
                    mhusika=request.user,
                    aina="mauzo",
                    kitendo__contains=marker,
                )
                .order_by("-pk")
                .first()
            )
            if existing_activity:
                sale_match = re.search(
                    r"\[SALE_IDS:([0-9,]+)\]",
                    existing_activity.kitendo or "",
                )
                if sale_match:
                    sale_ids = [
                        int(value)
                        for value in sale_match.group(1).split(",")
                        if value.isdigit()
                    ]
                    if sale_ids:
                        receipt_url = reverse("receipt", args=[sale_ids[-1]])
                        response = post_response(
                            True,
                            "Muamala huu ulikwishahifadhiwa.",
                            receipt_url=receipt_url,
                            status=200,
                            duplicate=True,
                        )
                        if response is not None:
                            return response
                        return redirect("receipt", sale_id=sale_ids[-1])

        cart_data_json = request.POST.get(
            "cart_data"
        )

        mteja_id = request.POST.get(
            "mteja_id"
        )

        punguzo_val = parse_decimal(
            request.POST.get("punguzo")
        )

        ongezeko_val = parse_decimal(
            request.POST.get("ongezeko")
        )

        njia_ya_malipo = (
            request.POST.get("njia_ya_malipo")
            or "CASH"
        )

        aina_ya_punguzo = (
            request.POST.get("aina_ya_punguzo")
            or "TZS"
        )

        aina_ya_ongezeko = (
            request.POST.get("aina_ya_ongezeko")
            or "TZS"
        )

        if not cart_data_json:
            response = post_response(False, "Kikapu cha mauzo kiko tupu.", status=400)
            if response is not None:
                return response
            messages.error(request, "Kikapu cha mauzo kiko tupu.")
            return redirect("mauzo")

        try:
            cart_data = json.loads(
                cart_data_json
            )
        except json.JSONDecodeError:
            response = post_response(False, "Taarifa za kikapu si sahihi.", status=400)
            if response is not None:
                return response
            messages.error(request, "Taarifa za kikapu si sahihi.")
            return redirect("mauzo")

        if not isinstance(cart_data, list) or not cart_data:
            response = post_response(False, "Hakuna bidhaa kwenye kikapu.", status=400)
            if response is not None:
                return response
            messages.error(request, "Hakuna bidhaa kwenye kikapu.")
            return redirect("mauzo")

        mteja_obj = None

        if mteja_id and str(mteja_id).strip():
            mteja_obj = (
                Mteja.objects
                .filter(pk=mteja_id)
                .first()
            )

            if not mteja_obj:
                response = post_response(False, "Mteja aliyechaguliwa hakupatikana.", status=400)
                if response is not None:
                    return response
                messages.error(request, "Mteja aliyechaguliwa hakupatikana.")
                return redirect("mauzo")

        try:
            with transaction.atomic():

                created_sales = []
                invoice_number = make_invoice_number()

                # Read every product from the database first. Never trust the
                # client-side price for financial calculations.
                prepared_items = []
                for item in cart_data:
                    bidhaa_id = item.get("id")
                    quantity = parse_positive_int(item.get("qty"))
                    if not bidhaa_id or quantity <= 0:
                        raise ValueError("Kuna bidhaa yenye taarifa zisizo sahihi.")

                    bidhaa = (
                        Bidhaa.objects
                        .select_for_update()
                        .filter(pk=bidhaa_id)
                        .first()
                    )
                    if not bidhaa:
                        raise ValueError("Bidhaa fulani haipo tena kwenye database.")
                    available = current_quantity(bidhaa, sale_shop_location) if sale_shop_location else int(bidhaa.idadi_stoo or 0)
                    if available < quantity:
                        raise ValueError(
                            f"Samahani, bidhaa '{bidhaa.jina_la_bidhaa}' "
                            f"imebaki {available} tu kwenye Duka!"
                        )

                    unit_price = bidhaa.bei_ya_kuuzia or Decimal("0.00")
                    gross_line = unit_price * quantity
                    prepared_items.append((bidhaa, quantity, gross_line))

                order_subtotal = sum(
                    (gross for _, _, gross in prepared_items),
                    Decimal("0.00"),
                )

                normalized_discount_input = max(punguzo_val, Decimal("0.00"))
                if aina_ya_punguzo == "PERCENT":
                    normalized_discount_input = min(normalized_discount_input, Decimal("100"))
                    order_discount = order_subtotal * normalized_discount_input / Decimal("100")
                else:
                    normalized_discount_input = min(normalized_discount_input, order_subtotal)
                    order_discount = normalized_discount_input

                after_discount = max(
                    order_subtotal - order_discount,
                    Decimal("0.00"),
                )

                normalized_markup_input = max(ongezeko_val, Decimal("0.00"))
                if aina_ya_ongezeko == "PERCENT":
                    normalized_markup_input = min(normalized_markup_input, Decimal("100"))
                    order_markup = after_discount * normalized_markup_input / Decimal("100")
                else:
                    order_markup = normalized_markup_input

                order_total = max(
                    after_discount + order_markup,
                    Decimal("0.00"),
                ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

                # Allocate the real net invoice total proportionally across
                # lines so stored line totals sum exactly to JUMLA KUU.
                allocated_totals = []
                remaining = order_total
                for index, (_, _, gross_line) in enumerate(prepared_items):
                    if index == len(prepared_items) - 1:
                        line_total = remaining
                    elif order_subtotal > 0:
                        line_total = (
                            order_total * gross_line / order_subtotal
                        ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                        line_total = min(max(line_total, Decimal("0.00")), remaining)
                        remaining -= line_total
                    else:
                        line_total = Decimal("0.00")
                    allocated_totals.append(line_total)

                for index, ((bidhaa, quantity, _gross_line), line_total) in enumerate(
                    zip(prepared_items, allocated_totals)
                ):
                    # Save the line with zero adjustments first. The model's
                    # create-time save then records the correct selling price
                    # and reduces stock safely.
                    sale = Mauzo(
                        bidhaa=bidhaa,
                        mteja=mteja_obj,
                        muuzaji=request.user,
                        duka=sale_duka,
                        stock_location=sale_shop_location,
                        idadi=quantity,
                        namba_ya_invoice=invoice_number,
                        njia_ya_malipo=njia_ya_malipo,
                        aina_ya_punguzo=aina_ya_punguzo,
                        punguzo=Decimal("0.00"),
                        aina_ya_ongezeko=aina_ya_ongezeko,
                        ongezeko=Decimal("0.00"),
                    )
                    sale.save()

                    # Attach the order-level adjustment only to the first line
                    # as metadata, while persisting its allocated final line
                    # amount. Existing one-line adjustment fields remain intact.
                    if index == 0:
                        sale.punguzo = normalized_discount_input
                        sale.ongezeko = normalized_markup_input
                    else:
                        sale.punguzo = Decimal("0.00")
                        sale.ongezeko = Decimal("0.00")
                    sale.aina_ya_punguzo = aina_ya_punguzo
                    sale.aina_ya_ongezeko = aina_ya_ongezeko
                    sale.jumla_pesa_iliyopokelewa = line_total
                    sale.save(
                        update_fields=[
                            "aina_ya_punguzo",
                            "punguzo",
                            "aina_ya_ongezeko",
                            "ongezeko",
                            "jumla_pesa_iliyopokelewa",
                        ]
                    )
                    created_sales.append(sale)

                sale_ids_text = ",".join(str(sale.id) for sale in created_sales)
                activity_message = (
                    f"Mauzo mapya yamekamilishwa "
                    f"({len(created_sales)} line)"
                )
                if client_transaction_id:
                    activity_message += (
                        f" [PWA:{client_transaction_id}]"
                        f" [SALE_IDS:{sale_ids_text}]"
                    )
                create_activity(
                    request.user,
                    activity_message,
                    "mauzo",
                )

                reviewers = (
                    User.objects
                    .filter(profile__duka=sale_duka, is_active=True)
                    .filter(
                        Q(is_superuser=True)
                        | Q(is_staff=True)
                    )
                    .exclude(
                        pk=request.user.pk
                    )
                    .distinct()
                )

                notifications = [
                    Notification(
                        recipient=user,
                        title="Mauzo Mapya",
                        message=(
                            f"{request.user.username} "
                            "amekamilisha mauzo mapya."
                        ),
                        notification_type="SALE",
                    )
                    for user in reviewers
                ]

                if notifications:
                    Notification.objects.bulk_create(
                        notifications
                    )

        except Exception as exc:
            logger.exception(
                "Sale creation failed for user=%s",
                request.user.username,
            )

            response = post_response(
                False,
                "Mauzo hayajakamilika. Tafadhali hakikisha taarifa za mauzo ni sahihi.",
                status=400,
            )
            if response is not None:
                return response

            messages.error(
                request,
                "Mauzo hayajakamilika. Tafadhali hakikisha taarifa za mauzo ni sahihi.",
            )

            return redirect("mauzo")

        if not ajax_post:
            messages.success(
                request,
                (
                    "Malipo yamekamilika, "
                    "risiti imeandaliwa na stoo imesasishwa!"
                ),
            )

        receipt_url = reverse("receipt", args=[created_sales[-1].id])
        response = post_response(
            True,
            "Malipo yamekamilika, risiti imeandaliwa na stoo imesasishwa!",
            receipt_url=receipt_url,
            status=201,
        )
        if response is not None:
            return response
        return redirect("receipt", sale_id=created_sales[-1].id)

    if sale_shop_location:
        available_ids = set(ProductStock.objects.filter(location=sale_shop_location, quantity__gt=0).values_list("bidhaa_id", flat=True))
        
        # SURGERY FIX: Tunavuta tu bidhaa ambazo zilishawahi kufika dukani na sasa hivi zimeisha (quantity <= 0)
        # Hii inazuia bidhaa za Store ambazo hazijawahi kuja dukani kutoonekana kabisa kwenye POS.
        sold_out_ids = set(ProductStock.objects.filter(location=sale_shop_location, quantity__lte=0).values_list("bidhaa_id", flat=True))
        
        bidhaa_zote = list(Bidhaa.objects.filter(id__in=available_ids).order_by("jina_la_bidhaa"))
        bidhaa_zilizokwisha = list(Bidhaa.objects.filter(id__in=sold_out_ids).order_by("jina_la_bidhaa"))
        
        for p in bidhaa_zote + bidhaa_zilizokwisha:
            p.pos_stock = current_quantity(p, sale_shop_location)
    else:
        bidhaa_zote = Bidhaa.objects.filter(idadi_stoo__gt=0).order_by("jina_la_bidhaa")
        bidhaa_zilizokwisha = Bidhaa.objects.filter(idadi_stoo__lte=0).order_by("jina_la_bidhaa")

    wateja_wote = (
        Mteja.objects
        .all()
        .order_by("-id")
    )

    mauzo_ya_nyuma = (
        Mauzo.objects
        .select_related(
            "bidhaa",
            "mteja",
            "muuzaji",
        )
        .order_by("-id")[:20]
    )

    context = {
        "title": "Point of Sale (POS)",
        "bidhaa_zote": bidhaa_zote,
        "bidhaa_zilizokwisha": bidhaa_zilizokwisha,
        "wateja_wote": wateja_wote,
        "mauzo_ya_nyuma": mauzo_ya_nyuma,
        "kategoria_zote": categories_for_store(get_store_profile(request)),
    }

    return render(
        request,
        "mauzo.html",
        context,
    )


# =========================================================
# 3. RISITI / RECEIPTS
# =========================================================

def _receipt_snapshot(request, sale):
    """Build one authoritative receipt dataset from the stored sale lines."""
    invoice = sale.namba_ya_invoice or f"INV-{sale.id:06d}"
    lines = list(
        Mauzo.objects.select_related("bidhaa")
        .filter(namba_ya_invoice=invoice)
        .order_by("id")
    ) or [sale]

    # IMPORTANT: receipt totals are recalculated from the gross product prices
    # and the order-level adjustment stored on the invoice's first line. This
    # prevents an adjustment from being applied once per line.
    subtotal_total = sum(
        (line.subtotal or Decimal("0.00"))
        for line in lines
    )

    adjustment_line = next(
        (
            line for line in lines
            if (line.punguzo or Decimal("0.00"))
            or (line.ongezeko or Decimal("0.00"))
        ),
        lines[0],
    )

    discount_input = max(
        adjustment_line.punguzo or Decimal("0.00"),
        Decimal("0.00"),
    )
    markup_input = max(
        adjustment_line.ongezeko or Decimal("0.00"),
        Decimal("0.00"),
    )
    discount_type = adjustment_line.aina_ya_punguzo or "TZS"
    markup_type = adjustment_line.aina_ya_ongezeko or "TZS"

    if discount_type == "PERCENT":
        discount_input = min(discount_input, Decimal("100"))
        discount_amount = subtotal_total * discount_input / Decimal("100")
    else:
        discount_amount = min(discount_input, subtotal_total)

    subtotal_after_discount = max(
        subtotal_total - discount_amount,
        Decimal("0.00"),
    )

    if markup_type == "PERCENT":
        markup_input = min(markup_input, Decimal("100"))
        markup_amount = subtotal_after_discount * markup_input / Decimal("100")
    else:
        markup_amount = markup_input

    receipt_total = max(
        subtotal_after_discount + markup_amount,
        Decimal("0.00"),
    ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

    # Build line totals proportionally from the authoritative invoice total.
    # This also repairs historical receipts created while discount/markup was
    # incorrectly applied to only the first line.
    receipt_lines = []
    remaining = receipt_total
    if subtotal_total > 0:
        for index, line in enumerate(lines):
            gross_line = line.subtotal or Decimal("0.00")
            if index == len(lines) - 1:
                line_total = remaining
            else:
                line_total = (
                    receipt_total * gross_line / subtotal_total
                ).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
                line_total = min(max(line_total, Decimal("0.00")), remaining)
                remaining -= line_total
            receipt_lines.append({"line": line, "total": line_total})
    else:
        receipt_lines = [
            {"line": line, "total": Decimal("0.00")}
            for line in lines
        ]

    duka = get_store_profile(request, create=True)
    customer = sale.mteja
    customer_name = (
        customer.majina_kamili
        if customer
        else "Walk-in Customer"
    )
    customer_phone = (
        getattr(customer, "whatsapp_no", None)
        or getattr(customer, "namba_ya_simu", "")
        or ""
    ).strip()

    location_parts = []
    if customer:
        eneo = (getattr(customer, "eneo", None) or "").strip()
        anwani = (getattr(customer, "anwani", None) or "").strip()
        if eneo:
            location_parts.append(eneo)
        if anwani and anwani.lower() != eneo.lower():
            location_parts.append(anwani)
    customer_location = " — ".join(location_parts)

    store_name = (
        getattr(duka, "jina_la_duka", None)
        or "TradeCore"
    ).strip()
    store_address = (
        getattr(duka, "anwani_au_mahali", None)
        or ""
    ).strip()
    store_phone = (getattr(duka, "simu", None) or "").strip()
    store_email = (getattr(duka, "email", None) or "").strip()
    store_tin = (getattr(duka, "tin", None) or "").strip()
    store_tagline = (getattr(duka, "tagline", None) or "").strip()
    receipt_footer = (
        getattr(duka, "receipt_footer", None)
        or "Asante kwa kufanya manunuzi nasi!"
    ).strip()
    currency = (getattr(duka, "currency", None) or "TZS").strip()

    def normalize_tz_phone(phone):
        digits = re.sub(r"\D", "", phone or "")
        if digits.startswith("0") and len(digits) >= 10:
            digits = "255" + digits[1:]
        elif digits.startswith("7") and len(digits) >= 9:
            digits = "255" + digits
        return digits

    phone_digits = normalize_tz_phone(customer_phone)
    items_text = [
        (
            f"{row['line'].bidhaa.jina_la_bidhaa} x{row['line'].idadi}"
            f" = {currency} {row['total']:,.0f}"
        )
        for row in receipt_lines
    ]

    whatsapp_lines = [
        store_name,
        "RISITI YA MAUZO",
        f"Invoice: {invoice}",
        f"Mteja: {customer_name}",
    ]
    if customer_location:
        whatsapp_lines.append(f"Mahali: {customer_location}")
    whatsapp_lines.extend(items_text)
    whatsapp_lines.append(
        f"Jumla ya bidhaa: {currency} {subtotal_total:,.0f}"
    )
    if discount_amount:
        label = (
            f"Punguzo ({discount_input:,.2f}%):"
            if discount_type == "PERCENT"
            else "Punguzo:"
        )
        whatsapp_lines.append(
            f"{label} - {currency} {discount_amount:,.0f}"
        )
    if markup_amount:
        label = (
            f"Ongezeko ({markup_input:,.2f}%):"
            if markup_type == "PERCENT"
            else "Ongezeko:"
        )
        whatsapp_lines.append(
            f"{label} + {currency} {markup_amount:,.0f}"
        )
    whatsapp_lines.extend([
        f"JUMLA KUU: {currency} {receipt_total:,.0f}",
        receipt_footer,
    ])
    whatsapp_message = "\n".join(whatsapp_lines)

    receipt_url = request.build_absolute_uri(
        reverse("receipt", args=[sale.id])
    )
    qr_payload = "\n".join([
        "TradeCore",
        store_name,
        f"Invoice: {invoice}",
        f"Mteja: {customer_name}",
        f"Jumla Kuu: {currency} {receipt_total:,.0f}",
        receipt_url,
    ])
    qr_png = BytesIO()
    qrcode.make(qr_payload).save(qr_png, format="PNG")
    qr_data_uri = (
        "data:image/png;base64,"
        + base64.b64encode(qr_png.getvalue()).decode("ascii")
    )

    return {
        "invoice": invoice,
        "lines": lines,
        "receipt_lines": receipt_lines,
        "subtotal_total": subtotal_total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        "receipt_total": receipt_total,
        "discount_input": discount_input,
        "discount_type": discount_type,
        "discount_amount": discount_amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        "markup_input": markup_input,
        "markup_type": markup_type,
        "markup_amount": markup_amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP),
        "duka": duka,
        "customer": customer,
        "customer_name": customer_name,
        "customer_phone": customer_phone,
        "customer_location": customer_location,
        "store_name": store_name,
        "store_address": store_address,
        "store_phone": store_phone,
        "store_email": store_email,
        "store_tin": store_tin,
        "store_tagline": store_tagline,
        "receipt_footer": receipt_footer,
        "currency": currency,
        "phone_digits": phone_digits,
        "whatsapp_message": whatsapp_message,
        "qr_data_uri": qr_data_uri,
        "receipt_url": receipt_url,
    }


@login_required(login_url="login")
def receipt_view(request, sale_id):
    denied = require_roles(request, "admin", "cashier")
    if denied:
        return denied

    sale = get_object_or_404(
        Mauzo.objects.select_related(
            "mteja",
            "muuzaji",
            "bidhaa",
        ),
        pk=sale_id,
    )

    # Receipt business information is editable here and persists in Duka.
    # Only admins can change it; cashiers can still view/print/download receipts.
    if request.method == "POST" and request.POST.get("action") == "save_receipt_profile":
        admin_denied = require_roles(request, "admin")
        if admin_denied:
            return admin_denied

        duka = get_store_profile(request, create=True)
        duka.jina_la_duka = (
            request.POST.get("jina_la_duka") or ""
        ).strip() or None
        duka.anwani_au_mahali = (
            request.POST.get("anwani_au_mahali") or ""
        ).strip() or None
        duka.simu = (
            request.POST.get("simu") or ""
        ).strip() or None
        duka.email = (
            request.POST.get("email") or ""
        ).strip() or None
        duka.tin = (
            request.POST.get("tin") or ""
        ).strip() or None
        duka.tagline = (
            request.POST.get("tagline") or ""
        ).strip() or None
        duka.receipt_footer = (
            request.POST.get("receipt_footer")
            or "Asante kwa kufanya manunuzi nasi!"
        ).strip()

        logo = request.FILES.get("logo")
        if logo:
            duka.logo = logo

        duka.save()
        messages.success(
            request,
            "Taarifa za risiti zimehifadhiwa kikamilifu.",
        )
        return redirect("receipt", sale_id=sale_id)

    data = _receipt_snapshot(request, sale)

    # Download uses the already-working /receipt/<id>/ URL with a query flag.
    # This avoids depending on a separate receipt_pdf URL pattern being present.
    if request.GET.get("download") == "pdf":
        return receipt_pdf(request, sale_id)

    return render(
        request,
        "receipt.html",
        {
            "title": f"Receipt {data['invoice']}",
            "sale": sale,
            "lines": data["lines"],
            "receipt_lines": data["receipt_lines"],
            "invoice": data["invoice"],
            "receipt_total": data["receipt_total"],
            "subtotal_total": data["subtotal_total"],
            "discount_input": data["discount_input"],
            "discount_type": data["discount_type"],
            "discount_amount": data["discount_amount"],
            "markup_input": data["markup_input"],
            "markup_type": data["markup_type"],
            "markup_amount": data["markup_amount"],
            "customer_phone": data["customer_phone"],
            "customer_location": data["customer_location"],
            "whatsapp_message": data["whatsapp_message"],
            "shop": data["duka"],
            "store_name": data["store_name"],
            "store_phone": data["store_phone"],
            "store_address": data["store_address"],
            "store_email": data["store_email"],
            "store_tin": data["store_tin"],
            "store_tagline": data["store_tagline"],
            "receipt_footer": data["receipt_footer"],
            "currency": data["currency"],
            "qr_data_uri": data["qr_data_uri"],
            "show_receipt_settings": request.GET.get("settings") == "1",
        },
    )


@login_required(login_url="login")
def receipt_pdf(request, sale_id):
    denied = require_roles(request, "admin", "cashier")
    if denied:
        return denied

    sale = get_object_or_404(
        Mauzo.objects.select_related("mteja", "muuzaji", "bidhaa"),
        pk=sale_id,
    )
    data = _receipt_snapshot(request, sale)

    from reportlab.lib.colors import HexColor, black, white
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import mm
    from reportlab.lib.utils import ImageReader

    # Thermal-receipt PDF: real PDF bytes generated from the actual sale/database.
    page_width = 80 * mm
    estimated_height = 152 * mm
    for row in data["receipt_lines"]:
        line = row["line"]
        name = str(line.bidhaa.jina_la_bidhaa or "")
        wrapped = max(1, (len(name) + 24) // 25)
        estimated_height += (wrapped * 4.5 + 5.5) * mm
    estimated_height += 63 * mm  # totals + thank-you + QR + footer
    page_height = max(170 * mm, min(400 * mm, estimated_height))

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="{data["invoice"]}.pdf"'
    )

    c = canvas.Canvas(
        response,
        pagesize=(page_width, page_height),
        pageCompression=1,
    )
    margin = 5.5 * mm
    center_x = page_width / 2
    right_x = page_width - margin
    y = page_height - 7.5 * mm

    purple = HexColor("#4c1d95")
    gray = HexColor("#6b7280")
    light = HexColor("#e5e7eb")

    def draw_center(text, font="Helvetica", size=8, color=black, gap=3.5 * mm):
        nonlocal y
        c.setFont(font, size)
        c.setFillColor(color)
        c.drawCentredString(center_x, y, str(text)[:120])
        y -= gap
        c.setFillColor(black)

    def draw_right(text, font="Helvetica", size=8, color=black, gap=4 * mm):
        nonlocal y
        c.setFont(font, size)
        c.setFillColor(color)
        c.drawRightString(right_x, y, str(text)[:120])
        y -= gap
        c.setFillColor(black)

    def draw_left(text, font="Helvetica", size=8, color=black, gap=4 * mm):
        nonlocal y
        c.setFont(font, size)
        c.setFillColor(color)
        c.drawString(margin, y, str(text)[:120])
        y -= gap
        c.setFillColor(black)

    # Optional logo from the configured Duka profile.
    logo = getattr(data["duka"], "logo", None)
    if logo:
        try:
            logo_buf = BytesIO(logo.read())
            logo.seek(0)
            c.drawImage(
                ImageReader(logo_buf),
                center_x - 8 * mm,
                y - 16 * mm,
                width=16 * mm,
                height=16 * mm,
                preserveAspectRatio=True,
                mask="auto",
            )
            y -= 18 * mm
        except Exception:
            pass

    draw_center("TRADECORE", "Helvetica-Bold", 9.5, purple, 4 * mm)
    draw_center(data["store_name"], "Helvetica-Bold", 13, black, 4.5 * mm)
    if data["store_tagline"]:
        draw_center(data["store_tagline"], "Helvetica", 7.2, gray, 3.5 * mm)
    if data["store_address"]:
        draw_center(data["store_address"], "Helvetica", 7.2, gray, 3.5 * mm)
    if data["store_phone"]:
        draw_center(f"Simu: {data['store_phone']}", "Helvetica", 7.2, gray, 3.5 * mm)
    if data["store_email"]:
        draw_center(data["store_email"], "Helvetica", 7.2, gray, 3.5 * mm)
    if data["store_tin"]:
        # TIN is intentionally below phone + email, matching the screen receipt.
        draw_center(f"TIN No: {data['store_tin']}", "Helvetica-Bold", 7.2, black, 5 * mm)

    c.setStrokeColor(light)
    c.line(margin, y, right_x, y)
    y -= 5 * mm

    # Normal-weight tracked title so the letters spread across the thermal receipt.
    title = "RISITI YA MAUZO"
    c.setFont("Helvetica", 12.2)
    c.setFillColor(purple)
    spacing = 1.1 * mm
    widths = [c.stringWidth(ch, "Helvetica", 12.2) for ch in title]
    total_width = sum(widths) + spacing * (len(title) - 1)
    cursor_x = center_x - (total_width / 2)
    for ch, char_width in zip(title, widths):
        c.drawString(cursor_x, y, ch)
        cursor_x += char_width + spacing
    y -= 5 * mm
    c.setFillColor(black)
    c.setFont("Helvetica-Bold", 8)
    c.setFillColor(black)
    c.drawRightString(right_x, y, f"Invoice No: {data['invoice']}")
    y -= 5 * mm

    draw_left(f"Mteja: {data['customer_name']}", "Helvetica-Bold", 8, black, 4 * mm)
    if data["customer_phone"]:
        draw_left(f"Simu: {data['customer_phone']}", "Helvetica", 7.5, gray, 4 * mm)
    if data["customer_location"]:
        draw_left(f"Mahali: {data['customer_location']}", "Helvetica", 7.5, gray, 4 * mm)
    draw_left(
        sale.tarehe_ya_mauzo.strftime("Tarehe: %d %b %Y %H:%M"),
        "Helvetica",
        7.5,
        gray,
        4.5 * mm,
    )
    draw_left(
        f"Njia ya malipo: {sale.get_njia_ya_malipo_display()}",
        "Helvetica",
        7.5,
        gray,
        5 * mm,
    )

    c.setStrokeColor(light)
    c.line(margin, y, right_x, y)
    y -= 4.5 * mm

    # Smart columns: product / qty / total.
    qty_x = page_width - 28 * mm
    c.setFont("Helvetica-Bold", 7.5)
    c.setFillColor(gray)
    c.drawString(margin, y, "Bidhaa")
    c.drawCentredString(qty_x, y, "Qty")
    c.drawRightString(right_x, y, "Jumla")
    y -= 4.5 * mm
    c.setFillColor(black)

    def wrap_chars(value, width=24):
        words = str(value or "").split()
        lines = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if current and len(candidate) > width:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
        return lines or [""]

    c.setFont("Helvetica", 7.5)
    for row in data["receipt_lines"]:
        line = row["line"]
        name_lines = wrap_chars(line.bidhaa.jina_la_bidhaa, 24)
        amount = row["total"]
        for idx, name_line in enumerate(name_lines):
            c.drawString(margin, y, name_line[:35])
            if idx == 0:
                c.drawCentredString(qty_x, y, str(line.idadi))
                c.drawRightString(right_x, y, f"{amount:,.0f}")
            y -= 4 * mm
        c.setFont("Helvetica", 6.6)
        unit = line.bei_ya_kuuzia_stoo or Decimal("0.00")
        c.setFillColor(gray)
        c.drawString(
            margin,
            y,
            f"{data['currency']} {unit:,.0f} / unit",
        )
        c.setFillColor(black)
        c.setFont("Helvetica", 7.5)
        y -= 4.3 * mm

    c.setStrokeColor(light)
    c.line(margin, y, right_x, y)
    y -= 5 * mm

    draw_right(
        f"Jumla ya bidhaa: {data['currency']} {data['subtotal_total']:,.0f}",
        "Helvetica",
        8,
        gray,
        4.2 * mm,
    )
    if data["discount_amount"]:
        discount_label = (
            f"Punguzo ({data['discount_input']:,.2f}%):"
            if data["discount_type"] == "PERCENT"
            else "Punguzo:"
        )
        draw_right(
            f"{discount_label} - {data['currency']} {data['discount_amount']:,.0f}",
            "Helvetica",
            8,
            HexColor("#b45309"),
            4.2 * mm,
        )
    if data["markup_amount"]:
        markup_label = (
            f"Ongezeko ({data['markup_input']:,.2f}%):"
            if data["markup_type"] == "PERCENT"
            else "Ongezeko:"
        )
        draw_right(
            f"{markup_label} + {data['currency']} {data['markup_amount']:,.0f}",
            "Helvetica",
            8,
            HexColor("#1d4ed8"),
            4.8 * mm,
        )

    c.setStrokeColor(light)
    c.line(margin, y, right_x, y)
    y -= 5 * mm
    draw_right(
        f"JUMLA KUU: {data['currency']} {data['receipt_total']:,.0f}",
        "Helvetica-Bold",
        11,
        purple,
        5.5 * mm,
    )

    # Thank-you message belongs immediately under the grand-total divider.
    c.setStrokeColor(light)
    c.line(margin, y, right_x, y)
    y -= 4.5 * mm
    draw_center(data["receipt_footer"], "Helvetica-Bold", 7.2, black, 4.2 * mm)

    # Centered QR code.
    qr_buf = BytesIO()
    qrcode.make(data["receipt_url"]).save(qr_buf, format="PNG")
    qr_buf.seek(0)
    qr_size = 30 * mm
    c.drawImage(
        ImageReader(qr_buf),
        center_x - (qr_size / 2),
        y - qr_size,
        width=qr_size,
        height=qr_size,
        preserveAspectRatio=True,
        mask="auto",
    )
    y -= qr_size + 4.5 * mm
    draw_center("Scan kuthibitisha risiti", "Helvetica", 6.8, gray, 4 * mm)

    draw_center("Powered by Meron Tech", "Helvetica-Bold", 7.2, purple, 3.8 * mm)
    draw_center("© 2026 Meron Tech. All rights reserved.", "Helvetica", 6.5, gray, 2 * mm)

    c.save()
    return response


# =========================================================
# 3. USIMAMIZI WA WATEJA
# =========================================================

@login_required(login_url="login")
def wateja_view(request):
    denied = require_roles(request, 'admin')
    if denied:
        return denied
    query = (
        request.GET.get("q")
        or ""
    ).strip()

    wateja = Mteja.objects.all()

    if query:
        wateja = wateja.filter(
            Q(
                majina_kamili__icontains=query
            )
            | Q(
                namba_ya_simu__icontains=query
            )
        )

    # CUSTOMER PROFILE HISTORY SURGERY:
    # Build the complete purchase/receipt history from the existing Mauzo
    # relationship. No new model or receipt system is introduced.
    customer_history = {}

    customer_rows = list(wateja)
    customer_ids = [
        customer.id
        for customer in customer_rows
    ]

    if customer_ids:
        history_sales = (
            Mauzo.objects
            .filter(mteja_id__in=customer_ids)
            .select_related(
                "mteja",
                "bidhaa",
                "muuzaji",
            )
            .order_by(
                "mteja_id",
                "-tarehe_ya_mauzo",
                "-id",
            )
        )

        for sale in history_sales:
            customer_id = str(sale.mteja_id)
            customer = sale.mteja
            invoice = sale.namba_ya_invoice or ""
            invoice_key = invoice or f"__sale_{sale.id}"

            bucket = customer_history.setdefault(
                customer_id,
                {
                    "customer": {
                        "id": customer.id,
                        "name": customer.majina_kamili or "",
                        "phone": customer.namba_ya_simu or "",
                        "whatsapp": customer.whatsapp_no or "",
                        "email": customer.email or "",
                        "type": customer.aina or "",
                        "location": customer.eneo or "",
                        "address": customer.anwani or "",
                        "registered": (
                            timezone.localtime(
                                customer.tarehe_ya_kusajiliwa
                            ).strftime("%d %b %Y")
                            if customer.tarehe_ya_kusajiliwa
                            else ""
                        ),
                        "marketing_opt_in": bool(
                            customer.marketing_opt_in
                        ),
                        "balance": float(
                            customer.outstanding_balance or 0
                        ),
                        "deposit": float(
                            customer.pre_order_deposit or 0
                        ),
                    },
                    "invoices": {},
                },
            )

            invoices = bucket["invoices"]

            if invoice_key not in invoices:
                local_dt = timezone.localtime(
                    sale.tarehe_ya_mauzo
                )

                invoices[invoice_key] = {
                    "invoice": invoice,
                    "sale_id": sale.id,
                    "date": local_dt.strftime("%d %b %Y"),
                    "time": local_dt.strftime("%I:%M %p"),
                    "staff": (
                        (
                            f"{sale.muuzaji.first_name} "
                            f"{sale.muuzaji.last_name}"
                        ).strip()
                        if sale.muuzaji
                        else ""
                    ),
                    "staff_username": (
                        sale.muuzaji.username
                        if sale.muuzaji
                        else ""
                    ),
                    "payment": sale.get_njia_ya_malipo_display(),
                    "total": 0.0,
                    "items": [],
                    "receipt_url": reverse(
                        "receipt",
                        args=[sale.id],
                    ),
                }

            invoice_row = invoices[invoice_key]
            line_total = float(
                sale.jumla_pesa_iliyopokelewa or 0
            )

            invoice_row["total"] += line_total
            invoice_row["items"].append(
                {
                    "product": (
                        sale.bidhaa.jina_la_bidhaa
                        if sale.bidhaa
                        else ""
                    ),
                    "quantity": int(sale.idadi or 0),
                    "unit_price": float(
                        sale.bei_ya_kuuzia_stoo or 0
                    ),
                    "line_total": line_total,
                }
            )

    # Convert invoice dictionaries to JSON-ready lists for the template.
    for customer_data in customer_history.values():
        customer_data["invoices"] = list(
            customer_data["invoices"].values()
        )
        customer_data["invoice_count"] = len(
            customer_data["invoices"]
        )
        customer_data["item_count"] = sum(
            item["quantity"]
            for invoice in customer_data["invoices"]
            for item in invoice["items"]
        )
        customer_data["total_spent"] = sum(
            invoice["total"]
            for invoice in customer_data["invoices"]
        )

    wateja = customer_rows

    now = timezone.now()

    start_today = now.replace(
        hour=0,
        minute=0,
        second=0,
        microsecond=0,
    )

    end_today = (
        start_today
        + timedelta(days=1)
    )

    total_customers_today = (
        Mteja.objects
        .filter(
            tarehe_ya_kusajiliwa__gte=start_today,
            tarehe_ya_kusajiliwa__lt=end_today,
        )
        .count()
    )

    new_customers = (
        Mteja.objects
        .filter(
            tarehe_ya_kusajiliwa__gte=start_today,
            tarehe_ya_kusajiliwa__lt=end_today,
        )
        .count()
    )

    vip_customers = (
        Mteja.objects
        .filter(
            outstanding_balance__gt=0,
        )
        .count()
    )

    outstanding_balance = (
        Mteja.objects
        .aggregate(
            total=Sum(
                "outstanding_balance"
            )
        )["total"]
        or Decimal("0.00")
    )

    context = {
        "title": "Usimamizi wa Wateja",
        "wateja": wateja,
        "query": query,
        "total_customers_today": total_customers_today,
        "new_customers": new_customers,
        "vip_customers": vip_customers,
        "outstanding_balance": outstanding_balance,
        "customer_history": customer_history,
    }

    return render(
        request,
        "wateja.html",
        context,
    )


@login_required(login_url="login")
def edit_mteja(request, customer_id):
    denied = require_roles(request, "admin")
    if denied:
        return denied
    if request.method != "POST":
        return redirect("wateja")
    mteja = get_object_or_404(Mteja, pk=customer_id)
    jina = (request.POST.get("majina_kamili") or "").strip()
    simu = (request.POST.get("namba_ya_simu") or "").strip()
    whatsapp = (request.POST.get("whatsapp_no") or "").strip()
    email = (request.POST.get("email") or "").strip()
    anwani = (request.POST.get("anwani") or request.POST.get("address") or "").strip()
    eneo = (request.POST.get("eneo") or request.POST.get("location") or request.POST.get("mahali") or "").strip()
    if not jina or not simu:
        messages.error(request, "Jina na namba ya simu ni lazima.")
        return redirect("wateja")
    mteja.majina_kamili = jina
    mteja.namba_ya_simu = simu
    mteja.whatsapp_no = whatsapp or None
    mteja.email = email or None
    mteja.anwani = anwani or None
    mteja.eneo = eneo or None
    mteja.save(
        update_fields=[
            "majina_kamili",
            "namba_ya_simu",
            "whatsapp_no",
            "email",
            "anwani",
            "eneo",
        ]
    )
    create_activity(request.user, f"Mteja amehaririwa: {mteja.majina_kamili}", "wateja")
    messages.success(request, "Taarifa za mteja zimesasishwa.")
    return redirect("wateja")


@login_required(login_url="login")
def import_customers(request):
    denied = require_roles(request, "admin")
    if denied:
        return denied

    is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"

    if request.method != "POST":
        payload = {
            "ok": False,
            "message": "Method hairuhusiwi.",
            "imported": 0,
            "duplicates": 0,
            "invalid": 0,
            "failed": 0,
            "total_rows": 0,
            "errors": ["Tuma POST request kwa import endpoint."],
        }
        if is_ajax:
            return JsonResponse(payload, status=405)
        return redirect("wateja")

    upload = request.FILES.get("customer_file")
    duka = get_store_profile(request)
    if duka is None:
        payload = {
            "ok": False,
            "message": "Biashara ya mtumiaji haijapatikana.",
            "imported": 0,
            "duplicates": 0,
            "invalid": 0,
            "failed": 1,
            "total_rows": 0,
            "errors": ["Account hii haijaunganishwa na biashara."],
        }
        if is_ajax:
            return JsonResponse(payload, status=400)
        messages.error(request, payload["message"])
        return redirect("wateja")

    if not upload:
        payload = {
            "ok": False,
            "message": "Chagua CSV file kwanza.",
            "imported": 0,
            "duplicates": 0,
            "invalid": 0,
            "failed": 0,
            "total_rows": 0,
            "errors": ["Hakuna customer_file iliyopokelewa."],
        }
        if is_ajax:
            return JsonResponse(payload, status=400)
        messages.error(request, payload["message"])
        return redirect("wateja")

    if not (upload.name or "").lower().endswith(".csv"):
        payload = {
            "ok": False,
            "message": "Faili linapaswa kuwa CSV.",
            "imported": 0,
            "duplicates": 0,
            "invalid": 0,
            "failed": 0,
            "total_rows": 0,
            "errors": ["Extension ya faili si .csv."],
        }
        if is_ajax:
            return JsonResponse(payload, status=400)
        messages.error(request, payload["message"])
        return redirect("wateja")

    try:
        content = upload.read().decode("utf-8-sig")

        try:
            dialect = csv.Sniffer().sniff(
                content[:4096],
                delimiters=",;|\t",
            )
        except csv.Error:
            dialect = csv.excel

        reader = csv.DictReader(
            io.StringIO(content),
            dialect=dialect,
        )

        rows = list(reader)

        total_rows = len(rows)
        imported = 0
        duplicates = 0
        invalid = 0
        failed = 0
        errors = []

        def pick(row, *aliases):
            normalized = {
                str(key).strip().lower(): (value or "").strip()
                for key, value in row.items()
                if key is not None
            }

            for alias in aliases:
                value = normalized.get(alias)
                if value:
                    return value

            return ""

        for row_number, row in enumerate(rows, start=2):

            name = pick(
                row,
                "majina_kamili",
                "jina",
                "name",
                "full_name",
            )

            phone = pick(
                row,
                "namba_ya_simu",
                "simu",
                "phone",
                "phone_number",
            )

            whatsapp = pick(
                row,
                "whatsapp_no",
                "whatsapp",
                "whatsapp_number",
            )

            email = pick(
                row,
                "email",
                "email_address",
            )

            eneo = pick(
                row,
                "eneo",
                "location",
                "mahali",
                "place",
            )

            anwani = pick(
                row,
                "anwani",
                "address",
                "residence",
            )

            if not name or not phone:
                invalid += 1

                if len(errors) < 6:
                    errors.append(
                        f"Row {row_number}: Jina na namba ya simu vinahitajika."
                    )

                continue

            if Mteja.objects.filter(
                namba_ya_simu=phone,
                majina_kamili__iexact=name,
            ).exists():

                duplicates += 1
                continue

            try:
                with transaction.atomic():
                    Mteja.objects.create(
                        duka=duka,
                        majina_kamili=name,
                        namba_ya_simu=phone,
                        whatsapp_no=whatsapp or None,
                        email=email or None,
                        eneo=eneo or None,
                        anwani=anwani or None,
                    )

                imported += 1

            except Exception as exc:

                failed += 1

                logger.exception(
                    "Customer import row failed: row=%s name=%s",
                    row_number,
                    name,
                )

                if len(errors) < 6:
                    errors.append(
                        f"Row {row_number}: {exc}"
                    )

        skipped = duplicates + invalid

        payload = {
            "ok": imported > 0 or failed == 0,
            "message": (
                f"Import imekamilika: contacts {imported} zimeingizwa."
                if imported > 0
                else "Hakuna contact mpya iliyoingizwa."
            ),
            "imported": imported,
            "duplicates": duplicates,
            "invalid": invalid,
            "failed": failed,
            "total_rows": total_rows,
            "skipped": skipped,
            "errors": errors,
        }

        if is_ajax:
            return JsonResponse(
                payload,
                status=200 if payload["ok"] else 400,
                json_dumps_params={"ensure_ascii": False},
            )

        if imported > 0 and (duplicates or invalid or failed):
            messages.warning(
                request,
                (
                    f"Import imekamilika kwa sehemu: "
                    f"{imported} imported, "
                    f"{duplicates} duplicates, "
                    f"{invalid} invalid, "
                    f"{failed} failed."
                ),
            )
        elif imported > 0:
            messages.success(
                request,
                f"Import imefanikiwa: contacts {imported} zimeongezwa.",
            )
        else:
            messages.error(
                request,
                "Import imekamilika lakini hakuna contact mpya iliyoongezwa.",
            )

    except UnicodeDecodeError:

        logger.exception(
            "Customer import failed: invalid encoding"
        )

        payload = {
            "ok": False,
            "message": "CSV haiwezi kusomeka kama UTF-8.",
            "imported": 0,
            "duplicates": 0,
            "invalid": 0,
            "failed": 1,
            "total_rows": 0,
            "errors": [
                "Save CSV yako kama UTF-8 CSV kisha ujaribu tena."
            ],
        }

        if is_ajax:
            return JsonResponse(
                payload,
                status=400,
                json_dumps_params={"ensure_ascii": False},
            )

        messages.error(
            request,
            payload["message"],
        )

    except Exception as exc:

        logger.exception(
            "Customer import failed"
        )

        payload = {
            "ok": False,
            "message": f"Import imeshindikana: {exc}",
            "imported": 0,
            "duplicates": 0,
            "invalid": 0,
            "failed": 1,
            "total_rows": 0,
            "errors": [str(exc)],
        }

        if is_ajax:
            return JsonResponse(
                payload,
                status=500,
                json_dumps_params={"ensure_ascii": False},
            )

        messages.error(
            request,
            f"Import imeshindikana: {exc}",
        )

    return redirect("wateja")


@login_required(login_url="login")
def sajili_mteja_haraka(request):
    denied = require_roles(request, 'admin', 'cashier')
    if denied:
        return denied
    if request.method != "POST":
        return redirect("mauzo")

    jina = (
        request.POST.get("majina_kamili")
        or ""
    ).strip()

    simu = (
        request.POST.get("namba_ya_simu")
        or ""
    ).strip()

    if not jina:
        messages.error(
            request,
            "Jina la mteja linahitajika.",
        )
        return redirect("mauzo")

    # CUSTOMER LOCATION SURGERY:
    # POS may submit the field as "location" while Wateja uses "eneo".
    # Accept both names and persist them into the existing Mteja fields.
    eneo = (
        request.POST.get("eneo")
        or request.POST.get("location")
        or request.POST.get("mahali")
        or ""
    ).strip()

    anwani = (
        request.POST.get("anwani")
        or request.POST.get("address")
        or ""
    ).strip()

    duka = get_store_profile(request)
    if duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("mauzo")

    mteja = Mteja.objects.create(
        duka=duka,
        majina_kamili=jina,
        namba_ya_simu=simu,
        eneo=eneo or None,
        anwani=anwani or None,
    )

    create_activity(
        request.user,
        (
            f"Mteja mpya amesajiliwa: "
            f"{mteja.majina_kamili}"
        ),
        "wateja",
    )

    messages.success(
        request,
        (
            f"Mteja {mteja.majina_kamili} "
            "amesajiliwa kikamilifu!"
        ),
    )

    return redirect("mauzo")


# =========================================================
# 4. DASHBOARD
# =========================================================

@login_required(login_url="login")
def dashboard(request):
    # Role-based landing: Cashier and Storekeeper use their dedicated workspaces
    # directly instead of entering the management dashboard.
    role = get_profile_role(request.user)

    if role == "cashier":
        return redirect("mauzo")

    if role == "stoo":
        return redirect("stoo_bidhaa")

    now = timezone.localtime(timezone.now())
    duka = get_store_profile(request)
    bidhaa_qs = Bidhaa.objects.filter(duka=duka) if duka else Bidhaa.objects.none()

    jumla_bidhaa = bidhaa_qs.count()

    inventory_store_map = {}
    inventory_shop_map = {}
    inventory_store_units = 0
    inventory_shop_units = 0
    if duka and getattr(duka, "inventory_location_mode", False):
        store_location, shop_location = ensure_inventory_locations(duka)
        if store_location:
            inventory_store_map = dict(
                ProductStock.objects.filter(
                    location=store_location,
                    bidhaa__duka=duka,
                ).values_list("bidhaa_id", "quantity")
            )
        if shop_location:
            inventory_shop_map = dict(
                ProductStock.objects.filter(
                    location=shop_location,
                    bidhaa__duka=duka,
                ).values_list("bidhaa_id", "quantity")
            )
        inventory_store_units = sum(int(v or 0) for v in inventory_store_map.values())
        inventory_shop_units = sum(int(v or 0) for v in inventory_shop_map.values())

        thamani_stoo = sum(
            (
                (p.bei_ya_kuuzia or Decimal("0.00"))
                * (
                    int(inventory_store_map.get(p.id, 0) or 0)
                    + int(inventory_shop_map.get(p.id, 0) or 0)
                )
            )
            for p in bidhaa_qs
        ) or Decimal("0.00")

        low_stock_count = 0
        out_of_stock_count = 0
        for p in bidhaa_qs:
            total_on_hand = (
                int(inventory_store_map.get(p.id, 0) or 0)
                + int(inventory_shop_map.get(p.id, 0) or 0)
            )
            if total_on_hand == 0:
                out_of_stock_count += 1
            elif total_on_hand <= 5:
                low_stock_count += 1
    else:
        thamani_stoo = (
            bidhaa_qs
            .aggregate(
                total=Sum(
                    ExpressionWrapper(
                        F("bei_ya_kuuzia") * F("idadi_stoo"),
                        output_field=DecimalField(max_digits=14, decimal_places=2),
                    )
                )
            )["total"]
            or Decimal("0.00")
        )
        low_stock_count = bidhaa_qs.filter(idadi_stoo__lte=5, idadi_stoo__gt=0).count()
        out_of_stock_count = bidhaa_qs.filter(idadi_stoo=0).count()

    goods_received_count = MzigoUlioingia.objects.filter(duka=duka).count() if duka else 0
    stock_issued_count = Mauzo.objects.filter(duka=duka).aggregate(total=Sum("idadi"))["total"] or 0

    pending_po_count = (
        PurchaseRequest.objects
        .filter(
            status="PENDING"
        )
        .count()
    )

    reorder_request_count = (
        pending_po_count
    )

    reorder_request = (
        PurchaseRequest.objects
        .filter(
            status="PENDING"
        )
        .select_related(
            "requested_by",
            "reviewed_by",
        )
        .prefetch_related(
            "items__bidhaa"
        )
        .first()
    )

    current_month = now.month
    current_year = now.year

    last_month = (
        current_month - 1
        if current_month > 1
        else 12
    )

    last_month_year = (
        current_year
        if current_month > 1
        else current_year - 1
    )

    mauzo_mwezi = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__year=current_year,
            tarehe_ya_mauzo__month=current_month,
        )
        .aggregate(
            total=Sum(
                "jumla_pesa_iliyopokelewa"
            )
        )["total"]
        or Decimal("0.00")
    )

    mauzo_mwezi_jana = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__year=last_month_year,
            tarehe_ya_mauzo__month=last_month,
        )
        .aggregate(
            total=Sum(
                "jumla_pesa_iliyopokelewa"
            )
        )["total"]
        or Decimal("0.00")
    )

    if mauzo_mwezi_jana > 0:
        growth_percentage = (
            (
                mauzo_mwezi
                - mauzo_mwezi_jana
            )
            / mauzo_mwezi_jana
        ) * 100
    else:
        growth_percentage = (
            100
            if mauzo_mwezi > 0
            else 0
        )

    mauzo_leo = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__date=now.date()
        )
        .aggregate(
            total=Sum(
                "jumla_pesa_iliyopokelewa"
            )
        )["total"]
        or Decimal("0.00")
    )

    # Target bado haijaunganishwa na model.
    target_ya_leo = Decimal("0.00")
    asilimia_target = 0
    bakiza_target = Decimal("0.00")

    if target_ya_leo > 0:
        asilimia_target = min(
            int(
                (
                    mauzo_leo
                    / target_ya_leo
                ) * 100
            ),
            100,
        )

        bakiza_target = max(
            target_ya_leo
            - mauzo_leo,
            Decimal("0.00"),
        )

    # MONTHLY FINANCIAL SNAPSHOT
    # Reuse the same sales/return/expense definitions used by the reporting layer:
    # Gross Sales = selling price × quantity
    # Net Sales = recorded sales − approved refunds
    # Net Profit = Net Sales − COGS − expenses
    # Cash Flow = recorded sales − approved refunds − paid expenses
    monthly_sales_qs = Mauzo.objects.filter(
        tarehe_ya_mauzo__year=current_year,
        tarehe_ya_mauzo__month=current_month,
    )
    monthly_returns_qs = SaleReturn.objects.filter(
        status="APPROVED",
        sale__tarehe_ya_mauzo__year=current_year,
        sale__tarehe_ya_mauzo__month=current_month,
    )

    monthly_sales_financials = monthly_sales_qs.annotate(
        actual_cogs=Coalesce(
            F("bei_ya_kununulia_stoo"),
            F("bidhaa__bei_ya_kununulia"),
            Value(Decimal("0.00")),
        )
    ).aggregate(
        gross_sales=Sum(
            ExpressionWrapper(
                F("bei_ya_kuuzia_stoo") * F("idadi"),
                output_field=DecimalField(max_digits=14, decimal_places=2),
            )
        ),
        recorded_sales=Sum("jumla_pesa_iliyopokelewa"),
        cogs=Sum(
            ExpressionWrapper(
                F("actual_cogs") * F("idadi"),
                output_field=DecimalField(max_digits=14, decimal_places=2),
            )
        ),
    )

    dashboard_gross_sales_mwezi = (
        monthly_sales_financials["gross_sales"]
        or Decimal("0.00")
    )
    dashboard_recorded_sales_mwezi = (
        monthly_sales_financials["recorded_sales"]
        or Decimal("0.00")
    )
    dashboard_cogs_mwezi = (
        monthly_sales_financials["cogs"]
        or Decimal("0.00")
    )

    dashboard_refund_mwezi = (
        monthly_returns_qs.aggregate(total=Sum("refund_amount"))["total"]
        or Decimal("0.00")
    )

    dashboard_net_sales_mwezi = (
        dashboard_recorded_sales_mwezi
        - dashboard_refund_mwezi
    )

    # Preserve the existing dashboard Gross Profit card semantics:
    # gross sales value minus historical COGS, before operating expenses.
    gross_profit_mwezi = (
        dashboard_gross_sales_mwezi
        - dashboard_cogs_mwezi
    )

    dashboard_expenses_qs = Matumizi.objects.filter(
        created_at__year=current_year,
        created_at__month=current_month,
    )
    dashboard_total_expenses_mwezi = (
        dashboard_expenses_qs.aggregate(total=Sum("kiasi"))["total"]
        or Decimal("0.00")
    )
    dashboard_paid_expenses_mwezi = (
        dashboard_expenses_qs
        .filter(status="PAID")
        .aggregate(total=Sum("kiasi"))["total"]
        or Decimal("0.00")
    )

    dashboard_net_profit_mwezi = (
        dashboard_net_sales_mwezi
        - dashboard_cogs_mwezi
        - dashboard_total_expenses_mwezi
    )
    dashboard_cash_flow_mwezi = (
        dashboard_recorded_sales_mwezi
        - dashboard_refund_mwezi
        - dashboard_paid_expenses_mwezi
    )

    if dashboard_gross_sales_mwezi > 0:
        profit_margin = (
            gross_profit_mwezi
            / dashboard_gross_sales_mwezi
        ) * 100
    else:
        profit_margin = Decimal("0.00")

    miamala_recent = (
        Mauzo.objects
        .select_related(
            "bidhaa",
            "mteja",
            "muuzaji",
        )
        .order_by("-id")[:10]
    )

    # Dashboard history datasets — reuse the existing models and relationships.
    # No new route, model, table or service is introduced.
    dashboard_sales_history = (
        Mauzo.objects
        .select_related(
            "bidhaa",
            "mteja",
            "muuzaji",
        )
        .order_by("-tarehe_ya_mauzo", "-id")[:200]
    )

    dashboard_expense_history = (
        Matumizi.objects
        .select_related("created_by")
        .order_by("-created_at")[:100]
    )

    dashboard_order_history = (
        PurchaseRequest.objects
        .select_related(
            "requested_by",
            "reviewed_by",
        )
        .prefetch_related("items__bidhaa")
        .order_by("-created_at")[:100]
    )

    activity_logs = (
        ActivityLog.objects
        .select_related("mhusika")
        .order_by("-muda")[:250]
    )

    activity_total = ActivityLog.objects.count()
    activity_summary = (
        ActivityLog.objects
        .values("aina")
        .annotate(total=Count("id"))
        .order_by("-total", "aina")
    )
    activity_type_count = activity_summary.count()

    top_selling = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__date=now.date()
        )
        .values(
            "bidhaa__jina_la_bidhaa"
        )
        .annotate(
            quantity=Sum("idadi")
        )
        .order_by("-quantity")
        .first()
    )

    top_selling_product = (
        top_selling[
            "bidhaa__jina_la_bidhaa"
        ]
        if top_selling
        else None
    )

    top_selling_product_quantity = (
        top_selling["quantity"]
        if top_selling
        else None
    )

    if duka and getattr(duka, "inventory_location_mode", False):
        healthy_count = sum(
            1
            for p in bidhaa_qs
            if (
                int(inventory_store_map.get(p.id, 0) or 0)
                + int(inventory_shop_map.get(p.id, 0) or 0)
            ) > 5
        )
    else:
        healthy_count = bidhaa_qs.filter(idadi_stoo__gt=5).count()

    critical_count = low_stock_count

    total_inventory_items = (
        jumla_bidhaa
    )

    if total_inventory_items:
        inventory_healthy_percentage = (
            healthy_count
            / total_inventory_items
        ) * 100

        inventory_critical_percentage = (
            critical_count
            / total_inventory_items
        ) * 100

        inventory_out_percentage = (
            out_of_stock_count
            / total_inventory_items
        ) * 100
    else:
        inventory_healthy_percentage = 0
        inventory_critical_percentage = 0
        inventory_out_percentage = 0

    inventory_health_percentage = (
        inventory_healthy_percentage
    )

    inventory_health_message = (
        f"{low_stock_count} bidhaa ziko kwenye low stock."
        if low_stock_count
        else "Inventory iko katika hali nzuri."
    )

    sales_rows = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__date=now.date()
        )
        .values("tarehe_ya_mauzo")
        .annotate(
            total=Sum(
                "jumla_pesa_iliyopokelewa"
            )
        )
        .order_by("tarehe_ya_mauzo")
    )

    sales_chart_data = {
        "labels": [
            row[
                "tarehe_ya_mauzo"
            ].strftime("%H:%M")
            for row in sales_rows
        ],
        "values": [
            float(
                row["total"] or 0
            )
            for row in sales_rows
        ],
    }

    # Notification has no direct Duka FK, so linked business data must be scoped
    # through the current tenant while system/user notifications remain usable.
    notification_tenant_filter = Q(purchase_request__isnull=True)
    if duka is not None:
        notification_tenant_filter |= Q(purchase_request__duka=duka)
    else:
        notification_tenant_filter = Q(pk__in=[])

    latest_notifications = (
        Notification.objects
        .filter(
            Q(recipient=request.user) & notification_tenant_filter
        )
        .select_related(
            "purchase_request"
        )
        .order_by(
            "-created_at"
        )[:10]
    )

    unread_notification_count = (
        Notification.objects
        .filter(
            Q(recipient=request.user)
            & Q(is_read=False)
            & notification_tenant_filter,
        )
        .count()
    )

    # REAL NOTIFICATION FEED: reuse the existing dashboard URL; no new route/model.
    notification_records = list(
        Notification.objects
        .filter(Q(recipient=request.user) & notification_tenant_filter)
        .select_related("purchase_request")
        .order_by("-created_at")[:100]
    )
    notified_purchase_request_ids = {
        n.purchase_request_id for n in notification_records if n.purchase_request_id
    }
    notification_alerts = []
    severity_by_type = {
        "PURCHASE_REQUEST": "order",
        "LOW_STOCK": "warning",
        "SALE": "success",
        "STOCK": "info",
        "SYSTEM": "info",
    }

    for notification in notification_records:
        notification_alerts.append({
            "kind": "notification",
            "severity": severity_by_type.get(notification.notification_type, "info"),
            "title": notification.title,
            "message": notification.message,
            "created_at": timezone.localtime(notification.created_at).isoformat(),
            "is_read": notification.is_read,
            "notification_id": notification.id,
            "read_url": reverse("notification_mark_read", args=[notification.id]),
            "action_url": notification.link_url or (reverse("dashboard") if notification.purchase_request_id else ""),
        })

    live_orders = (
        PurchaseRequest.objects
        .filter(status="PENDING")
        .select_related("requested_by")
        .prefetch_related("items__bidhaa")
        .order_by("-created_at")[:50]
    )
    for order in live_orders:
        if order.id in notified_purchase_request_ids:
            continue
        if order.requested_by and order.requested_by.get_full_name():
            requester = order.requested_by.get_full_name().strip()
        elif order.requested_by:
            requester = order.requested_by.username
        else:
            requester = "System"
        item_count = len(order.items.all())
        notification_alerts.append({
            "kind": "order",
            "severity": "order",
            "title": f"New Order Request · {order.display_number}",
            "message": f"{requester} ameleta order yenye {item_count} item{"s" if item_count != 1 else ""} inayosubiri kupitiwa.",
            "created_at": timezone.localtime(order.created_at).isoformat(),
            "is_read": False,
            "notification_id": None,
            "read_url": "",
            "action_url": reverse("dashboard"),
        })

    live_now_iso = timezone.localtime(now).isoformat()
    if duka and getattr(duka, "inventory_location_mode", False):
        inventory_alert_rows = []
        for bidhaa in bidhaa_qs:
            total_on_hand = (
                int(inventory_store_map.get(bidhaa.id, 0) or 0)
                + int(inventory_shop_map.get(bidhaa.id, 0) or 0)
            )
            if total_on_hand <= 5:
                inventory_alert_rows.append((total_on_hand, bidhaa))
        for stock_qty, bidhaa in sorted(
            inventory_alert_rows,
            key=lambda pair: (pair[0], pair[1].jina_la_bidhaa.lower()),
        ):
            is_out = stock_qty == 0
            notification_alerts.append({
                "kind": "stock",
                "severity": "danger" if is_out else "warning",
                "title": (
                    f"Stock Imeisha · {bidhaa.jina_la_bidhaa}"
                    if is_out
                    else f"Stock Inakaribia Kuisha · {bidhaa.jina_la_bidhaa}"
                ),
                "message": (
                    "Bidhaa hii imefika stock 0. Inahitaji replenishment."
                    if is_out
                    else f"Zimebaki {stock_qty} pcs stoo. Kiwango cha low stock ni 5 au chini."
                ),
                "time_label": "LIVE · OUT OF STOCK" if is_out else "LIVE · LOW STOCK",
                "created_at": live_now_iso,
                "is_read": False,
                "notification_id": None,
                "read_url": "",
                "action_url": reverse("stoo_bidhaa"),
            })
    else:
        for bidhaa in bidhaa_qs.filter(idadi_stoo=0).order_by("jina_la_bidhaa"):
            notification_alerts.append({
                "kind": "stock",
                "severity": "danger",
                "title": f"Stock Imeisha · {bidhaa.jina_la_bidhaa}",
                "message": "Bidhaa hii imefika stock 0. Inahitaji replenishment.",
                "time_label": "LIVE · OUT OF STOCK",
                "created_at": live_now_iso,
                "is_read": False,
                "notification_id": None,
                "read_url": "",
                "action_url": reverse("stoo_bidhaa"),
            })
        for bidhaa in bidhaa_qs.filter(idadi_stoo__gt=0, idadi_stoo__lte=5).order_by("idadi_stoo", "jina_la_bidhaa"):
            notification_alerts.append({
                "kind": "stock",
                "severity": "warning",
                "title": f"Stock Inakaribia Kuisha · {bidhaa.jina_la_bidhaa}",
                "message": f"Zimebaki {int(bidhaa.idadi_stoo or 0)} pcs stoo. Kiwango cha low stock ni 5 au chini.",
                "time_label": "LIVE · LOW STOCK",
                "created_at": live_now_iso,
                "is_read": False,
                "notification_id": None,
                "read_url": "",
                "action_url": reverse("stoo_bidhaa"),
            })

    notification_alerts.sort(key=lambda item: item.get("created_at") or "", reverse=True)
    live_alert_count = sum(1 for item in notification_alerts if item.get("kind") in {"order", "stock"})
    notification_alert_count = unread_notification_count + live_alert_count

    if request.GET.get("notification_feed") == "1":
        return JsonResponse({
            "ok": True,
            "alerts": notification_alerts,
            "alert_count": notification_alert_count,
            "unread_count": unread_notification_count,
        }, json_dumps_params={"ensure_ascii": False})

    context = {
        "title": "TradeCore Dashboard",
        "jumla_bidhaa": jumla_bidhaa,
        "thamani_stoo": thamani_stoo,
        "low_stock_count": low_stock_count,
        "out_of_stock_count": out_of_stock_count,
        "goods_received_count": goods_received_count,
        "pending_po_count": pending_po_count,
        "stock_issued_count": stock_issued_count,
        "mauzo_leo": mauzo_leo,
        "mauzo_mwezi": mauzo_mwezi,
        "growth_percentage": round(
            float(growth_percentage),
            1,
        ),
        "target_ya_leo": target_ya_leo,
        "asilimia_target": asilimia_target,
        "bakiza_target": bakiza_target,
        "gross_profit_mwezi": gross_profit_mwezi,
        "profit_margin": round(float(profit_margin), 1),
        "dashboard_gross_sales_mwezi": dashboard_gross_sales_mwezi,
        "dashboard_net_sales_mwezi": dashboard_net_sales_mwezi,
        "dashboard_total_expenses_mwezi": dashboard_total_expenses_mwezi,
        "dashboard_net_profit_mwezi": dashboard_net_profit_mwezi,
        "dashboard_cash_flow_mwezi": dashboard_cash_flow_mwezi,
        "dashboard_refund_mwezi": dashboard_refund_mwezi,
        "dashboard_paid_expenses_mwezi": dashboard_paid_expenses_mwezi,
        "top_selling_product": top_selling_product,
        "top_selling_product_quantity": (
            top_selling_product_quantity
        ),
        "miamala_recent": miamala_recent,
        "latest_sale_id": miamala_recent[0].id if miamala_recent else None,
        "dashboard_sales_history": dashboard_sales_history,
        "dashboard_expense_history": dashboard_expense_history,
        "dashboard_order_history": dashboard_order_history,
        "activity_logs": activity_logs,
        "activity_total": activity_total,
        "activity_summary": activity_summary,
        "activity_type_count": activity_type_count,
        "sales_chart_data": sales_chart_data,
        "reorder_request": reorder_request,
        "reorder_request_count": (
            reorder_request_count
        ),
        "latest_notifications": (
            latest_notifications
        ),
        "unread_notification_count": (
            unread_notification_count
        ),
        "notification_alert_count": notification_alert_count,
        "inventory_health_percentage": (
            inventory_health_percentage
        ),
        "inventory_healthy_percentage": (
            inventory_healthy_percentage
        ),
        "inventory_critical_percentage": (
            inventory_critical_percentage
        ),
        "inventory_out_percentage": (
            inventory_out_percentage
        ),
        "inventory_health_message": (
            inventory_health_message
        ),
    }

    return render(
        request,
        "dashboard.html",
        context,
    )


# =========================================================
# 5. USIMAMIZI WA MATUMIZI
# =========================================================

@login_required(login_url="login")
def matumizi_view(request):
    # SURGERY: Tunamruhusu Cashier na Storekeeper (Stoo) kuingia kwenye ukurasa wa Matumizi
    denied = require_roles(request, "admin", "cashier", "stoo")
    if denied:
        return denied

    duka = get_store_profile(request)
    if duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("dashboard")

    now = timezone.localtime(
        timezone.now()
    )

    start_today = now.replace(
        hour=0,
        minute=0,
        second=0,
        microsecond=0,
    )

    if start_today.month == 12:
        start_next_month = start_today.replace(
            year=start_today.year + 1,
            month=1,
            day=1,
        )
    else:
        start_next_month = start_today.replace(
            month=start_today.month + 1,
            day=1,
        )

    start_month = start_today.replace(
        day=1
    )

    if request.method == "POST":

        action = (
            request.POST.get("action")
            or ""
        ).strip().lower()

        role = get_profile_role(
            request.user
        )

        # SURGERY: Tunamruhusu Cashier kusajili matumizi ya duka
        can_manage = (
            request.user.is_superuser
            or request.user.is_staff
            or role in ["admin", "cashier","stoo"]
        )

        if not can_manage:
            messages.error(
                request,
                "Huna mamlaka ya kusajili au kulipa matumizi.",
            )
            return redirect("matumizi")

        if action == "create":

            category = (
                request.POST.get(
                    "category"
                )
                or "Others"
            ).strip()

            maelezo = (
                request.POST.get(
                    "maelezo"
                )
                or ""
            ).strip()

            reference = (
                request.POST.get(
                    "reference"
                )
                or ""
            ).strip()

            kiasi = parse_decimal(
                request.POST.get("kiasi")
            )

            status = (
                request.POST.get(
                    "status"
                )
                or "PAID"
            ).strip().upper()

            valid_categories = {
                value
                for value, _label
                in Matumizi.CATEGORY_CHOICES
            }

            valid_statuses = {
                "PAID",
                "PENDING",
            }

            if category not in valid_categories:
                messages.error(
                    request,
                    "Category ya matumizi si sahihi.",
                )
                return redirect("matumizi")

            if status not in valid_statuses:
                messages.error(
                    request,
                    "Status ya matumizi si sahihi.",
                )
                return redirect("matumizi")

            if not maelezo:
                messages.error(
                    request,
                    "Maelezo ya matumizi yanahitajika.",
                )
                return redirect("matumizi")

            if kiasi <= 0:
                messages.error(
                    request,
                    "Kiasi lazima kiwe zaidi ya sifuri.",
                )
                return redirect("matumizi")

            matumizi = Matumizi.objects.create(
                duka=duka,
                category=category,
                maelezo=maelezo,
                reference=reference,
                kiasi=kiasi,
                status=status,
                created_by=request.user,
                paid_at=(
                    now
                    if status == "PAID"
                    else None
                ),
            )

            create_activity(
                request.user,
                (
                    f"Matumizi mapya yamesajiliwa: "
                    f"{matumizi.maelezo} - "
                    f"TZS {matumizi.kiasi:,.2f}"
                ),
                "matumizi",
            )

            messages.success(
                request,
                "Matumizi yamesajiliwa kikamilifu.",
            )

            return redirect("matumizi")

        if action == "pay":

            matumizi_id = request.POST.get(
                "matumizi_id"
            )

            matumizi = get_object_or_404(
                Matumizi,
                pk=matumizi_id,
                duka=duka,
            )

            if matumizi.status == "PAID":
                messages.info(
                    request,
                    "Matumizi haya tayari yameshalipwa.",
                )
                return redirect("matumizi")

            matumizi.status = "PAID"
            matumizi.paid_at = now

            matumizi.save(
                update_fields=[
                    "status",
                    "paid_at",
                ]
            )

            create_activity(
                request.user,
                (
                    f"Matumizi yamelipwa: "
                    f"{matumizi.maelezo} - "
                    f"TZS {matumizi.kiasi:,.2f}"
                ),
                "matumizi",
            )

            messages.success(
                request,
                "Bili imelipwa na status imebadilishwa kuwa Paid.",
            )

            return redirect("matumizi")

        messages.error(
            request,
            "Action ya matumizi haijulikani.",
        )

        return redirect("matumizi")

    query = (
        request.GET.get("q")
        or ""
    ).strip()

    status_filter = (
        request.GET.get("status")
        or ""
    ).strip().upper()

    matumizi_qs = (
        Matumizi.objects
        .select_related("created_by")
        .all()
    )

    if query:
        matumizi_qs = matumizi_qs.filter(
            Q(category__icontains=query)
            | Q(
                maelezo__icontains=query
            )
            | Q(
                reference__icontains=query
            )
            | Q(
                created_by__username__icontains=query
            )
        )

    if status_filter in {
        "PAID",
        "PENDING",
    }:
        matumizi_qs = matumizi_qs.filter(
            status=status_filter
        )

    all_expenses = Matumizi.objects.all()

    today_qs = all_expenses.filter(
        created_at__gte=start_today,
        created_at__lt=(
            start_today
            + timedelta(days=1)
        ),
    )

    monthly_qs = all_expenses.filter(
        created_at__gte=start_month,
        created_at__lt=start_next_month,
    )

    pending_qs = all_expenses.filter(
        status="PENDING"
    )

    today_total = (
        today_qs.aggregate(
            total=Sum("kiasi")
        )["total"]
        or Decimal("0.00")
    )

    monthly_total = (
        monthly_qs.aggregate(
            total=Sum("kiasi")
        )["total"]
        or Decimal("0.00")
    )

    pending_total = (
        pending_qs.aggregate(
            total=Sum("kiasi")
        )["total"]
        or Decimal("0.00")
    )

    largest_expense = (
        monthly_qs
        .order_by(
            "-kiasi",
            "-created_at",
        )
        .first()
    )

    context = {
        "title": "Usimamizi wa Matumizi",
        "matumizi": matumizi_qs[:200],
        "query": query,
        "status_filter": status_filter,
        "today_total": today_total,
        "monthly_total": monthly_total,
        "pending_total": pending_total,
        "largest_expense": largest_expense,
        "expense_categories": (
            Matumizi.CATEGORY_CHOICES
        ),
        "now": now,
    }

    return render(
        request,
        "matumizi.html",
        context,
    )


# =========================================================
# 6. RIPOTI & UCHAMBUZI
# =========================================================


def _daily_report_data(duka, report_date):
    """Build one authoritative, product-level operational report for a business day."""
    sales_qs = (
        Mauzo.objects
        .filter(tarehe_ya_mauzo__date=report_date)
        .select_related("bidhaa", "muuzaji", "mteja")
    )
    if duka:
        sales_qs = sales_qs.filter(duka=duka)

    sales_rows = []
    sales_product = {}
    for sale in sales_qs.order_by("tarehe_ya_mauzo", "id"):
        key = sale.bidhaa_id
        row = sales_product.setdefault(
            key,
            {
                "name": sale.bidhaa.jina_la_bidhaa,
                "quantity": 0,
                "sales": Decimal("0"),
                "cogs": Decimal("0"),
                "profit": Decimal("0"),
            },
        )
        qty = int(sale.idadi or 0)
        amount = Decimal(sale.jumla_pesa_iliyopokelewa or 0)
        # SURGERY FIX: Tumia Historical Cost kwenye Daily Report
        historical_cost = sale.bei_ya_kununulia_stoo or sale.bidhaa.bei_ya_kununulia or Decimal("0.00")
        cogs = historical_cost * qty
        row["quantity"] += qty
        row["sales"] += amount
        row["cogs"] += cogs
        row["profit"] += amount - cogs
        local = timezone.localtime(sale.tarehe_ya_mauzo)
        sales_rows.append(
            {
                "time": local.strftime("%H:%M"),
                "invoice": sale.namba_ya_invoice or f"INV-{sale.id:06d}",
                "product": sale.bidhaa.jina_la_bidhaa,
                "quantity": qty,
                "amount": amount,
                "user": sale.muuzaji.username if sale.muuzaji else "Mfumo",
            }
        )
    product_rows = sorted(
        sales_product.values(),
        key=lambda x: (-x["quantity"], x["name"].lower()),
    )
    gross_sales = sum((x["sales"] for x in product_rows), Decimal("0"))
    cogs = sum((x["cogs"] for x in product_rows), Decimal("0"))
    profit = sum((x["profit"] for x in product_rows), Decimal("0"))

    expenses_qs = Matumizi.objects.filter(created_at__date=report_date)
    if duka:
        expenses_qs = expenses_qs.filter(duka=duka)
    expense_rows = [
        {
            "time": timezone.localtime(x.created_at).strftime("%H:%M"),
            "category": x.category,
            "description": x.maelezo,
            "amount": Decimal(x.kiasi or 0),
            "user": x.created_by.username if x.created_by else "Mfumo",
        }
        for x in expenses_qs.order_by("created_at")
    ]
    expenses_total = sum((x["amount"] for x in expense_rows), Decimal("0"))

    movement_qs = (
        StockMovement.objects
        .filter(created_at__date=report_date)
        .select_related("bidhaa", "actor", "location", "from_location", "to_location", "transfer")
    )
    if duka:
        movement_qs = movement_qs.filter(duka=duka)

    movements = []
    touched_product_ids = set(sales_product.keys())
    transfer_map = {}
    for m in movement_qs.order_by("created_at", "id"):
        touched_product_ids.add(m.bidhaa_id)
        local = timezone.localtime(m.created_at)
        if m.movement_type == "TRANSFER":
            detail = f"{m.from_location.jina if m.from_location else '-'} → {m.to_location.jina if m.to_location else '-'} · {m.quantity}"
            if m.reference not in transfer_map:
                transfer_map[m.reference] = {
                    "time": local.strftime("%H:%M"),
                    "reference": m.reference,
                    "from": m.from_location.jina if m.from_location else "-",
                    "to": m.to_location.jina if m.to_location else "-",
                    "reason": m.reason or "",
                    "user": m.actor.username if m.actor else "Mfumo",
                    "items": [],
                }
            transfer_map[m.reference]["items"].append(
                {"product": m.bidhaa.jina_la_bidhaa, "quantity": int(m.quantity or 0)}
            )
        else:
            detail = f"{m.location.jina if m.location else '-'} · {m.quantity}"
        movements.append(
            {
                "time": local.strftime("%H:%M"),
                "type": dict(StockMovement.TYPE_CHOICES).get(m.movement_type, m.movement_type),
                "product": m.bidhaa.jina_la_bidhaa,
                "detail": detail,
                "reason": m.reason,
                "reference": m.reference,
                "user": m.actor.username if m.actor else "Mfumo",
            }
        )

    transfer_rows = list(transfer_map.values())
    stock_in_rows = [m for m in movements if m["type"] == "Stock In"]
    adjustment_rows = [m for m in movements if m["type"] in {"Adjustment", "Damage", "Stock Take"}]
    stock_in_count = len(stock_in_rows)
    adjustment_count = len(adjustment_rows)

    closing_stock_rows = []
    if duka:
        store_location, shop_location = ensure_inventory_locations(duka)
        if store_location and shop_location and touched_product_ids:
            products = (
                Bidhaa.objects
                .filter(id__in=touched_product_ids)
                .order_by("jina_la_bidhaa")
            )
            store_map = {
                x.bidhaa_id: int(x.quantity or 0)
                for x in ProductStock.objects.filter(location=store_location, bidhaa_id__in=touched_product_ids)
            }
            shop_map = {
                x.bidhaa_id: int(x.quantity or 0)
                for x in ProductStock.objects.filter(location=shop_location, bidhaa_id__in=touched_product_ids)
            }
            for product in products:
                store_qty = store_map.get(product.id, 0)
                shop_qty = shop_map.get(product.id, 0)
                closing_stock_rows.append(
                    {
                        "name": product.jina_la_bidhaa,
                        "store": store_qty,
                        "shop": shop_qty,
                        "total": store_qty + shop_qty,
                    }
                )

    transfer_count = len(transfer_rows)
    total_items = sum(x["quantity"] for x in product_rows)
    return {
        "date": report_date,
        "product_rows": product_rows,
        "sales_rows": sales_rows,
        "gross_sales": gross_sales,
        "cogs": cogs,
        "profit": profit,
        "expense_rows": expense_rows,
        "expenses_total": expenses_total,
        "balance": gross_sales - expenses_total,
        "movements": movements,
        "transfer_rows": transfer_rows,
        "stock_in_rows": stock_in_rows,
        "adjustment_rows": adjustment_rows,
        "closing_stock_rows": closing_stock_rows,
        "transfer_count": transfer_count,
        "stock_in_count": stock_in_count,
        "adjustment_count": adjustment_count,
        "total_items": total_items,
        "sales_transactions": sales_qs.count(),
    }


def _daily_whatsapp_message(report, store_name, kind="full"):
    d = report["date"].strftime("%d %b %Y")
    lines = []
    if kind == "morning":
        lines += [
            "TRADECORE — HABARI ZA ASUBUHI",
            store_name or "TradeCore",
            f"Taarifa ya {d}",
            "",
            "MAUZO YA JANA",
            f"• Items: {report['total_items']}",
            f"• Mauzo: TZS {report['gross_sales']:,.0f}",
            f"• Faida: TZS {report['profit']:,.0f}",
            f"• Matumizi: TZS {report['expenses_total']:,.0f}",
            f"• Bakio ya Mauzo − Matumizi: TZS {report['balance']:,.0f}",
            "",
            "BIDHAA ZILIZOUZWA",
        ]
        if report["product_rows"]:
            for x in report["product_rows"]:
                lines.append(f"• {x['name']} × {x['quantity']} — TZS {x['sales']:,.0f} — Faida TZS {x['profit']:,.0f}")
        else:
            lines.append("• Hakuna bidhaa iliyouzwa.")
        lines += [
            "",
            f"Stock In: {report['stock_in_count']}",
            f"Transfers: {report['transfer_count']}",
            f"Adjustments/Checks: {report['adjustment_count']}",
            "",
            "Ripoti kamili: PDF",
        ]
    else:
        lines += [
            "TRADECORE — FULL DAILY REPORT",
            store_name or "TradeCore",
            f"Tarehe: {d}",
            "",
            "1. MAUZO — PRODUCT DETAIL",
        ]
        if report["product_rows"]:
            for x in report["product_rows"]:
                lines.append(f"• {x['name']} × {x['quantity']} | Mauzo TZS {x['sales']:,.0f} | COGS TZS {x['cogs']:,.0f} | Faida TZS {x['profit']:,.0f}")
        else:
            lines.append("• Hakuna bidhaa iliyouzwa siku hii.")
        lines += [
            f"Jumla Items: {report['total_items']}",
            f"Jumla Mauzo: TZS {report['gross_sales']:,.0f}",
            f"Jumla COGS: TZS {report['cogs']:,.0f}",
            f"Jumla Faida: TZS {report['profit']:,.0f}",
            "",
            "2. SALES TRANSACTIONS",
        ]
        if report["sales_rows"]:
            for x in report["sales_rows"]:
                lines.append(f"• {x['time']} | {x['invoice']} | {x['product']} × {x['quantity']} | TZS {x['amount']:,.0f} | {x['user']}")
        else:
            lines.append("• Hakuna invoice/line ya mauzo.")
        lines += ["", "3. STOCK IN"]
        if report["stock_in_rows"]:
            for m in report["stock_in_rows"]:
                lines.append(f"• {m['time']} | {m['product']} | {m['detail']} | {m['reference']} | {m['user']} | {m['reason'] or '—'}")
        else:
            lines.append("• Hakuna Stock In.")
        lines += ["", "4. TRANSFERS"]
        if report["transfer_rows"]:
            for t in report["transfer_rows"]:
                lines.append(f"• {t['time']} | {t['reference']} | {t['from']} → {t['to']} | {t['reason'] or '—'} | By {t['user']}")
                for item in t["items"]:
                    lines.append(f"   - {item['product']} × {item['quantity']}")
        else:
            lines.append("• Hakuna transfer.")
        lines += ["", "5. RETURNS / DAMAGE / ADJUSTMENTS"]
        detail_rows = [m for m in report["movements"] if m["type"] in {"Return", "Damage", "Adjustment", "Stock Take"}]
        if detail_rows:
            for m in detail_rows:
                lines.append(f"• {m['time']} | {m['type']} | {m['product']} | {m['detail']} | {m['reference']} | {m['user']} | {m['reason'] or '—'}")
        else:
            lines.append("• Hakuna return/damage/adjustment.")
        lines += ["", "6. MATUMIZI"]
        if report["expense_rows"]:
            for e in report["expense_rows"]:
                lines.append(f"• {e['time']} | {e['category']} | {e['description']} | TZS {e['amount']:,.0f} | {e['user']}")
        else:
            lines.append("• Hakuna matumizi yaliyorekodiwa.")
        lines += [
            f"Jumla Matumizi: TZS {report['expenses_total']:,.0f}",
            "",
            "7. CLOSING STOCK — PRODUCT ZILIZOGUSWA LEO",
        ]
        if report["closing_stock_rows"]:
            for x in report["closing_stock_rows"]:
                lines.append(f"• {x['name']} | Store {x['store']} | Duka {x['shop']} | Total {x['total']}")
        else:
            lines.append("• Hakuna stock position ya bidhaa zilizoguswa leo.")
        lines += [
            "",
            "8. MUHTASARI WA FEDHA",
            f"• Mauzo: TZS {report['gross_sales']:,.0f}",
            f"• COGS: TZS {report['cogs']:,.0f}",
            f"• Faida: TZS {report['profit']:,.0f}",
            f"• Matumizi: TZS {report['expenses_total']:,.0f}",
            f"• Bakio ya Mauzo − Matumizi: TZS {report['balance']:,.0f}",
        ]
    return "\n".join(lines)


def _build_daily_report_pdf(duka, report_date):
    """Render the authoritative daily report as a PDF byte string."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.units import mm

    report = _daily_report_data(duka, report_date)
    store_name = getattr(duka, "jina_la_duka", None) or "TradeCore"
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"TradeCore Daily Report {report_date}",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "TCTitle",
        parent=styles["Title"],
        fontSize=18,
        leading=21,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#5b21b6"),
    )
    sub = ParagraphStyle(
        "TCSub",
        parent=styles["Normal"],
        fontSize=8,
        textColor=colors.HexColor("#6b7280"),
        alignment=TA_CENTER,
    )
    h = ParagraphStyle(
        "TCH",
        parent=styles["Heading2"],
        fontSize=10,
        leading=12,
        textColor=colors.HexColor("#5b21b6"),
        spaceBefore=8,
        spaceAfter=5,
    )
    story = [
        Paragraph("TRADECORE", title),
        Paragraph("Full Daily Business Report", sub),
        Paragraph(f"{store_name} · {report_date.strftime('%d %b %Y')}", sub),
        Spacer(1, 7 * mm),
    ]

    summary = [
        ["Metric", "Value"],
        ["Sales", f"TZS {report['gross_sales']:,.0f}"],
        ["Items", str(report["total_items"])],
        ["COGS", f"TZS {report['cogs']:,.0f}"],
        ["Profit", f"TZS {report['profit']:,.0f}"],
        ["Expenses", f"TZS {report['expenses_total']:,.0f}"],
        ["Balance", f"TZS {report['balance']:,.0f}"],
    ]
    t = Table(summary, colWidths=[65 * mm, 65 * mm])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#ede9fe")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#4c1d95")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#e5e7eb")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [t, Spacer(1, 5 * mm), Paragraph("1. PRODUCTS SOLD", h)]

    sales = [["Product", "Qty", "Sales", "Profit"]] + [
        [x["name"], x["quantity"], f"TZS {x['sales']:,.0f}", f"TZS {x['profit']:,.0f}"]
        for x in report["product_rows"]
    ]
    if len(sales) == 1:
        sales.append(["Hakuna bidhaa iliyouzwa", "0", "TZS 0", "TZS 0"])
    t = Table(sales, colWidths=[72 * mm, 18 * mm, 40 * mm, 40 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f0f7")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(t)

    story += [Spacer(1, 5 * mm), Paragraph("2. STOCK MOVEMENTS", h)]
    mv = [["Time", "Type", "Product", "Detail", "User"]] + [
        [
            m["time"],
            m["type"],
            m["product"],
            m["detail"] + (f" | {m['reason']}" if m["reason"] else ""),
            m["user"],
        ]
        for m in report["movements"]
    ]
    if len(mv) == 1:
        mv.append(["—", "—", "—", "Hakuna stock movement", "—"])
    t = Table(mv, colWidths=[15 * mm, 28 * mm, 45 * mm, 65 * mm, 25 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f0f7")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 6.3),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(t)

    story += [Spacer(1, 5 * mm), Paragraph("3. TRANSFERS", h)]
    transfer_rows = [["Time", "Ref", "From", "To", "Product", "Qty", "Reason"]]
    for tr in report["transfer_rows"]:
        for item in tr["items"]:
            transfer_rows.append([
                tr["time"], tr["reference"], tr["from"], tr["to"],
                item["product"], str(item["quantity"]), tr["reason"] or "—",
            ])
    if len(transfer_rows) == 1:
        transfer_rows.append(["—", "—", "—", "—", "Hakuna transfer", "0", "—"])
    t = Table(
        transfer_rows,
        colWidths=[14 * mm, 23 * mm, 27 * mm, 27 * mm, 47 * mm, 14 * mm, 30 * mm],
        repeatRows=1,
    )
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f0f7")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 6.1),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(t)

    story += [Spacer(1, 5 * mm), Paragraph("4. CLOSING STOCK", h)]
    closing = [["Product", "Store", "Duka", "Total"]] + [
        [x["name"], str(x["store"]), str(x["shop"]), str(x["total"])]
        for x in report["closing_stock_rows"]
    ]
    if len(closing) == 1:
        closing.append(["Hakuna stock position", "0", "0", "0"])
    t = Table(closing, colWidths=[100 * mm, 24 * mm, 24 * mm, 24 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f0f7")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(t)

    story += [Spacer(1, 5 * mm), Paragraph("5. EXPENSES", h)]
    ex = [["Time", "Category", "Description", "Amount", "User"]] + [
        [e["time"], e["category"], e["description"], f"TZS {e['amount']:,.0f}", e["user"]]
        for e in report["expense_rows"]
    ]
    if len(ex) == 1:
        ex.append(["—", "—", "Hakuna matumizi yaliyorekodiwa", "TZS 0", "—"])
    t = Table(ex, colWidths=[15 * mm, 30 * mm, 70 * mm, 32 * mm, 28 * mm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f0f7")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#e5e7eb")),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 6.5),
        ("ALIGN", (3, 1), (3, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
    ]))
    story.append(t)

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


def _make_daily_report_token(duka, report_date):
    payload = {"duka_id": duka.pk, "date": report_date.isoformat()}
    return signing.TimestampSigner(salt="tradecore-daily-report").sign_object(payload)


@login_required(login_url="login")
def daily_report_view(request):
    denied = require_roles(request, "admin")
    if denied:
        return denied
    duka = get_store_profile(request)
    if duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("dashboard")
    date_raw = (request.GET.get("date") or "").strip()
    try:
        report_date = timezone.datetime.strptime(date_raw, "%Y-%m-%d").date() if date_raw else timezone.localdate()
    except ValueError:
        report_date = timezone.localdate()
    report = _daily_report_data(duka, report_date)
    store_name = getattr(duka, "jina_la_duka", None) or "TradeCore"
    return render(
        request,
        "daily_report.html",
        {
            "title": "Daily Business Report",
            "store": store_name,
            "report": report,
            "whatsapp_message": _daily_whatsapp_message(report, store_name, "full"),
            "morning_message": _daily_whatsapp_message(report, store_name, "morning"),
        },
    )


@login_required(login_url="login")
def daily_report_pdf(request):
    denied = require_roles(request, "admin")
    if denied:
        return denied
    duka = get_store_profile(request)
    if duka is None:
        return HttpResponse("Biashara ya mtumiaji haijapatikana.", status=400)
    raw = (request.GET.get("date") or "").strip()
    try:
        report_date = timezone.datetime.strptime(raw, "%Y-%m-%d").date() if raw else timezone.localdate()
    except ValueError:
        report_date = timezone.localdate()
    pdf_bytes = _build_daily_report_pdf(duka, report_date)
    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = f'inline; filename="TradeCore_Daily_Report_{report_date.isoformat()}.pdf"'
    return response


@require_GET
def daily_report_pdf_public(request, token):
    signer = signing.TimestampSigner(salt="tradecore-daily-report")
    try:
        payload = signer.unsign_object(token, max_age=settings.TRADECORE_REPORT_TOKEN_MAX_AGE)
    except SignatureExpired:
        return HttpResponse("Kiungo cha ripoti kimekwisha muda. Omba report mpya.", status=410)
    except BadSignature:
        return HttpResponse("Kiungo cha ripoti si sahihi.", status=404)

    if not isinstance(payload, dict) or "duka_id" not in payload or "date" not in payload:
        return HttpResponse("Kiungo cha ripoti si sahihi.", status=404)

    try:
        report_date = timezone.datetime.strptime(str(payload["date"]), "%Y-%m-%d").date()
        duka = Duka.objects.select_related("mwenye_duka").get(pk=int(payload["duka_id"]))
    except (ValueError, TypeError, Duka.DoesNotExist):
        return HttpResponse("Ripoti haijapatikana.", status=404)

    tenant_token = set_current_duka_id(duka.pk)
    try:
        pdf_bytes = _build_daily_report_pdf(duka, report_date)
    finally:
        reset_current_duka(tenant_token)

    response = HttpResponse(pdf_bytes, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="TradeCore_Daily_Report_{report_date.isoformat()}.pdf"'
    response["Cache-Control"] = "private, no-store, max-age=0"
    return response


@login_required(login_url="login")
def ripoti_view(request):
    denied = require_roles(request, "admin")
    if denied:
        return denied

    now = timezone.localtime(
        timezone.now()
    )

    role = get_profile_role(
        request.user
    )

    # Ripoti za kifedha ni za management.
    if role != "admin" and not request.user.is_superuser:
        messages.error(
            request,
            "Huna ruhusa ya kuona ripoti za kifedha.",
        )
        return redirect("dashboard")

    period = (
        request.GET.get("period")
        or "month"
    ).lower()

    if period not in {
        "week",
        "month",
        "year",
    }:
        period = "month"

    today = now.date()

    # -----------------------------------------------------
    # PERIOD RANGE
    # -----------------------------------------------------

    if period == "week":
        start_date = (
            today
            - timedelta(days=6)
        )
        end_date = today

    elif period == "year":
        start_date = today.replace(
            month=1,
            day=1,
        )
        end_date = today.replace(
            month=12,
            day=31,
        )

    else:
        start_date = today.replace(
            day=1
        )

        if start_date.month == 12:
            next_month = start_date.replace(
                year=start_date.year + 1,
                month=1,
                day=1,
            )
        else:
            next_month = start_date.replace(
                month=start_date.month + 1,
                day=1,
            )

        end_date = (
            next_month
            - timedelta(days=1)
        )

    # -----------------------------------------------------
    # SALES
    # -----------------------------------------------------

    sales_qs = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__date__gte=start_date,
            tarehe_ya_mauzo__date__lte=end_date,
        )
        .select_related(
            "bidhaa",
            "mteja",
            "muuzaji",
        )
    )

    # -----------------------------------------------------
    # APPROVED RETURNS / REFUNDS
    # -----------------------------------------------------

    approved_returns = (
        SaleReturn.objects
        .filter(
            status="APPROVED",
            sale__tarehe_ya_mauzo__date__gte=start_date,
            sale__tarehe_ya_mauzo__date__lte=end_date,
        )
        .select_related(
            "sale",
            "sale__bidhaa",
        )
    )

    refund_total = (
        approved_returns.aggregate(
            total=Sum(
                "refund_amount"
            )
        )["total"]
        or Decimal("0.00")
    )

    # -----------------------------------------------------
    # GROSS SALES
    # Formula:
    # SUM(Selling Price × Quantity)
    # -----------------------------------------------------

    gross_sales = (
        sales_qs.aggregate(
            total=Sum(
                ExpressionWrapper(
                    F("bei_ya_kuuzia_stoo")
                    * F("idadi"),
                    output_field=DecimalField(
                        max_digits=14,
                        decimal_places=2,
                    ),
                )
            )
        )["total"]
        or Decimal("0.00")
    )

    # -----------------------------------------------------
    # RECORDED SALES
    # -----------------------------------------------------

    recorded_sales = (
        sales_qs.aggregate(
            total=Sum(
                "jumla_pesa_iliyopokelewa"
            )
        )["total"]
        or Decimal("0.00")
    )

    # -----------------------------------------------------
    # NET SALES
    # Formula:
    # Recorded Sales - Approved Refunds
    # -----------------------------------------------------

    net_sales = (
        recorded_sales
        - refund_total
    )

    # -----------------------------------------------------
    # COGS BEFORE RETURNS
    # Formula:
    # Purchase Price × Sold Quantity
    # -----------------------------------------------------

    cogs_before_returns = (
        sales_qs.annotate(
            # SURGERY FIX: Tumia Coalesce kuhakikisha inapata Historical Cost kwanza, akikosa anavuta Live Price (Kwa data za zamani)
            actual_cogs=Coalesce(F("bei_ya_kununulia_stoo"), F("bidhaa__bei_ya_kununulia"))
        ).aggregate(
            total=Sum(
                ExpressionWrapper(
                    F("actual_cogs") * F("idadi"),
                    output_field=DecimalField(
                        max_digits=14,
                        decimal_places=2,
                    ),
                )
            )
        )["total"]
        or Decimal("0.00")
    )

    # -----------------------------------------------------
    # COGS OF RETURNED GOODS
    # -----------------------------------------------------

    return_cogs = (
        approved_returns.annotate(
            # SURGERY FIX: Hapa napo tunatumia Historical cost kwa bidhaa zilizorudishwa
            actual_return_cogs=Coalesce(F("sale__bei_ya_kununulia_stoo"), F("sale__bidhaa__bei_ya_kununulia"))
        ).aggregate(
            total=Sum(
                ExpressionWrapper(
                    F("actual_return_cogs") * F("quantity_returned"),
                    output_field=DecimalField(
                        max_digits=14,
                        decimal_places=2,
                    ),
                )
            )
        )["total"]
        or Decimal("0.00")
    )

    # Net COGS after returned goods
    cogs = (
        cogs_before_returns
        - return_cogs
    )

    # -----------------------------------------------------
    # EXPENSES
    # -----------------------------------------------------

    expenses_qs = (
        Matumizi.objects
        .filter(
            created_at__date__gte=start_date,
            created_at__date__lte=end_date,
        )
    )

    total_expenses = (
        expenses_qs.aggregate(
            total=Sum("kiasi")
        )["total"]
        or Decimal("0.00")
    )

    # -----------------------------------------------------
    # NET PROFIT
    # Formula:
    # Net Sales - Net COGS - Expenses
    # -----------------------------------------------------

    net_profit = (
        net_sales
        - cogs
        - total_expenses
    )

    # -----------------------------------------------------
    # SALES CHART
    # -----------------------------------------------------

    chart_labels = []
    chart_values = []

    if period == "year":

        rows = (
            sales_qs
            .values(
                "tarehe_ya_mauzo__month"
            )
            .annotate(
                total=Sum(
                    "jumla_pesa_iliyopokelewa"
                )
            )
            .order_by(
                "tarehe_ya_mauzo__month"
            )
        )

        monthly = {
            row[
                "tarehe_ya_mauzo__month"
            ]: float(
                row["total"] or 0
            )
            for row in rows
        }

        for month in range(1, 13):

            label = timezone.datetime(
                today.year,
                month,
                1,
            ).strftime("%b")

            chart_labels.append(
                label
            )

            chart_values.append(
                monthly.get(
                    month,
                    0,
                )
            )

    else:

        rows = (
            sales_qs
            .values(
                "tarehe_ya_mauzo__date"
            )
            .annotate(
                total=Sum(
                    "jumla_pesa_iliyopokelewa"
                )
            )
            .order_by(
                "tarehe_ya_mauzo__date"
            )
        )

        daily = {
            row[
                "tarehe_ya_mauzo__date"
            ]: float(
                row["total"] or 0
            )
            for row in rows
        }

        current = start_date

        while current <= end_date:

            chart_labels.append(
                current.strftime(
                    "%d %b"
                )
            )

            chart_values.append(
                daily.get(
                    current,
                    0,
                )
            )

            current += timedelta(
                days=1
            )

    sales_chart_data = {
        "labels": chart_labels,
        "values": chart_values,
    }

    # -----------------------------------------------------
    # CASH FLOW
    # -----------------------------------------------------

    payment_rows = (
        sales_qs
        .values(
            "njia_ya_malipo"
        )
        .annotate(
            total=Sum(
                "jumla_pesa_iliyopokelewa"
            )
        )
        .order_by("-total")
    )

    cashflow_data = {
        "labels": [
            row[
                "njia_ya_malipo"
            ]
            for row in payment_rows
        ],
        "values": [
            float(
                row["total"] or 0
            )
            for row in payment_rows
        ],
    }

    # -----------------------------------------------------
    # TRENDING PRODUCTS
    # -----------------------------------------------------

    week_start = (
        today
        - timedelta(days=6)
    )

    week_trending = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__date__gte=week_start,
            tarehe_ya_mauzo__date__lte=today,
        )
        .values(
            "bidhaa__jina_la_bidhaa"
        )
        .annotate(
            quantity=Sum(
                "idadi"
            )
        )
        .order_by(
            "-quantity",
            "bidhaa__jina_la_bidhaa",
        )[:3]
    )

    month_trending = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__year=today.year,
            tarehe_ya_mauzo__month=today.month,
        )
        .values(
            "bidhaa__jina_la_bidhaa"
        )
        .annotate(
            quantity=Sum(
                "idadi"
            )
        )
        .order_by(
            "-quantity",
            "bidhaa__jina_la_bidhaa",
        )[:3]
    )

    def _serialize_trending(rows):
        return [
            {
                "name": row["bidhaa__jina_la_bidhaa"],
                "quantity": int(row["quantity"] or 0),
            }
            for row in rows
        ]

    trending_data = {
        "week": _serialize_trending(week_trending),
        "month": _serialize_trending(month_trending),
    }

    # -----------------------------------------------------
    # CUSTOMER DAILY REPORT
    # -----------------------------------------------------

    customer_daily_report = []

    today_sales = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__date=today
        )
        .select_related(
            "mteja",
            "bidhaa",
        )
        .order_by(
            "-tarehe_ya_mauzo"
        )
    )

    for sale in today_sales:

        local_time = timezone.localtime(
            sale.tarehe_ya_mauzo
        )

        customer_daily_report.append({
            "customer": (
                sale.mteja.majina_kamili
                if sale.mteja
                else "Walk-in Customer"
            ),
            "day": local_time.strftime("%A"), # Inaleta jina la siku (Mf. Monday)
            "date": local_time.strftime("%d %b %Y"), # Inaleta tarehe (Mf. 26 Sep 2026)
            "time": local_time.strftime("%I:%M %p"), # Hii inaleta ule muda vizuri
            "product": (
                sale.bidhaa.jina_la_bidhaa
            ),
            "quantity": sale.idadi,
            "amount": float(
                sale.jumla_pesa_iliyopokelewa
                or 0
            ),
        })

    # -----------------------------------------------------
    # EMPLOYEE DAILY REPORT
    # -----------------------------------------------------

    employee_rows = (
        Mauzo.objects
        .filter(
            tarehe_ya_mauzo__date=today
        )
        .values(
            "muuzaji__username",
            "muuzaji__first_name",
            "muuzaji__last_name",
        )
        .annotate(
            transactions=Count("id"),
            total_sales=Sum(
                "jumla_pesa_iliyopokelewa"
            ),
        )
        .order_by(
            "-total_sales"
        )
    )

    employee_daily_report = []

    for row in employee_rows:

        full_name = (
            f'{row["muuzaji__first_name"] or ""} '
            f'{row["muuzaji__last_name"] or ""}'
        ).strip()

        employee_daily_report.append({
            "name": (
                full_name
                or row[
                    "muuzaji__username"
                ]
                or "Unknown"
            ),
            "transactions": int(
                row["transactions"] or 0
            ),
            "total_sales": float(
                row["total_sales"] or 0
            ),
        })

    # -----------------------------------------------------
    # CONTEXT
    # -----------------------------------------------------

    context = {
        "title": "Ripoti & Uchambuzi",
        "report_period": period,

        "gross_sales": gross_sales,
        "recorded_sales": recorded_sales,
        "refund_total": refund_total,
        "net_sales": net_sales,

        "cogs_before_returns": (
            cogs_before_returns
        ),
        "return_cogs": return_cogs,
        "cogs": cogs,

        "total_expenses": (
            total_expenses
        ),

        "net_profit": net_profit,

        "sales_chart_data": (
            sales_chart_data
        ),

        "cashflow_data": (
            cashflow_data
        ),

        "trending_data": (
            trending_data
        ),

        "customer_daily_report": (
            customer_daily_report
        ),

        "employee_daily_report": (
            employee_daily_report
        ),
    }

    return render(
        request,
        "ripoti.html",
        context,
    )


# =========================================================
# 7. MAWASILIANO
# =========================================================
# =========================================================
# 7. MAWASILIANO
# =========================================================

@login_required(login_url="login")
def ujumbe_view(request):
    denied = require_roles(request, "admin")
    if denied:
        return denied

    # WATEJA WOTE KWA AJILI YA COMPOSER
    wateja_list = (
        Mteja.objects
        .all()
        .order_by("majina_kamili")
    )

    # SMS BALANCE
    sms_balance = 0

    try:
        vocha = VochaYaDuka.objects.first()

        if vocha:
            sms_balance = vocha.salio_la_sms

    except Exception:
        sms_balance = 0


    # COMMUNICATION HISTORY
    communication_history = []

    try:
        communication_history = (
            CommunicationLog.objects
            .select_related(
                "customer",
                "created_by",
            )
            .order_by("-created_at")[:20]
        )

    except Exception:
        # Ikiwa CommunicationLog bado haija-migrate,
        # page isi-crash.
        communication_history = []


    # DAILY REPORT STATUS / REAL COMMUNICATION & SALES STATS
    duka = get_store_profile(request)
    daily_report_settings, _ = DailyReportSettings.objects.get_or_create(duka=duka)
    daily_report_active = daily_report_settings.enabled
    today = timezone.localtime(timezone.now()).date()
    sales_today = Mauzo.objects.filter(tarehe_ya_mauzo__date=today)
    transaction_count_today = sales_today.values("namba_ya_invoice").distinct().count()
    items_sold_today = sales_today.aggregate(total=Sum("idadi"))["total"] or 0
    customers_today = sales_today.filter(mteja__isnull=False).values("mteja_id").distinct().count()
    payment_breakdown = {
        code: sales_today.filter(njia_ya_malipo=code).aggregate(total=Sum("jumla_pesa_iliyopokelewa"))["total"] or Decimal("0.00")
        for code, _label in Mauzo.PAYMENT_CHOICES
    }
    sales_by_staff = (
        sales_today.filter(muuzaji__isnull=False)
        .values("muuzaji__username")
        .annotate(total=Sum("jumla_pesa_iliyopokelewa"))
        .order_by("-total")[:5]
    )
    latest_transaction = sales_today.select_related("mteja", "bidhaa").order_by("-id").first()

    context = {
        "title": "Kituo cha Mawasiliano",
        "wateja_list": wateja_list,
        "sms_balance": sms_balance,
        "communication_history": communication_history,
        "daily_report_active": daily_report_active,
        "transaction_count_today": transaction_count_today,
        "items_sold_today": items_sold_today,
        "customers_today": customers_today,
        "payment_breakdown": payment_breakdown,
        "sales_by_staff": sales_by_staff,
        "latest_transaction": latest_transaction,
    }


    return render(
        request,
        "ujumbe.html",
        context,
    )



# =========================================================
# 8. PURCHASE REQUESTS + NOTIFICATIONS
# =========================================================

@login_required(login_url="login")
def create_smart_reorder_request(request):
    if request.method != "POST":
        return redirect("dashboard")
        
    # SURGERY FIX: Kutafuta duka la user kabla ya kulitumia
    duka = get_store_profile(request)
    if not duka:
        messages.error(request, "Biashara haijatambulika.")
        return redirect("dashboard")

    role = get_profile_role(
        request.user
    )

    if (
        role not in {"stoo", "admin"}
        and not request.user.is_superuser
    ):
        messages.error(
            request,
            "Huna mamlaka ya kutengeneza smart reorder.",
        )
        return redirect("dashboard")

    if getattr(duka, "inventory_location_mode", False):
        store_location, _ = ensure_inventory_locations(duka)
        low_stock_rows = (
            ProductStock.objects
            .filter(
                location=store_location,
                bidhaa__duka=duka,
                quantity__gt=0,
                quantity__lte=5,
            )
            .select_related("bidhaa")
            .order_by("quantity", "bidhaa__jina_la_bidhaa")
        )
        low_stock_products = [row.bidhaa for row in low_stock_rows]
        current_stock_map = {row.bidhaa_id: int(row.quantity or 0) for row in low_stock_rows}
    else:
        low_stock_products = list(
            Bidhaa.objects
            .filter(
                duka=duka,
                idadi_stoo__lte=5,
                idadi_stoo__gt=0,
            )
            .order_by("idadi_stoo", "jina_la_bidhaa")
        )
        current_stock_map = {
            product.id: int(product.idadi_stoo or 0)
            for product in low_stock_products
        }

    if not low_stock_products:
        messages.info(
            request,
            (
                "Hakuna bidhaa yenye low stock "
                "inayohitaji reorder kwa sasa."
            ),
        )
        return redirect("dashboard")

    with transaction.atomic():

        existing = (
            PurchaseRequest.objects
            .filter(
                status="PENDING",
                requested_by=request.user,
            )
            .first()
        )

        if existing:
            messages.info(
                request,
                (
                    f"{existing.display_number} "
                    "bado ipo pending."
                ),
            )
            return redirect("dashboard")

        purchase_request = (
            PurchaseRequest.objects.create(
                duka=duka,
                requested_by=request.user,
                source_department="Stoo",
                notes=(
                    "Smart reorder generated "
                    "from current low-stock levels."
                ),
            )
        )

        for product in low_stock_products:

            recommended_quantity = max(1, 10 - current_stock_map.get(product.id, 0))

            PurchaseRequestItem.objects.create(
                purchase_request=purchase_request,
                bidhaa=product,
                quantity=recommended_quantity,
                unit_price=product.bei_ya_kununulia,
                current_stock=current_stock_map.get(product.id, 0),
            )

        reviewers = (
            User.objects
            .filter(profile__duka=duka, is_active=True)
            .filter(
                Q(is_superuser=True)
                | Q(is_staff=True)
            )
            .distinct()
        )

        notifications = [
            Notification(
                recipient=user,
                title="Smart Reorder Request Mpya",
                message=(
                    f"{request.user.username} "
                    f"ametengeneza "
                    f"{purchase_request.display_number} "
                    "kutoka low-stock data."
                ),
                notification_type="PURCHASE_REQUEST",
                purchase_request=purchase_request,
            )
            for user in reviewers
            if user.pk != request.user.pk
        ]

        if notifications:
            Notification.objects.bulk_create(
                notifications
            )

        create_activity(
            request.user,
            (
                f"Smart reorder "
                f"{purchase_request.display_number} "
                "imetengenezwa"
            ),
            "purchase_request",
        )

    messages.success(
        request,
        (
            f"{purchase_request.display_number} "
            "imetengenezwa kutoka low-stock data."
        ),
    )

    return redirect("dashboard")


@login_required(login_url="login")
def create_purchase_request(request):
    if request.method != "POST":
        return redirect("dashboard")
        
    # SURGERY FIX: Kutafuta duka la user kabla ya kulitumia
    duka = get_store_profile(request)
    if not duka:
        messages.error(request, "Biashara haijatambulika.")
        return redirect("dashboard")

    role = get_profile_role(
        request.user
    )

    if (
        role not in {"stoo", "admin"}
        and not request.user.is_superuser
    ):
        messages.error(
            request,
            "Huna mamlaka ya kutuma purchase request.",
        )
        return redirect("dashboard")

    product_ids = request.POST.getlist(
        "bidhaa_id"
    )

    quantities = request.POST.getlist(
        "quantity"
    )

    notes = (
        request.POST.get("notes")
        or ""
    ).strip()

    if not product_ids:
        messages.error(
            request,
            "Chagua angalau bidhaa moja.",
        )
        return redirect("dashboard")

    try:
        store_location = None
        current_stock_map = {}
        if getattr(duka, "inventory_location_mode", False):
            store_location, _ = ensure_inventory_locations(duka)
            if store_location:
                current_stock_map = dict(
                    ProductStock.objects.filter(
                        location=store_location,
                        bidhaa_id__in=product_ids,
                    ).values_list("bidhaa_id", "quantity")
                )

        with transaction.atomic():

            purchase_request = (
                PurchaseRequest.objects.create(
                    duka=duka,
                    requested_by=request.user,
                    source_department="Stoo",
                    notes=notes or None,
                )
            )

            for product_id, quantity_value in zip(
                product_ids,
                quantities,
            ):

                quantity = parse_positive_int(
                    quantity_value
                )

                if quantity <= 0:
                    continue

                product = get_object_or_404(
                    Bidhaa.objects.select_for_update(),
                    pk=product_id,
                    duka=duka,
                )

                PurchaseRequestItem.objects.create(
                    purchase_request=purchase_request,
                    bidhaa=product,
                    quantity=quantity,
                    unit_price=product.bei_ya_kununulia,
                    current_stock=(current_stock_map.get(product.id, 0) if store_location else int(product.idadi_stoo or 0)),
                )

            if not purchase_request.items.exists():
                raise ValueError(
                    "Hakuna item halali kwenye purchase request."
                )

            reviewers = (
                User.objects
                .filter(profile__duka=duka, is_active=True)
                .filter(
                    Q(is_superuser=True)
                    | Q(is_staff=True)
                )
                .distinct()
            )

            notifications = [
                Notification(
                    recipient=user,
                    title="Purchase Request Mpya",
                    message=(
                        f"{request.user.username} "
                        f"ametuma "
                        f"{purchase_request.display_number} "
                        "kwa ajili ya mapitio."
                    ),
                    notification_type="PURCHASE_REQUEST",
                    purchase_request=purchase_request,
                )
                for user in reviewers
                if user.pk != request.user.pk
            ]

            if notifications:
                Notification.objects.bulk_create(
                    notifications
                )

            create_activity(
                request.user,
                (
                    f"Purchase request "
                    f"{purchase_request.display_number} "
                    "imetumwa"
                ),
                "purchase_request",
            )

    except Exception:
        logger.exception(
            "Purchase request creation failed for user=%s",
            request.user.username,
        )

        messages.error(
            request,
            "Purchase request haijatumwa. Tafadhali hakikisha taarifa ni sahihi.",
        )

        return redirect("dashboard")

    messages.success(
        request,
        (
            f"{purchase_request.display_number} "
            "imetumwa kwa mapitio."
        ),
    )

    return redirect("dashboard")


@login_required(login_url="login")
def approve_purchase_request(
    request,
    request_id,
):
    if request.method != "POST":
        return redirect("dashboard")

    role = get_profile_role(
        request.user
    )

    if (
        role != "admin"
        and not request.user.is_staff
        and not request.user.is_superuser
    ):
        messages.error(
            request,
            (
                "Huna mamlaka ya "
                "ku-approve purchase request."
            ),
        )
        return redirect("dashboard")

    duka = get_store_profile(request)
    if duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("dashboard")

    purchase_request = get_object_or_404(
        PurchaseRequest.objects.prefetch_related(
            "items__bidhaa"
        ),
        pk=request_id,
        duka=duka,
    )

    if purchase_request.status != "PENDING":
        messages.error(
            request,
            "Purchase request hii tayari imechakatwa.",
        )
        return redirect("dashboard")

    purchase_request.status = "APPROVED"
    purchase_request.reviewed_by = request.user
    purchase_request.reviewed_at = timezone.now()

    purchase_request.save(
        update_fields=[
            "status",
            "reviewed_by",
            "reviewed_at",
        ]
    )

    if purchase_request.requested_by_id and Profile.objects.filter(
        user_id=purchase_request.requested_by_id,
        duka=purchase_request.duka,
    ).exists():

        Notification.objects.create(
            recipient=purchase_request.requested_by,
            title="Purchase Request Imeidhinishwa",
            message=(
                f"{purchase_request.display_number} "
                f"imeidhinishwa na "
                f"{request.user.username}."
            ),
            notification_type="PURCHASE_REQUEST",
            purchase_request=purchase_request,
        )

    create_activity(
        request.user,
        (
            f"Purchase request "
            f"{purchase_request.display_number} "
            "imeidhinishwa"
        ),
        "purchase_request",
    )

    messages.success(
        request,
        (
            f"{purchase_request.display_number} "
            "imeidhinishwa."
        ),
    )

    return redirect("dashboard")


@login_required(login_url="login")
def reject_purchase_request(
    request,
    request_id,
):
    if request.method != "POST":
        return redirect("dashboard")

    role = get_profile_role(
        request.user
    )

    if (
        role != "admin"
        and not request.user.is_staff
        and not request.user.is_superuser
    ):
        messages.error(
            request,
            (
                "Huna mamlaka ya "
                "kukataa purchase request."
            ),
        )
        return redirect("dashboard")

    duka = get_store_profile(request)
    if duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("dashboard")

    purchase_request = get_object_or_404(
        PurchaseRequest.objects.prefetch_related(
            "items__bidhaa"
        ),
        pk=request_id,
        duka=duka,
    )

    if purchase_request.status != "PENDING":
        messages.error(
            request,
            "Purchase request hii tayari imechakatwa.",
        )
        return redirect("dashboard")

    purchase_request.status = "REJECTED"
    purchase_request.reviewed_by = request.user
    purchase_request.reviewed_at = timezone.now()

    purchase_request.save(
        update_fields=[
            "status",
            "reviewed_by",
            "reviewed_at",
        ]
    )

    if purchase_request.requested_by_id and Profile.objects.filter(
        user_id=purchase_request.requested_by_id,
        duka=purchase_request.duka,
    ).exists():

        Notification.objects.create(
            recipient=purchase_request.requested_by,
            title="Purchase Request Imekataliwa",
            message=(
                f"{purchase_request.display_number} "
                f"imekataliwa na "
                f"{request.user.username}."
            ),
            notification_type="PURCHASE_REQUEST",
            purchase_request=purchase_request,
        )

    create_activity(
        request.user,
        (
            f"Purchase request "
            f"{purchase_request.display_number} "
            "imekataliwa"
        ),
        "purchase_request",
    )

    messages.info(
        request,
        (
            f"{purchase_request.display_number} "
            "imekataliwa."
        ),
    )

    return redirect("dashboard")


@login_required(login_url="login")
def mark_notification_read(
    request,
    notification_id,
):
    duka = get_store_profile(request)
    tenant_filter = Q(purchase_request__isnull=True)
    if duka is not None:
        tenant_filter |= Q(purchase_request__duka=duka)
    else:
        tenant_filter = Q(pk__in=[])

    notification = get_object_or_404(
        Notification.objects.filter(tenant_filter),
        pk=notification_id,
        recipient=request.user,
    )

    notification.is_read = True

    notification.save(
        update_fields=[
            "is_read"
        ]
    )

    return redirect("dashboard")


@login_required(login_url="login")
def mark_all_notifications_read(request):
    if request.method == "POST":

        duka = get_store_profile(request)
        tenant_filter = Q(purchase_request__isnull=True)
        if duka is not None:
            tenant_filter |= Q(purchase_request__duka=duka)
        else:
            tenant_filter = Q(pk__in=[])

        Notification.objects.filter(
            Q(recipient=request.user)
            & Q(is_read=False)
            & tenant_filter,
        ).update(
            is_read=True
        )

        messages.success(
            request,
            "Notifications zote zimesomwa.",
        )

    return redirect("dashboard")


# =========================================================
# 9. USIMAMIZI WA WAFANYAKAZI
# =========================================================

@login_required(login_url="login")
def wafanyakazi_view(request):
    denied = require_roles(request, "admin")
    if denied:
        return denied

    duka = get_store_profile(request)
    wafanyakazi_wote = (
        User.objects
        .filter(profile__duka=duka)
        .select_related("profile")
        .order_by("username")
    ) if duka else User.objects.none()

    idadi_online = sum(
        1
        for mfanyakazi in wafanyakazi_wote
        if hasattr(
            mfanyakazi,
            "profile",
        )
        and mfanyakazi.profile.yuko_online
    )

    active_employees = (
        wafanyakazi_wote
        .filter(
            is_active=True
        )
        .count()
    )

    blocked_employees = (
        wafanyakazi_wote
        .filter(
            is_active=False
        )
        .count()
    )

    context = {
        "title": "Usimamizi wa Wafanyakazi",
        "wafanyakazi": wafanyakazi_wote,
        "idadi_online": idadi_online,
        "total_employees": (
            wafanyakazi_wote.count()
        ),
        "active_employees": (
            active_employees
        ),
        "blocked_employees": (
            blocked_employees
        ),
    }

    return render(
        request,
        "wafanyakazi.html",
        context,
    )


@login_required(login_url="login")
def sajili_mfanyakazi(request):
    if (
        get_profile_role(request.user) != "admin"
        and not request.user.is_superuser
    ):
        messages.error(
            request,
            "Huna mamlaka ya kusajili mfanyakazi.",
        )
        return redirect("wafanyakazi")

    if request.method != "POST":
        return redirect("wafanyakazi")

    jina = (request.POST.get("username") or "").strip()
    password = request.POST.get("password") or ""
    jukumu = (
        request.POST.get("jukumu") or "cashier"
    ).strip().lower()

    allowed_roles = {"admin", "cashier", "stoo"}
    if jukumu not in allowed_roles:
        messages.error(
            request,
            "Jukumu lililochaguliwa si sahihi.",
        )
        return redirect("wafanyakazi")

    if not jina or not password:
        messages.error(
            request,
            "Username na password vinahitajika.",
        )
        return redirect("wafanyakazi")

    # Read the permission controls that are already posted by the existing UI.
    # The selected role provides a safe fallback when a role checkbox is absent.
    ni_admin = (
        request.POST.get("ni_admin", "").strip().lower()
        in {"true", "on", "1", "yes"}
    )

    def posted_flag(name):
        return (
            request.POST.get(name, "").strip().lower()
            in {"true", "on", "1", "yes"}
        )

    ruhusa_mauzo = posted_flag("ruhusa_mauzo") or jukumu == "cashier"
    ruhusa_stoo = posted_flag("ruhusa_stoo") or jukumu == "stoo"
    ruhusa_wateja = posted_flag("ruhusa_wateja")
    ruhusa_matumizi = posted_flag("ruhusa_matumizi")
    ruhusa_ripoti = posted_flag("ruhusa_ripoti")
    ruhusa_mawasiliano = posted_flag("ruhusa_mawasiliano")

    if jukumu == "admin":
        ni_admin = True

    if ni_admin:
        ruhusa_mauzo = True
        ruhusa_stoo = True
        ruhusa_wateja = True
        ruhusa_matumizi = True
        ruhusa_ripoti = True
        ruhusa_mawasiliano = True

    owner_duka = get_store_profile(request)
    if owner_duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("wafanyakazi")

    if User.objects.filter(username=jina).exists():
        messages.error(
            request,
            (
                f"Jina la '{jina}' tayari linatumika. "
                "Tumia jina jingine."
            ),
        )
        return redirect("wafanyakazi")

    with transaction.atomic():
        user_mpya = User.objects.create_user(
            username=jina,
            password=password,
        )

        # Profile is normally created by the model signal; keep a safe fallback.
        profile = getattr(user_mpya, "profile", None)
        if not profile:
            profile = Profile.objects.create(user=user_mpya)

        # Preserve both the new permission flags and the legacy role flags.
        profile.ni_admin = ni_admin
        profile.ruhusa_mauzo = ruhusa_mauzo
        profile.ruhusa_stoo = ruhusa_stoo
        profile.ruhusa_wateja = ruhusa_wateja
        profile.ruhusa_matumizi = ruhusa_matumizi
        profile.ruhusa_ripoti = ruhusa_ripoti
        profile.ruhusa_mawasiliano = ruhusa_mawasiliano

        profile.ni_cashier = ruhusa_mauzo
        profile.ni_stoo = ruhusa_stoo

        update_fields = [
            "ni_admin",
            "ruhusa_mauzo",
            "ruhusa_stoo",
            "ruhusa_wateja",
            "ruhusa_matumizi",
            "ruhusa_ripoti",
            "ruhusa_mawasiliano",
            "ni_cashier",
            "ni_stoo",
        ]

        # Attach the staff member to the admin's existing Duka.
        if hasattr(profile, "duka_id") and owner_duka is not None:
            profile.duka = owner_duka
            update_fields.append("duka")

        profile.save(update_fields=update_fields)

    create_activity(
        request.user,
        f"Mfanyakazi mpya amesajiliwa: {user_mpya.username}",
        "wafanyakazi",
    )

    messages.success(
        request,
        (
            f"Mfanyakazi {user_mpya.username} "
            "amesajiliwa kikamilifu!"
        ),
    )

    return redirect("wafanyakazi")


@login_required(login_url="login")
def badili_hali_mfanyakazi(
    request,
    user_id,
):
    if (
        get_profile_role(request.user)
        != "admin"
        and not request.user.is_superuser
    ):
        messages.error(
            request,
            (
                "Huna mamlaka ya "
                "kubadilisha hali ya mfanyakazi."
            ),
        )

        return redirect(
            "wafanyakazi"
        )

    duka = get_store_profile(request)
    mfanyakazi_user = get_object_or_404(
        User.objects.filter(profile__duka=duka),
        pk=user_id,
    )

    if mfanyakazi_user == request.user:
        messages.error(
            request,
            "Huwezi kujizuia mwenyewe.",
        )

        return redirect(
            "wafanyakazi"
        )

    mfanyakazi_user.is_active = (
        not mfanyakazi_user.is_active
    )

    mfanyakazi_user.save(
        update_fields=[
            "is_active"
        ]
    )

    hali = (
        "Amezuiwa"
        if not mfanyakazi_user.is_active
        else "Ameruhusiwa"
    )

    create_activity(
        request.user,
        (
            f"Hali ya "
            f"{mfanyakazi_user.username} "
            f"imebadilishwa kuwa "
            f"{hali}"
        ),
        "wafanyakazi",
    )

    messages.info(
        request,
        (
            f"Hali ya "
            f"{mfanyakazi_user.username} "
            f"imebadilishwa kuwa: "
            f"{hali}"
        ),
    )

    return redirect(
        "wafanyakazi"
    )


@login_required(login_url="login")
def futa_mfanyakazi(
    request,
    user_id,
):
    if (
        get_profile_role(request.user)
        != "admin"
        and not request.user.is_superuser
    ):
        messages.error(
            request,
            "Huna mamlaka ya kufuta mfanyakazi.",
        )

        return redirect(
            "wafanyakazi"
        )

    if request.method != "POST":
        messages.error(
            request,
            "Tumia POST kufuta mfanyakazi.",
        )

        return redirect(
            "wafanyakazi"
        )

    duka = get_store_profile(request)
    mfanyakazi = get_object_or_404(
        User.objects.filter(profile__duka=duka),
        pk=user_id,
    )

    if mfanyakazi == request.user:
        messages.error(
            request,
            (
                "Huwezi kujifuta mwenyewe "
                "kwenye mfumo!"
            ),
        )

        return redirect(
            "wafanyakazi"
        )

    jina = mfanyakazi.username

    # SURGERY FIX: Kamata ProtectedError kuzuia mfumo ku-crash
    try:
        mfanyakazi.delete()
    except ProtectedError:
        messages.error(
            request,
            (
                f"Huwezi kumfuta '{jina}' kabisa kwa sababu tayari "
                "ana historia ya miamala kwenye mfumo (k.m. aliwahi kuuza au kuingiza stock). "
                "Tafadhali tumia kitufe cha 'Zulia / Block' kumsimamisha kazi badala ya kumfuta."
            ),
        )
        return redirect("wafanyakazi")
    except Exception as e:
        messages.error(
            request,
            f"Imeshindikana kumfuta mfanyakazi. Kosa: {str(e)}",
        )
        return redirect("wafanyakazi")

    create_activity(
        request.user,
        (
            f"Mfanyakazi amefutwa: "
            f"{jina}"
        ),
        "wafanyakazi",
    )

    messages.success(
        request,
        (
            f"Mfanyakazi '{jina}' "
            "amefutwa kikamilifu kwenye mfumo!"
        ),
    )

    return redirect(
        "wafanyakazi"
    )


# =========================================================
# =========================================================
# 10. REGISTRATION / ONBOARDING
# =========================================================

def register_view(request):
    """
    TradeCore public account + business onboarding.

    Creates, atomically:
        User -> Profile (auto-created by the model signal) -> Duka

    The registration UI keeps Username as the short account identity used
    by the avatar, while Email is also accepted by login_view as the login
    identity.  Business information is persisted in Duka for the header,
    receipts and store-level settings.
    """
    if request.user.is_authenticated:
        return redirect_by_role(request.user)

    if request.method != "POST":
        # REGISTRATION UI SURGERY: discard stale flash messages left by
        # another page/session action (e.g. product delete/status notices).
        # Validation errors created below are rendered immediately and remain
        # visible to the user without polluting the clean onboarding screen.
        list(get_messages(request))
        return render(
            request,
            "registration.html",
            {
                "business_type_choices": BUSINESS_TYPE_CHOICES,
            },
        )

    # ---------------------------------------------------------
    # ACCOUNT
    # ---------------------------------------------------------
    full_name = (request.POST.get("full_name") or "").strip()
    phone = (request.POST.get("phone") or "").strip()
    email = (request.POST.get("email") or "").strip().lower()
    username = (request.POST.get("username") or "").strip().lower()
    password = request.POST.get("password") or ""
    confirm_password = request.POST.get("confirm_password") or ""

    # ---------------------------------------------------------
    # BUSINESS
    # ---------------------------------------------------------
    business_name = (request.POST.get("business_name") or "").strip()
    business_type = (request.POST.get("business_type") or "GENERAL").strip().upper()
    location = (request.POST.get("location") or "").strip()
    business_address = (request.POST.get("business_address") or "").strip()
    business_phone = (request.POST.get("business_phone") or "").strip()
    business_email = (request.POST.get("business_email") or "").strip().lower()
    tin = (request.POST.get("tin") or "").strip()

    # ---------------------------------------------------------
    # RECEIPT / STORE IDENTITY
    # ---------------------------------------------------------
    logo = request.FILES.get("logo")
    tagline = (request.POST.get("tagline") or "").strip()
    terminal_name = (request.POST.get("terminal_name") or "").strip() or "POS-01"
    receipt_footer = (
        (request.POST.get("receipt_footer") or "").strip()
        or "Asante kwa kufanya manunuzi nasi!"
    )
    currency = (request.POST.get("currency") or "TZS").strip().upper() or "TZS"
    agreed = request.POST.get("agree") == "on"

    # ---------------------------------------------------------
    # SERVER-SIDE VALIDATION
    # ---------------------------------------------------------
    errors = []

    if not full_name:
        errors.append("Jaza jina lako kamili.")

    # Personal phone is collected by the current UI. The current Profile
    # model has no dedicated phone column, so we validate it here but do not
    # silently write it into an unrelated field.
    phone_digits = re.sub(r"\D+", "", phone)
    if not phone or len(phone_digits) < 7:
        errors.append("Weka namba sahihi ya simu.")

    if not email:
        errors.append("Email ni lazima kwa login ya account.")
    elif not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", email):
        errors.append("Weka email sahihi.")

    if not username:
        errors.append("Weka username ya account.")
    elif not re.fullmatch(r"[a-z0-9_.-]{3,30}", username):
        errors.append("Username itumie herufi ndogo, namba, underscore, dot au dash (3–30 characters).")

    if len(password) < 8:
        errors.append("Nenosiri liwe angalau characters 8.")
    if password != confirm_password:
        errors.append("Nenosiri na confirmation hazilingani.")

    if not business_name:
        errors.append("Weka jina la biashara.")
    if business_type not in dict(BUSINESS_TYPE_CHOICES):
        errors.append("Aina ya biashara haijatambuliwa.")
    if not location:
        errors.append("Weka location / eneo la biashara.")

    business_phone_digits = re.sub(r"\D+", "", business_phone)
    if not business_phone or len(business_phone_digits) < 7:
        errors.append("Weka namba sahihi ya simu ya biashara.")

    if business_email and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", business_email):
        errors.append("Email ya biashara si sahihi.")

    if not agreed:
        errors.append("Kubali masharti ya demo kabla ya kutengeneza account.")

    if logo:
        # Small guard for the public registration endpoint.
        if getattr(logo, "size", 0) > 5 * 1024 * 1024:
            errors.append("Logo isiwe kubwa kuliko 5MB.")
        content_type = (getattr(logo, "content_type", "") or "").lower()
        if content_type and content_type not in {"image/png", "image/jpeg", "image/webp"}:
            errors.append("Logo itumie PNG, JPG au WEBP.")

    if errors:
        for error in errors:
            messages.error(request, error)
        return render(
            request,
            "registration.html",
            {
                "business_type_choices": BUSINESS_TYPE_CHOICES,
            },
        )

    # ---------------------------------------------------------
    # DUPLICATE CHECKS
    # ---------------------------------------------------------
    if User.objects.filter(username__iexact=username).exists():
        messages.error(request, f"Username '{username}' tayari inatumika. Tumia username nyingine.")
        return render(
            request,
            "registration.html",
            {"business_type_choices": BUSINESS_TYPE_CHOICES},
        )

    if User.objects.filter(email__iexact=email).exists():
        messages.error(request, "Email hiyo tayari imesajiliwa. Tumia email nyingine au ingia kwenye account iliyopo.")
        return render(
            request,
            "registration.html",
            {"business_type_choices": BUSINESS_TYPE_CHOICES},
        )

    if business_email and User.objects.filter(email__iexact=business_email).exists() and business_email != email:
        # Business email does not have to equal an account email, so this is
        # intentionally NOT a duplicate blocker. Duka.email is independent.
        pass

    # ---------------------------------------------------------
    # CREATE ACCOUNT + STORE ATOMICALLY
    # ---------------------------------------------------------
    try:
        with transaction.atomic():
            user = User.objects.create_user(
                username=username,
                email=email,
                password=password,
            )

            # Full name is account/profile identity. Avatar/header can keep
            # using user.username without exposing the full legal name.
            name_parts = full_name.split()
            user.first_name = name_parts[0] if name_parts else ""
            user.last_name = " ".join(name_parts[1:]) if len(name_parts) > 1 else ""
            user.save(update_fields=["first_name", "last_name", "email"])

            # The post_save signal normally creates Profile. Get-or-create
            # keeps this safe even if the signal is changed later.
            profile, _ = Profile.objects.get_or_create(user=user)
            profile.ni_admin = True
            profile.ni_cashier = False
            profile.ni_stoo = False
            profile.ruhusa_mauzo = True
            profile.ruhusa_stoo = True
            profile.ruhusa_wateja = True
            profile.ruhusa_matumizi = True
            profile.ruhusa_ripoti = True
            profile.ruhusa_mawasiliano = True
            profile.save(
                update_fields=[
                    "ni_admin",
                    "ni_cashier",
                    "ni_stoo",
                    "ruhusa_mauzo",
                    "ruhusa_stoo",
                    "ruhusa_wateja",
                    "ruhusa_matumizi",
                    "ruhusa_ripoti",
                    "ruhusa_mawasiliano",
                ]
            )

            store_address = location
            if business_address:
                store_address = f"{location}, {business_address}"

            duka = Duka.objects.create(
                mwenye_duka=user,
                jina_la_duka=business_name,
                anwani_au_mahali=store_address,
                simu=business_phone,
                tin=tin or None,
                email=business_email or None,
                logo=logo,
                tagline=tagline or None,
                receipt_footer=receipt_footer,
                terminal_name=terminal_name,
                currency=currency,
                business_type=business_type,
                inventory_location_mode=True,
            )

            # Owner Profile -> Duka. This is the link used by staff/store
            # resolution in the rest of TradeCore.
            if hasattr(profile, "duka_id"):
                profile.duka = duka
                profile.save(update_fields=["duka"])

            # Prepare categories and the default Store/Shop locations for the
            # new business so the first dashboard/POS visit has a ready setup.
            ensure_business_categories(business_type)
            try:
                ensure_inventory_locations(duka)
            except Exception:
                logger.exception("Default inventory locations could not be prepared for new store id=%s", duka.id)

        # -----------------------------------------------------
        # AUTO-LOGIN + FIRST-WELCOME STATE
        # -----------------------------------------------------
        login(request, user)
        request.session.set_expiry(0)

        LoginHistory.objects.create(
            user=user,
            username_attempt=email,
            ip_address=get_client_ip(request),
            user_agent=(request.META.get("HTTP_USER_AGENT") or "")[:255],
            status="SUCCESS",
        )

        create_activity(
            user,
            f"Amefungua TradeCore na kusajili duka: {business_name}",
            "auth",
        )

        request.session["tradecore_store_name"] = business_name
        request.session["tradecore_welcome_user"] = username
        request.session["tradecore_welcome_role"] = "admin"
        request.session["tradecore_welcome_type"] = "trial"

        target = redirect_by_role(user)
        target_url = target.url if hasattr(target, "url") else reverse("dashboard")
        separator = "&" if "?" in target_url else "?"
        return redirect(f"{target_url}{separator}welcome=1")

    except Exception as exc:
        logger.exception("TradeCore registration failed")
        messages.error(
            request,
            "Account haikuundwa kwa sasa. Tafadhali kagua taarifa zako kisha ujaribu tena.",
        )
        return render(
            request,
            "registration.html",
            {"business_type_choices": BUSINESS_TYPE_CHOICES},
        )


# 10. LOGIN / LOGOUT
# =========================================================

def login_view(request):
    if request.user.is_authenticated:
        return redirect_by_role(
            request.user
        )

    if request.method == "POST":

        username_input = (
            request.POST.get(
                "username"
            )
            or ""
        ).strip()

        password_input = (
            request.POST.get(
                "password"
            )
            or ""
        )

        remember_me = request.POST.get(
            "remember_me"
        )

        ip_address = get_client_ip(
            request
        )

        user_agent = (
            request.META.get(
                "HTTP_USER_AGENT"
            )
            or ""
        )[:255]

        # Login accepts Email OR Username. New registrations use a unique email
        # as the primary login identity while retaining username for the
        # avatar/profile and backward compatibility with existing accounts.
        login_identity = username_input
        matched_email_user = (
            User.objects
            .filter(email__iexact=username_input)
            .order_by("id")
            .first()
        )
        if matched_email_user is not None:
            login_identity = matched_email_user.username

        user = authenticate(
            request,
            username=login_identity,
            password=password_input,
        )

        if user is not None:

            if user.is_active:

                login(
                    request,
                    user,
                )

                if remember_me:
                    request.session.set_expiry(
                        1209600
                    )
                else:
                    request.session.set_expiry(
                        0
                    )

                LoginHistory.objects.create(
                    user=user,
                    username_attempt=username_input,
                    ip_address=ip_address,
                    user_agent=user_agent,
                    status="SUCCESS",
                )

                if hasattr(
                    user,
                    "profile"
                ):

                    user.profile.muda_wa_mwisho = (
                        timezone.now()
                    )

                    user.profile.save(
                        update_fields=[
                            "muda_wa_mwisho"
                        ]
                    )

                create_activity(
                    user,
                    "Ameingia kwenye mfumo",
                    "auth",
                )

                # Per-login UI state: store label + one-time welcome popup.
                role = get_profile_role(user)
                store_name = ""
                try:
                    duka = get_store_profile(request)
                except Exception:
                    duka = None

                if duka is None:
                    profile = getattr(user, "profile", None)
                    profile_duka = getattr(profile, "duka", None) if profile else None
                    if profile_duka is not None:
                        duka = profile_duka

                # Legacy/single-store fallback. With multiple stores we never guess.
                if duka is None:
                    try:
                        if Duka.objects.count() == 1:
                            duka = Duka.objects.first()
                    except Exception:
                        duka = None

                if duka is not None:
                    store_name = str(getattr(duka, "jina_la_duka", "") or "").strip()

                request.session["tradecore_store_name"] = store_name
                request.session["tradecore_welcome_user"] = str(
                    user.get_full_name() or user.username or ""
                ).strip()
                request.session["tradecore_welcome_role"] = role or "staff"
                # Normal login uses the standard welcome message. Registration/payment
                # flows set a more specific welcome type after a successful event.
                request.session["tradecore_welcome_type"] = "login"

                target = redirect_by_role(user)
                target_url = target.url if hasattr(target, "url") else reverse("dashboard")
                separator = "&" if "?" in target_url else "?"
                return redirect(f"{target_url}{separator}welcome=1")

            LoginHistory.objects.create(
                username_attempt=username_input,
                ip_address=ip_address,
                user_agent=user_agent,
                status="LOCKED",
            )

            return render(
                request,
                "login.html",
                {
                    "error": True
                },
            )

        LoginHistory.objects.create(
            username_attempt=username_input,
            ip_address=ip_address,
            user_agent=user_agent,
            status="FAILED",
        )

        return render(
            request,
            "login.html",
            {
                "error": True
            },
        )

    return render(
        request,
        "login.html",
    )


@login_required(login_url="login")
def logout_view(request):
    if hasattr(
        request.user,
        "profile",
    ):

        request.user.profile.muda_wa_mwisho = (
            timezone.now()
        )

        request.user.profile.save(
            update_fields=[
                "muda_wa_mwisho"
            ]
        )

    create_activity(
        request.user,
        "Ametoka kwenye mfumo",
        "auth",
    )

    request.session.pop("tradecore_store_name", None)
    request.session.pop("tradecore_welcome_user", None)
    request.session.pop("tradecore_welcome_role", None)

    logout(request)

    return redirect(
        "login"
    )


# =========================================================
# 11. SALE RETURNS / REFUNDS
# =========================================================

@login_required(login_url="login")
def approve_sale_return(
    request,
    return_id,
):
    if request.method != "POST":
        messages.error(
            request,
            "Tumia POST kuidhinisha refund.",
        )
        return redirect("ripoti")

    role = get_profile_role(
        request.user
    )

    if role != "admin" and not request.user.is_superuser:
        messages.error(
            request,
            "Huna mamlaka ya kuidhinisha refund.",
        )
        return redirect("ripoti")

    duka = get_store_profile(request)
    if duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("ripoti")

    try:
        with transaction.atomic():

            sale_return = (
                SaleReturn.objects
                .select_for_update()
                .select_related(
                    "sale",
                    "sale__bidhaa",
                )
                .get(pk=return_id, sale__duka=duka)
            )

            if sale_return.status != "PENDING":
                messages.error(
                    request,
                    (
                        "Refund hii tayari "
                        "imekwisha fanyiwa kazi."
                    ),
                )

                return redirect(
                    "ripoti"
                )

            sale = (
                Mauzo.objects
                .select_for_update()
                .select_related(
                    "bidhaa"
                )
                .get(
                    pk=sale_return.sale_id,
                    duka=duka,
                )
            )

            if sale.idadi <= 0:
                messages.error(
                    request,
                    (
                        "Mauzo haya yana quantity "
                        "isiyoweza kurejeshwa."
                    ),
                )
                return redirect(
                    "ripoti"
                )

            already_returned = (
                SaleReturn.objects
                .filter(
                    sale=sale,
                    status="APPROVED",
                )
                .exclude(
                    pk=sale_return.pk
                )
                .aggregate(
                    total=Sum(
                        "quantity_returned"
                    )
                )["total"]
                or 0
            )

            available_to_return = (
                sale.idadi
                - already_returned
            )

            if (
                sale_return.quantity_returned
                > available_to_return
            ):
                messages.error(
                    request,
                    (
                        "Idadi ya bidhaa "
                        "inayorejeshwa imezidi "
                        "iliyobaki."
                    ),
                )

                return redirect(
                    "ripoti"
                )

            unit_refund = (
                sale.jumla_pesa_iliyopokelewa
                / sale.idadi
            )

            refund_amount = (
                unit_refund
                * sale_return.quantity_returned
            )

            sale_return.refund_amount = (
                refund_amount.quantize(
                    Decimal("0.01")
                )
            )

            sale_return.status = "APPROVED"

            sale_return.approved_by = (
                request.user
            )

            sale_return.approved_at = (
                timezone.now()
            )

            sale_return.save(
                update_fields=[
                    "refund_amount",
                    "status",
                    "approved_by",
                    "approved_at",
                ]
            )

            bidhaa = (
                Bidhaa.objects
                .select_for_update()
                .get(
                    pk=sale.bidhaa_id
                )
            )

            if sale.stock_location_id and sale.duka_id and getattr(sale.duka, "inventory_location_mode", False):
                add_stock(
                    duka=sale.duka,
                    bidhaa=bidhaa,
                    location=sale.stock_location,
                    quantity=sale_return.quantity_returned,
                    actor=request.user,
                    reference=f"RETURN-{sale_return.id}",
                    reason=sale_return.reason or "Sale Return",
                    movement_type="RETURN",
                )
            else:
                bidhaa.idadi_stoo += (
                    sale_return.quantity_returned
                )

                bidhaa.save(
                    update_fields=[
                        "idadi_stoo"
                    ]
                )

            create_activity(
                request.user,
                (
                    "Return imeidhinishwa: "
                    f"{sale.bidhaa.jina_la_bidhaa} × "
                    f"{sale_return.quantity_returned}, "
                    f"Refund TZS "
                    f"{sale_return.refund_amount:,.2f}"
                ),
                "refund",
            )

    except SaleReturn.DoesNotExist:
        messages.error(
            request,
            "Return hiyo haikupatikana.",
        )
        return redirect("ripoti")

    except Exception:
        logger.exception(
            "Sale return approval failed: return_id=%s user=%s",
            return_id,
            request.user.username,
        )

        messages.error(
            request,
            (
                "Refund haijaidhinishwa "
                "kwa sasa. Tafadhali jaribu tena."
            ),
        )

        return redirect(
            "ripoti"
        )

    messages.success(
        request,
        (
            "Return imeidhinishwa, "
            "stock imerudishwa na "
            "refund imerekodiwa."
        ),
    )

    return redirect(
        "ripoti"
    )


@login_required(login_url="login")
def reject_sale_return(
    request,
    return_id,
):
    if request.method != "POST":
        messages.error(
            request,
            "Tumia POST kukataa refund.",
        )
        return redirect("ripoti")

    role = get_profile_role(
        request.user
    )

    if role != "admin" and not request.user.is_superuser:
        messages.error(
            request,
            "Huna mamlaka ya kukataa refund.",
        )
        return redirect(
            "ripoti"
        )

    duka = get_store_profile(request)
    if duka is None:
        messages.error(request, "Biashara ya mtumiaji haijapatikana.")
        return redirect("ripoti")

    sale_return = get_object_or_404(
        SaleReturn,
        pk=return_id,
        status="PENDING",
        sale__duka=duka,
    )

    sale_return.status = "REJECTED"

    sale_return.approved_by = (
        request.user
    )

    sale_return.approved_at = (
        timezone.now()
    )

    sale_return.save(
        update_fields=[
            "status",
            "approved_by",
            "approved_at",
        ]
    )

    create_activity(
        request.user,
        (
            f"Refund imekataliwa: "
            f"Return #{sale_return.id}"
        ),
        "refund",
    )

    messages.info(
        request,
        "Refund imekataliwa.",
    )

    return redirect(
        "ripoti"
    )
    
@login_required(login_url="login")
@require_POST
def upload_profile_picture(request):
    try:
        image_data = (request.POST.get("image_data") or "").strip()
        if not image_data:
            return JsonResponse({
                "status": "error",
                "message": "Hakuna picha iliyopokelewa",
            }, status=400)

        # Keep the existing data-URI upload contract, but validate it strictly.
        try:
            header, imgstr = image_data.split(";base64,", 1)
        except ValueError:
            return JsonResponse({
                "status": "error",
                "message": "Format ya picha si sahihi.",
            }, status=400)

        mime_type = header.removeprefix("data:").strip().lower()
        allowed_types = {
            "image/jpeg": "jpg",
            "image/png": "png",
            "image/webp": "webp",
        }
        ext = allowed_types.get(mime_type)
        if not ext:
            return JsonResponse({
                "status": "error",
                "message": "Tumia JPG, PNG au WEBP.",
            }, status=400)

        try:
            raw = base64.b64decode(imgstr, validate=True)
        except (ValueError, base64.binascii.Error):
            return JsonResponse({
                "status": "error",
                "message": "Picha iliyotumwa si sahihi.",
            }, status=400)

        max_bytes = 5 * 1024 * 1024
        if not raw or len(raw) > max_bytes:
            return JsonResponse({
                "status": "error",
                "message": "Picha isiwe kubwa kuliko 5MB.",
            }, status=400)

        profile = getattr(request.user, "profile", None)
        if profile is None:
            profile = Profile.objects.create(user=request.user)

        profile.picha.save(
            f"profile_{request.user.id}.{ext}",
            ContentFile(raw),
            save=True,
        )

        return JsonResponse({"status": "success"})
    except Exception:
        logger.exception(
            "Profile picture upload failed for user=%s",
            getattr(request.user, "id", None),
        )
        return JsonResponse({
            "status": "error",
            "message": "Imeshindikana kuhifadhi picha.",
        }, status=500)


def _set_welcome_state(request, welcome_type, user=None):
    """Prepare the existing TradeCore welcome modal for a one-time success state."""
    request.session["tradecore_welcome_type"] = welcome_type
    if user is not None:
        request.session["tradecore_welcome_user"] = str(
            user.get_full_name() or user.username or ""
        ).strip()
    request.session.modified = True


def _flutterwave_secret_key():
    return (getattr(settings, "FLW_SECRET_KEY", None) or os.getenv("FLW_SECRET_KEY", "")).strip()


def _flutterwave_secret_hash():
    return (getattr(settings, "FLW_SECRET_HASH", None) or os.getenv("FLW_SECRET_HASH", "")).strip()


def _payment_api_headers():
    secret_key = _flutterwave_secret_key()
    if not secret_key:
        raise RuntimeError("FLW_SECRET_KEY haijawekwa kwenye environment.")
    return {
        "Authorization": f"Bearer {secret_key}",
        "Content-Type": "application/json",
        "Accept": "application/json",
    }


def _normalize_tanzania_phone(value):
    digits = re.sub(r"\D+", "", str(value or ""))
    if digits.startswith("+255"):
        digits = digits[1:]
    if digits.startswith("255"):
        local = digits[3:]
    elif digits.startswith("0"):
        local = digits[1:]
    else:
        local = digits
    if len(local) != 9 or local[0] not in "67":
        return ""
    return f"255{local}"


_FLUTTERWAVE_NETWORKS = {
    "AIRTEL": {"label": "Airtel Money", "api": "Airtel"},
    "TIGO": {"label": "Tigo / Mixx by Yas", "api": "Tigo"},
    "HALOPESA": {"label": "HaloPesa", "api": "Halopesa"},
    "VODAFONE": {"label": "Vodacom M-Pesa", "api": "Vodafone"},
}


def _activate_payment_transaction(payment, gateway_payload, transaction_id=None):
    """Activate Premium exactly once after a verified successful gateway payment."""
    data = gateway_payload.get("data") if isinstance(gateway_payload, dict) else None
    data = data if isinstance(data, dict) else {}

    tx_ref = str(data.get("tx_ref") or payment.tx_ref).strip()
    status = str(data.get("status") or "").strip().lower()
    currency = str(data.get("currency") or payment.currency).upper()
    try:
        paid_amount = Decimal(str(data.get("amount")))
    except (InvalidOperation, TypeError, ValueError):
        paid_amount = Decimal("0.00")

    if tx_ref != payment.tx_ref:
        payment.status = "FAILED"
        payment.failure_reason = "Transaction reference haifanani."
        payment.gateway_payload = gateway_payload
        payment.gateway_status = status
        payment.save(update_fields=["status", "failure_reason", "gateway_payload", "gateway_status", "updated_at"])
        return False

    if status != "successful":
        mapped = "FAILED" if status in {"failed", "cancelled", "cancelled_by_user"} else "PENDING"
        payment.status = mapped
        payment.gateway_status = status
        payment.gateway_payload = gateway_payload
        if mapped == "FAILED":
            payment.failure_reason = str(data.get("processor_response") or "Malipo hayakukamilika.")[:255]
        payment.save(update_fields=["status", "gateway_status", "gateway_payload", "failure_reason", "updated_at"])
        return False

    if currency != payment.currency or paid_amount != payment.amount:
        payment.status = "FAILED"
        payment.failure_reason = "Amount au currency ya malipo haijathibitishwa."
        payment.gateway_payload = gateway_payload
        payment.gateway_status = status
        payment.save(update_fields=["status", "failure_reason", "gateway_payload", "gateway_status", "updated_at"])
        return False

    with transaction.atomic():
        locked = PaymentTransaction.objects.select_for_update().select_related("subscription").get(pk=payment.pk)
        if locked.status == "SUCCESSFUL":
            return True

        stored_meta = (
            locked.gateway_payload.get("_tradecore", {})
            if isinstance(locked.gateway_payload, dict)
            else {}
        )
        try:
            billing_months = int(stored_meta.get("billing_months") or 1)
        except (TypeError, ValueError):
            billing_months = 1

        expected_pricing = calculate_premium_pricing(billing_months)
        if locked.amount != expected_pricing["total"]:
            locked.status = "FAILED"
            locked.failure_reason = "Transaction amount haifanani na duration iliyochaguliwa."
            locked.save(update_fields=["status", "failure_reason", "updated_at"])
            return False

        now = timezone.now()
        merged_gateway_payload = dict(gateway_payload) if isinstance(gateway_payload, dict) else {}
        merged_gateway_payload["_tradecore"] = {
            "billing_months": billing_months,
            "unit_price": str(expected_pricing["unit_price"]),
            "discount": str(expected_pricing["discount"]),
            "standard_total": str(expected_pricing["standard_total"]),
        }

        locked.status = "SUCCESSFUL"
        locked.gateway_transaction_id = str(transaction_id or data.get("id") or "")
        locked.gateway_reference = str(data.get("flw_ref") or "")[:150]
        locked.gateway_status = status
        locked.gateway_payload = merged_gateway_payload
        locked.failure_reason = ""
        locked.completed_at = now
        locked.save(update_fields=[
            "status",
            "gateway_transaction_id",
            "gateway_reference",
            "gateway_status",
            "gateway_payload",
            "failure_reason",
            "completed_at",
            "updated_at",
        ])

        subscription = locked.subscription
        subscription.activate_premium(payment_time=now, months=billing_months)
        subscription.save(update_fields=[
            "plan_type",
            "is_active",
            "premium_start_date",
            "premium_end_date",
            "last_payment_at",
        ])

    return True


def _verify_flutterwave_transaction(transaction_id, expected_tx_ref):
    response = requests.get(
        f"https://api.flutterwave.com/v3/transactions/{int(transaction_id)}/verify",
        headers=_payment_api_headers(),
        timeout=20,
    )
    try:
        payload = response.json()
    except ValueError:
        payload = {"status": "error", "message": response.text[:255]}
    if response.status_code >= 400:
        raise RuntimeError(str(payload.get("message") or "Flutterwave verification failed."))

    data = payload.get("data") if isinstance(payload, dict) else None
    if not isinstance(data, dict) or str(data.get("tx_ref") or "") != expected_tx_ref:
        raise RuntimeError("Verification reference haifanani na transaction ya TradeCore.")
    return payload


PREMIUM_STANDARD_MONTHLY = Decimal("20000.00")
PREMIUM_LONG_TERM_MONTHLY = Decimal("15000.00")
PREMIUM_LONG_TERM_THRESHOLD = 6
PREMIUM_MAX_MONTHS = 60


def calculate_premium_pricing(months):
    """Server-owned Premium pricing: 1-6 months at 20K/month, 7-60 at 15K/month."""
    try:
        months = int(months)
    except (TypeError, ValueError):
        raise ValueError("Idadi ya miezi si sahihi.")
    if months < 1 or months > PREMIUM_MAX_MONTHS:
        raise ValueError(f"Chagua kati ya mwezi 1 na miezi {PREMIUM_MAX_MONTHS}.")
    unit_price = (
        PREMIUM_LONG_TERM_MONTHLY
        if months > PREMIUM_LONG_TERM_THRESHOLD
        else PREMIUM_STANDARD_MONTHLY
    )
    standard_total = PREMIUM_STANDARD_MONTHLY * months
    total = unit_price * months
    return {
        "months": months,
        "unit_price": unit_price,
        "standard_total": standard_total,
        "discount": standard_total - total,
        "total": total,
        "long_term": months > PREMIUM_LONG_TERM_THRESHOLD,
    }


def premium_pricing_options():
    options = []
    for months, label, badge in (
        (1, "1 Mwezi", ""),
        (3, "3 Miezi", ""),
        (6, "6 Miezi", ""),
        (12, "12 Miezi", "Best Value"),
    ):
        pricing = calculate_premium_pricing(months)
        options.append({
            "months": pricing["months"],
            "label": label,
            "badge": badge,
            "unit_price": pricing["unit_price"],
            "standard_total": pricing["standard_total"],
            "discount": pricing["discount"],
            "total": pricing["total"],
            "monthly_display": f"{pricing['unit_price']:,.0f}",
            "total_display": f"{pricing['total']:,.0f}",
            "discount_display": f"{pricing['discount']:,.0f}",
        })
    return options


# =========================================================
# TRADECORE CONTROL ROOM — OWNER / PLATFORM MONITORING
# =========================================================
def _control_money(value):
    """Compact TZS display for the private TradeCore owner dashboard."""
    value = Decimal(str(value or "0"))
    absolute = abs(value)
    if absolute >= Decimal("1000000000"):
        text = f"{value / Decimal('1000000000'):.1f}B"
    elif absolute >= Decimal("1000000"):
        text = f"{value / Decimal('1000000'):.1f}M"
    elif absolute >= Decimal("1000"):
        text = f"{value / Decimal('1000'):.1f}K"
    else:
        text = f"{value:.0f}"
    return text.replace(".0B", "B").replace(".0M", "M").replace(".0K", "K")


@login_required(login_url="login")
@require_GET
def control_room(request):
    """Private owner-only operational dashboard for TradeCore."""
    if not request.user.is_superuser:
        return HttpResponseForbidden("Access denied.")

    now = timezone.now()
    today = now.date()
    seven_days_ago = now - timedelta(days=7)
    thirty_days_ago = now - timedelta(days=30)
    five_minutes_ago = now - timedelta(minutes=5)

    # Platform-wide owner dashboard; do not change tenant-scoped business modules.
    users_qs = User.objects.filter(is_superuser=False)
    businesses_qs = Duka.objects.select_related("mwenye_duka").filter(mwenye_duka__is_superuser=False)
    subscriptions_qs = BusinessSubscription.objects.select_related("user")
    payments_qs = PaymentTransaction.objects.select_related("user", "subscription")

    q = (request.GET.get("q") or "").strip()
    if q:
        businesses_qs = businesses_qs.filter(
            Q(jina_la_duka__icontains=q)
            | Q(mwenye_duka__username__icontains=q)
            | Q(mwenye_duka__first_name__icontains=q)
            | Q(mwenye_duka__last_name__icontains=q)
            | Q(mwenye_duka__email__icontains=q)
        )

    total_users = users_qs.count()
    total_businesses = businesses_qs.count()
    registrations_today = users_qs.filter(date_joined__date=today).count()
    registrations_7d = users_qs.filter(date_joined__gte=seven_days_ago).count()
    registrations_30d = users_qs.filter(date_joined__gte=thirty_days_ago).count()

    online_users = Profile.objects.filter(
        user__is_superuser=False,
        muda_wa_mwisho__gte=five_minutes_ago,
    ).count()

    active_trial = subscriptions_qs.filter(
        plan_type="TRIAL",
        is_active=True,
        trial_start_date__lte=now,
        trial_start_date__gte=thirty_days_ago,
    ).count()
    active_premium = subscriptions_qs.filter(
        plan_type="PREMIUM",
        is_active=True,
        premium_end_date__gt=now,
    ).count()
    expired_trial = subscriptions_qs.filter(
        plan_type="TRIAL",
        trial_start_date__lt=thirty_days_ago,
    ).count()
    expired_premium = subscriptions_qs.filter(
        plan_type="PREMIUM",
    ).filter(Q(premium_end_date__isnull=True) | Q(premium_end_date__lte=now)).count()

    successful_payments = payments_qs.filter(status="SUCCESSFUL")
    revenue_total = successful_payments.aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
    revenue_today = successful_payments.filter(created_at__date=today).aggregate(total=Sum("amount"))["total"] or Decimal("0.00")
    pending_payments = payments_qs.filter(status__in=["INITIATED", "PENDING", "PENDING_BANK"]).count()
    failed_payments = payments_qs.filter(status__in=["FAILED", "CANCELLED"]).count()

    recent_businesses = list(
        businesses_qs.select_related("mwenye_duka")
        .order_by("-mwenye_duka__date_joined")[:50]
    )
    recent_payments = list(payments_qs.order_by("-created_at")[:20])
    live_users = list(
        Profile.objects.select_related("user")
        .filter(user__is_superuser=False, muda_wa_mwisho__gte=five_minutes_ago)
        .order_by("-muda_wa_mwisho")[:20]
    )

    for business in recent_businesses:
        owner = business.mwenye_duka
        business.owner_name = owner.get_full_name().strip() or owner.username
        business.owner_email = owner.email or "—"
        business.joined_local = timezone.localtime(owner.date_joined)

    for payment in recent_payments:
        payment.display_amount = _control_money(payment.amount)
        payment.created_local = timezone.localtime(payment.created_at)
        payment.customer_name = payment.user.get_full_name().strip() or payment.user.username

    for profile in live_users:
        profile.live_name = profile.user.get_full_name().strip() or profile.user.username
        profile.live_since = timezone.localtime(profile.muda_wa_mwisho)

    context = {
        "title": "TradeCore Control Room",
        "query": q,
        "total_users": total_users,
        "total_businesses": total_businesses,
        "registrations_today": registrations_today,
        "registrations_7d": registrations_7d,
        "registrations_30d": registrations_30d,
        "online_users": online_users,
        "active_trial": active_trial,
        "active_premium": active_premium,
        "expired_trial": expired_trial,
        "expired_premium": expired_premium,
        "revenue_total": _control_money(revenue_total),
        "revenue_today": _control_money(revenue_today),
        "pending_payments": pending_payments,
        "failed_payments": failed_payments,
        "recent_businesses": recent_businesses,
        "recent_payments": recent_payments,
        "live_users": live_users,
        "last_refresh": timezone.localtime(now),
    }
    return render(request, "control_room.html", context)


@login_required(login_url="login")
def billing_page(request):
    """Render the commercial TradeCore billing page."""
    monthly_price = PREMIUM_STANDARD_MONTHLY
    subscription = getattr(request.user, "subscription", None)
    return render(
        request,
        "billing.html",
        {
            "monthly_price": monthly_price,
            "monthly_price_display": f"{monthly_price:,.0f}",
            "subscription": subscription,
            "pricing_options": premium_pricing_options(),
            "premium_long_term_monthly_display": f"{PREMIUM_LONG_TERM_MONTHLY:,.0f}",
            "premium_max_months": PREMIUM_MAX_MONTHS,
            "trial_read_only": bool(
                subscription
                and subscription.plan_type == "TRIAL"
                and subscription.is_trial_expired
            ),
            "bank_name": (
                getattr(settings, "TRADECORE_BANK_NAME", None)
                or os.getenv("TRADECORE_BANK_NAME", "")
            ).strip(),
            "bank_account_name": (
                getattr(settings, "TRADECORE_BANK_ACCOUNT_NAME", None)
                or os.getenv("TRADECORE_BANK_ACCOUNT_NAME", "Meron Tech")
            ).strip(),
            "bank_account_number": (
                getattr(settings, "TRADECORE_BANK_ACCOUNT_NUMBER", None)
                or os.getenv("TRADECORE_BANK_ACCOUNT_NUMBER", "")
            ).strip(),
        },
    )


@login_required(login_url="login")
@require_POST
def initiate_mobile_payment(request):
    """Create a Tanzania mobile-money charge for the selected Premium duration."""
    network = (request.POST.get("network") or "").strip().upper()
    phone = _normalize_tanzania_phone(request.POST.get("phone_number"))
    try:
        pricing = calculate_premium_pricing(request.POST.get("months", "1"))
    except ValueError as exc:
        return JsonResponse({"status": "error", "message": str(exc)}, status=400)
    if network not in _FLUTTERWAVE_NETWORKS:
        return JsonResponse({"status": "error", "message": "Chagua mtandao sahihi wa simu."}, status=400)
    if not phone:
        return JsonResponse({"status": "error", "message": "Weka namba sahihi ya Tanzania (06/07...)."}, status=400)

    subscription, _ = BusinessSubscription.objects.get_or_create(user=request.user)
    if subscription.plan_type == "PREMIUM" and not subscription.is_premium_expired:
        return JsonResponse({"status": "success", "state": "already_active", "message": "TradeCore Premium yako bado iko active."})

    recent_pending = (
        PaymentTransaction.objects
        .filter(
            user=request.user,
            method="MOBILE_MONEY",
            status__in=["INITIATED", "PENDING"],
            created_at__gte=timezone.now() - timedelta(minutes=10),
        )
        .order_by("-created_at")
        .first()
    )
    if recent_pending:
        return JsonResponse({
            "status": "success",
            "state": "pending",
            "tx_ref": recent_pending.tx_ref,
            "payment_id": recent_pending.id,
            "message": "Tayari kuna ombi la malipo linalosubiriwa. Kamilisha prompt ya simu yako.",
        })

    email = (request.user.email or "").strip()
    if not email:
        return JsonResponse({"status": "error", "message": "Account yako haina email. Ongeza email kabla ya kulipia Premium."}, status=400)

    tx_ref = f"TC-PREM-{timezone.now():%Y%m%d%H%M%S}-{uuid.uuid4().hex[:10].upper()}"
    payment = PaymentTransaction.objects.create(
        user=request.user,
        subscription=subscription,
        tx_ref=tx_ref,
        amount=pricing["total"],
        currency="TZS",
        method="MOBILE_MONEY",
        network=network,
        phone_number=phone,
        status="INITIATED",
    )

    payload = {
        "tx_ref": tx_ref,
        "amount": str(payment.amount),
        "currency": "TZS",
        "email": email,
        "phone_number": phone,
        "network": _FLUTTERWAVE_NETWORKS[network]["api"],
        "fullname": (request.user.get_full_name() or request.user.username or "TradeCore Customer").strip(),
        "meta": {
            "tradecore_user_id": str(request.user.id),
            "payment_id": str(payment.id),
            "plan": "PREMIUM",
            "billing_months": pricing["months"],
            "unit_price": str(pricing["unit_price"]),
            "discount": str(pricing["discount"]),
            "standard_total": str(pricing["standard_total"]),
        },
    }

    try:
        response = requests.post(
            "https://api.flutterwave.com/v3/charges?type=mobile_money_tanzania",
            headers=_payment_api_headers(),
            json=payload,
            timeout=20,
        )
        try:
            result = response.json()
        except ValueError:
            result = {"status": "error", "message": response.text[:255]}
    except requests.RequestException:
        logger.exception("Flutterwave mobile-money initiation failed for payment=%s", payment.id)
        payment.status = "FAILED"
        payment.failure_reason = "Huduma ya malipo haikupatikana kwa sasa."
        payment.save(update_fields=["status", "failure_reason", "updated_at"])
        return JsonResponse({"status": "error", "message": "Huduma ya malipo haikupatikana kwa sasa. Jaribu tena."}, status=502)
    except RuntimeError as exc:
        payment.status = "FAILED"
        payment.failure_reason = str(exc)[:255]
        payment.save(update_fields=["status", "failure_reason", "updated_at"])
        return JsonResponse({"status": "error", "message": "Payment gateway haijawekwa vizuri kwenye server."}, status=503)

    data = result.get("data") if isinstance(result, dict) else {}
    gateway_status = str(result.get("status") or "").lower()
    payment.gateway_payload = dict(result) if isinstance(result, dict) else {"gateway_result": result}
    payment.gateway_payload["_tradecore"] = {
        "billing_months": pricing["months"],
        "unit_price": str(pricing["unit_price"]),
        "discount": str(pricing["discount"]),
        "standard_total": str(pricing["standard_total"]),
    }
    payment.gateway_transaction_id = str(data.get("id") or "")
    payment.gateway_reference = str(data.get("flw_ref") or "")[:150]
    payment.gateway_status = str(data.get("status") or gateway_status)
    payment.status = "PENDING" if gateway_status == "success" else "FAILED"
    if payment.status == "FAILED":
        payment.failure_reason = str(result.get("message") or data.get("processor_response") or "Malipo hayakuanzishwa.")[:255]
    payment.save(update_fields=[
        "gateway_payload",
        "gateway_transaction_id",
        "gateway_reference",
        "gateway_status",
        "status",
        "failure_reason",
        "updated_at",
    ])

    if payment.status == "FAILED":
        return JsonResponse({"status": "error", "message": payment.failure_reason or "Malipo hayakuanzishwa."}, status=400)

    return JsonResponse({
        "status": "success",
        "state": "pending",
        "tx_ref": payment.tx_ref,
        "payment_id": payment.id,
        "message": "Ombi la malipo limetumwa kwenye simu yako. Ingiza PIN kwenye prompt ya mtandao wako kukamilisha.",
    })


@login_required(login_url="login")
@require_GET
def mobile_payment_status(request):
    tx_ref = (request.GET.get("tx_ref") or "").strip()
    payment = get_object_or_404(PaymentTransaction, tx_ref=tx_ref, user=request.user)
    return JsonResponse({
        "status": "success",
        "payment_status": payment.status,
        "message": {
            "SUCCESSFUL": "Malipo yamepokelewa. Premium yako imeamilishwa.",
            "PENDING": "Bado tunasubiri uthibitisho wa malipo kutoka kwa mtandao.",
            "FAILED": payment.failure_reason or "Malipo hayajakamilika.",
        }.get(payment.status, "Tunasubiri majibu ya malipo."),
        "redirect_url": reverse("dashboard") if payment.status == "SUCCESSFUL" else "",
    })


@login_required(login_url="login")
@require_POST
def recheck_mobile_payment(request):
    tx_ref = (request.POST.get("tx_ref") or "").strip()
    payment = get_object_or_404(PaymentTransaction, tx_ref=tx_ref, user=request.user)
    if not payment.gateway_transaction_id:
        return JsonResponse({"status": "error", "message": "Transaction bado haijapata reference ya gateway."}, status=400)
    if payment.status == "SUCCESSFUL":
        return JsonResponse({"status": "success", "payment_status": "SUCCESSFUL", "redirect_url": reverse("dashboard")})
    try:
        payload = _verify_flutterwave_transaction(payment.gateway_transaction_id, payment.tx_ref)
        verified = _activate_payment_transaction(payment, payload, transaction_id=payment.gateway_transaction_id)
    except (requests.RequestException, RuntimeError, ValueError):
        logger.exception("Flutterwave payment recheck failed for payment=%s", payment.id)
        return JsonResponse({"status": "error", "message": "Imeshindikana kuhakiki malipo kwa sasa. Jaribu tena."}, status=502)

    return JsonResponse({
        "status": "success" if verified else "pending",
        "payment_status": payment.status,
        "message": "Malipo yamepokelewa. Premium yako imeamilishwa." if verified else "Bado hatujapata uthibitisho wa malipo.",
        "redirect_url": reverse("dashboard") if verified else "",
    })


@csrf_exempt
@require_POST
def flutterwave_webhook(request):
    """Verify Flutterwave signature, then server-verify and apply a payment exactly once."""
    secret_hash = _flutterwave_secret_hash()
    raw_body = request.body
    legacy_signature = request.headers.get("verif-hash")
    modern_signature = request.headers.get("flutterwave-signature")

    signature_valid = False
    if secret_hash and legacy_signature:
        signature_valid = hmac.compare_digest(str(legacy_signature), secret_hash)
    elif secret_hash and modern_signature:
        expected_signature = hmac.new(
            secret_hash.encode("utf-8"),
            raw_body,
            hashlib.sha256,
        ).digest()
        expected_signature_b64 = base64.b64encode(expected_signature).decode("utf-8")
        signature_valid = hmac.compare_digest(str(modern_signature), expected_signature_b64)

    if not signature_valid:
        return HttpResponse(status=401)

    try:
        payload = json.loads(raw_body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return HttpResponse(status=400)

    data = payload.get("data") if isinstance(payload, dict) else None
    if not isinstance(data, dict):
        return HttpResponse(status=200)

    tx_ref = str(data.get("tx_ref") or "").strip()
    if not tx_ref:
        return HttpResponse(status=200)

    payment = PaymentTransaction.objects.filter(tx_ref=tx_ref).first()
    if payment is None:
        return HttpResponse(status=200)
    if payment.status == "SUCCESSFUL":
        return HttpResponse(status=200)

    transaction_id = data.get("id")
    try:
        if transaction_id:
            verified_payload = _verify_flutterwave_transaction(transaction_id, payment.tx_ref)
        else:
            verified_payload = payload
        _activate_payment_transaction(payment, verified_payload, transaction_id=transaction_id)
    except (requests.RequestException, RuntimeError, ValueError):
        logger.exception("Flutterwave webhook verification failed for tx_ref=%s", tx_ref)
        return HttpResponse(status=500)

    return HttpResponse(status=200)


@login_required(login_url="login")
@require_POST
def submit_bank_payment(request):
    """Record a bank transfer as pending; it never self-activates Premium."""
    try:
        pricing = calculate_premium_pricing(request.POST.get("months", "1"))
    except ValueError as exc:
        return JsonResponse({"status": "error", "message": str(exc)}, status=400)

    subscription, _ = BusinessSubscription.objects.get_or_create(user=request.user)
    if subscription.plan_type == "PREMIUM" and not subscription.is_premium_expired:
        return JsonResponse({"status": "success", "state": "already_active", "message": "TradeCore Premium yako bado iko active."})

    recent_bank = (
        PaymentTransaction.objects
        .filter(
            user=request.user,
            method="BANK_TRANSFER",
            status="PENDING_BANK",
            created_at__gte=timezone.now() - timedelta(days=1),
        )
        .order_by("-created_at")
        .first()
    )
    if recent_bank:
        return JsonResponse({
            "status": "success",
            "state": "pending_bank",
            "reference": recent_bank.tx_ref,
            "message": "Tayari kuna bank transfer pending yenye reference hii.",
        })

    tx_ref = f"TC-BANK-{timezone.now():%Y%m%d%H%M%S}-{uuid.uuid4().hex[:10].upper()}"
    payment = PaymentTransaction.objects.create(
        user=request.user,
        subscription=subscription,
        tx_ref=tx_ref,
        amount=pricing["total"],
        currency="TZS",
        method="BANK_TRANSFER",
        status="PENDING_BANK",
        gateway_payload={
            "_tradecore": {
                "billing_months": pricing["months"],
                "unit_price": str(pricing["unit_price"]),
                "discount": str(pricing["discount"]),
                "standard_total": str(pricing["standard_total"]),
            }
        },
    )
    return JsonResponse({
        "status": "success",
        "state": "pending_bank",
        "reference": payment.tx_ref,
        "message": "Transfer yako imewekwa pending. Premium itaamilishwa baada ya payment kuthibitishwa.",
    })



@login_required(login_url="login")
@require_http_methods(["GET", "POST"])
def daily_report_settings_api(request):
    """Read or persist the tenant's Daily Business Report preferences."""
    denied = require_roles(request, "admin")
    if denied:
        return denied

    duka = get_store_profile(request)
    if not duka:
        return JsonResponse(
            {"status": "error", "message": "Biashara haijapatikana."},
            status=400,
        )

    report_settings, _ = DailyReportSettings.objects.get_or_create(duka=duka)

    if request.method == "GET":
        return JsonResponse(
            {
                "status": "success",
                "settings": {
                    "enabled": report_settings.enabled,
                    "report_time": report_settings.report_time.strftime("%H:%M"),
                    "recipient_boss": report_settings.recipient_boss,
                    "recipient_manager": report_settings.recipient_manager,
                    "include_sales_summary": report_settings.include_sales_summary,
                    "include_payment_breakdown": report_settings.include_payment_breakdown,
                    "include_staff_performance": report_settings.include_staff_performance,
                    "include_products_sold": report_settings.include_products_sold,
                    "include_discounts": report_settings.include_discounts,
                    "include_markups": report_settings.include_markups,
                    "include_debts": report_settings.include_debts,
                    "include_transaction_details": report_settings.include_transaction_details,
                },
            }
        )

    raw_time = (request.POST.get("report_time") or "").strip()
    if raw_time:
        try:
            parsed_time = datetime.strptime(raw_time, "%H:%M").time()
        except ValueError:
            return JsonResponse(
                {"status": "error", "message": "Muda wa ripoti si sahihi. Tumia mfumo wa Saa:Dakika."},
                status=400,
            )
        report_settings.report_time = parsed_time

    def posted_bool(name):
        return (request.POST.get(name) or "").strip().lower() in {"1", "true", "yes", "on"}

    if "enabled" in request.POST:
        report_settings.enabled = posted_bool("enabled")

    checkbox_fields = (
        "recipient_boss",
        "recipient_manager",
        "include_sales_summary",
        "include_payment_breakdown",
        "include_staff_performance",
        "include_products_sold",
        "include_discounts",
        "include_markups",
        "include_debts",
        "include_transaction_details",
    )
    for field in checkbox_fields:
        if field in request.POST:
            setattr(report_settings, field, posted_bool(field))

    report_settings.save()

    return JsonResponse(
        {
            "status": "success",
            "message": "Mipangilio ya Daily Business Report imehifadhiwa.",
            "settings": {
                "enabled": report_settings.enabled,
                "report_time": report_settings.report_time.strftime("%H:%M"),
            },
        }
    )


@login_required(login_url="login")
@require_POST
def whatsapp_send_api(request):
    """Send one customer WhatsApp message through Momo Business."""
    denied = require_roles(request, "admin")
    if denied:
        return denied

    duka = get_store_profile(request)
    if not duka:
        return JsonResponse(
            {"ok": False, "message": "Biashara ya mtumiaji haijapatikana."},
            status=400,
        )

    api_base = (
        getattr(settings, "MOMO_API_BASE_URL", "https://business.momo.tz/api/v3")
        or "https://business.momo.tz/api/v3"
    ).strip().rstrip("/")
    token = (getattr(settings, "MOMO_API_TOKEN", "") or "").strip()
    sender_id = (getattr(settings, "MOMO_WHATSAPP_SENDER_ID", "") or "").strip()

    if not token:
        return JsonResponse(
            {"ok": False, "message": "Momo WhatsApp haijawekwa: MOMO_API_TOKEN haipo kwenye environment variables."},
            status=503,
        )

    target = (request.POST.get("recipient") or "").strip()
    message_text = (request.POST.get("message") or "").strip()

    if target in {"", "all", "leo", "deni"}:
        return JsonResponse({"ok": False, "message": "Chagua mteja mmoja kwa kutuma WhatsApp."}, status=400)
    if not message_text:
        return JsonResponse({"ok": False, "message": "Andika ujumbe kwanza."}, status=400)
    if len(message_text) > 4096:
        return JsonResponse({"ok": False, "message": "Ujumbe wa WhatsApp haupaswi kuzidi characters 4096."}, status=400)

    digits = re.sub(r"\D", "", target)
    if digits.startswith("0") and len(digits) >= 10:
        digits = "255" + digits[1:]
    elif digits.startswith("7") and len(digits) >= 9:
        digits = "255" + digits
    if not digits.startswith("255") or len(digits) != 12:
        return JsonResponse({"ok": False, "message": "Namba ya WhatsApp si sahihi. Tumia namba ya Tanzania."}, status=400)

    variants = {target, digits, f"+{digits}", f"0{digits[3:]}"}
    customer_q = Q()
    for value in variants:
        customer_q |= Q(namba_ya_simu__iexact=value)
        customer_q |= Q(whatsapp_no__iexact=value)

    customer = Mteja.objects.filter(duka=duka).filter(customer_q).first()
    if not customer:
        return JsonResponse({"ok": False, "message": "Mteja huyo hayupo kwenye biashara hii."}, status=404)

    recipient = re.sub(r"\D", "", customer.whatsapp_no or customer.namba_ya_simu or "")
    if recipient.startswith("0") and len(recipient) >= 10:
        recipient = "255" + recipient[1:]
    elif recipient.startswith("7") and len(recipient) >= 9:
        recipient = "255" + recipient
    if not recipient.startswith("255") or len(recipient) != 12:
        return JsonResponse({"ok": False, "message": "Namba ya WhatsApp ya mteja si sahihi."}, status=400)

    store_name = (getattr(duka, "jina_la_duka", None) or "TradeCore").strip()
    message_text = (
        message_text.replace("{Jina}", customer.majina_kamili or "")
        .replace("{Duka}", store_name)
        .replace("{Deni}", f"{customer.outstanding_balance or 0:,.0f}")
    )

    payload = {"recipient": recipient, "message": message_text}
    if sender_id:
        payload["sender_id"] = sender_id

    try:
        response = requests.post(
            f"{api_base}/whatsapp/send",
            headers={
                "Authorization": f"Bearer {token}",
                "Accept": "application/json",
                "Content-Type": "application/json",
            },
            json=payload,
            timeout=30,
        )
        try:
            response_data = response.json()
        except ValueError:
            response_data = {"message": (response.text or "")[:500]}
    except requests.RequestException:
        logger.exception("Momo WhatsApp send failed: customer_id=%s user=%s", customer.id, request.user.username)
        CommunicationLog.objects.create(
            duka=duka, customer=customer, channel="WHATSAPP", message_type="MANUAL",
            recipient=f"+{recipient}", message=message_text, status="FAILED",
            error_message="Imeshindikana kuwasiliana na Momo WhatsApp API.", created_by=request.user,
        )
        return JsonResponse({"ok": False, "message": "Imeshindikana kuwasiliana na Momo WhatsApp API kwa sasa."}, status=502)

    messages_data = []
    if isinstance(response_data, dict):
        messages_data = (response_data.get("data") or {}).get("messages") or []
    first = messages_data[0] if messages_data else {}
    provider_id = str(first.get("uid") or first.get("id") or first.get("gateway_message_id") or "")[:255]
    provider_status = str(first.get("status") or "").lower()
    api_success = response.status_code in {200, 201, 202}
    delivery_ok = provider_status not in {"failed", "error"}
    log_status = "SENT" if api_success and delivery_ok else "FAILED"

    CommunicationLog.objects.create(
        duka=duka, customer=customer, channel="WHATSAPP", message_type="MANUAL",
        recipient=f"+{recipient}", message=message_text, status=log_status,
        provider_message_id=provider_id or None,
        provider_response={"status_code": response.status_code, "response": response_data},
        error_message=None if log_status == "SENT" else str(
            first.get("error_message") or (response_data.get("message") if isinstance(response_data, dict) else "") or "Momo WhatsApp imekataa ujumbe."
        ),
        created_by=request.user,
    )

    if log_status == "SENT":
        msg = (
            f"WhatsApp imewekwa kwenye queue kwa {customer.majina_kamili}."
            if provider_status == "queued"
            else f"WhatsApp imetumwa kwa {customer.majina_kamili}."
        )
        return JsonResponse({"ok": True, "message": msg, "customer": customer.majina_kamili})

    provider_message = str(
        first.get("error_message")
        or (response_data.get("message") if isinstance(response_data, dict) else "")
        or "Momo WhatsApp imekataa ujumbe."
    )[:500]
    return JsonResponse({"ok": False, "message": provider_message}, status=400)


@login_required(login_url="login")
@require_POST
def tuma_daily_report_whatsapp(request):
    """Send the approved daily-report WhatsApp template through Momo Business."""
    denied = require_roles(request, "admin")
    if denied:
        return denied

    duka = get_store_profile(request)
    if not duka:
        return JsonResponse({"status": "error", "message": "Biashara haijapatikana."}, status=400)

    api_base = (
        getattr(settings, "MOMO_API_BASE_URL", "https://business.momo.tz/api/v3")
        or "https://business.momo.tz/api/v3"
    ).strip().rstrip("/")
    token = (getattr(settings, "MOMO_API_TOKEN", "") or "").strip()
    sender_id = (getattr(settings, "MOMO_WHATSAPP_SENDER_ID", "") or "").strip()
    template_name = (getattr(settings, "MOMO_DAILY_REPORT_TEMPLATE_NAME", "tradecore_daily_business_report") or "tradecore_daily_business_report").strip()
    template_language = (getattr(settings, "MOMO_DAILY_REPORT_TEMPLATE_LANGUAGE", "sw") or "sw").strip()

    if not token:
        return JsonResponse({"status": "error", "message": "Momo WhatsApp haijawekwa: MOMO_API_TOKEN haipo kwenye environment variables."}, status=503)

    leo = timezone.localdate()
    report = _daily_report_data(duka, leo)
    wateja_count = (
        Mauzo.objects.filter(duka=duka, tarehe_ya_mauzo__date=leo, mteja__isnull=False)
        .values("mteja_id").distinct().count()
    )

    try:
        if getattr(duka, "inventory_location_mode", False):
            ensure_inventory_locations(duka)
            stock_rows = ProductStock.objects.filter(location__duka=duka).select_related("bidhaa")
            stock_iliyopo = sum(
                (row.bidhaa.bei_ya_kuuzia or Decimal("0")) * int(row.quantity or 0)
                for row in stock_rows
            )
        else:
            stock_iliyopo = (
                Bidhaa.objects.filter(duka=duka)
                .aggregate(total=Sum(ExpressionWrapper(F("bei_ya_kuuzia") * F("idadi_stoo"), output_field=DecimalField(max_digits=14, decimal_places=2))))["total"]
                or Decimal("0")
            )
    except Exception:
        logger.exception("Daily report stock valuation failed")
        stock_iliyopo = Decimal("0")

    owner = getattr(duka, "mwenye_duka", None)
    jina_la_mpokeaji = (
        (owner.get_full_name().strip() if owner else "")
        or (owner.first_name.strip() if owner else "")
        or (owner.username if owner else "")
        or "Mteja"
    )
    biashara = (getattr(duka, "jina_la_duka", None) or "TradeCore").strip()
    tarehe_str = leo.strftime("%d %b %Y")
    mauzo_str = f"{report['gross_sales']:,.0f}"
    faida_str = f"{report['profit']:,.0f}"
    wateja_str = str(wateja_count)
    bidhaa_str = str(report["total_items"])
    matumizi_str = f"{report['expenses_total']:,.0f}"
    stock_str = f"{stock_iliyopo:,.0f}"
    report_token = _make_daily_report_token(duka, leo)

    recipient = re.sub(r"\D", "", (getattr(duka, "simu", None) or "").strip())
    if recipient.startswith("0"):
        recipient = "255" + recipient[1:]
    elif recipient.startswith("7"):
        recipient = "255" + recipient
    if not recipient.startswith("255") or len(recipient) != 12:
        return JsonResponse({"status": "error", "message": "Namba ya WhatsApp ya biashara haijawekwa kwa mfumo sahihi wa Tanzania (+255...)."}, status=400)

    hour = timezone.localtime(timezone.now()).hour
    greeting = "Habari za asubuhi" if 5 <= hour < 12 else ("Habari za mchana" if 12 <= hour < 18 else "Habari za jioni")
    report_link = request.build_absolute_uri(reverse("daily_report_pdf_public", args=[report_token]))

    payload = {
        "recipient": recipient,
        "message_type": "template",
        "template": {
            "name": template_name,
            "language": template_language,
            "components": [
                {
                    "type": "body",
                    "parameters": [
                        {"type": "text", "text": f"{greeting} {jina_la_mpokeaji}"},
                        {"type": "text", "text": biashara},
                        {"type": "text", "text": tarehe_str},
                        {"type": "text", "text": mauzo_str},
                        {"type": "text", "text": faida_str},
                        {"type": "text", "text": wateja_str},
                        {"type": "text", "text": bidhaa_str},
                        {"type": "text", "text": matumizi_str},
                        {"type": "text", "text": stock_str},
                    ],
                },
                {
                    "type": "button",
                    "sub_type": "url",
                    "index": "0",
                    "parameters": [{"type": "text", "text": report_token}],
                },
            ],
        },
    }
    if sender_id:
        payload["sender_id"] = sender_id

    try:
        response = requests.post(
            f"{api_base}/whatsapp/send",
            headers={"Authorization": f"Bearer {token}", "Accept": "application/json", "Content-Type": "application/json"},
            json=payload,
            timeout=30,
        )
        try:
            response_data = response.json()
        except ValueError:
            response_data = {"message": (response.text or "")[:500]}
    except requests.RequestException:
        logger.exception("Momo daily report request failed")
        return JsonResponse({"status": "error", "message": "Imeshindikana kuwasiliana na Momo WhatsApp API kwa sasa."}, status=502)

    messages_data = (response_data.get("data") or {}).get("messages") if isinstance(response_data, dict) else []
    messages_data = messages_data or []
    first = messages_data[0] if messages_data else {}
    provider_status = str(first.get("status") or "").lower()
    provider_id = str(first.get("uid") or first.get("id") or first.get("gateway_message_id") or "")[:255]
    api_success = response.status_code in {200, 201, 202}
    delivery_ok = provider_status not in {"failed", "error"}

    if api_success and delivery_ok:
        CommunicationLog.objects.create(
            duka=duka,
            customer=None,
            channel="WHATSAPP",
            message_type="DAILY_REPORT",
            recipient=f"+{recipient}",
            message=f"Daily Business Report — {biashara} — {tarehe_str}",
            status="SENT",
            provider_message_id=provider_id or None,
            provider_response={
                "status_code": response.status_code,
                "response": response_data,
            },
            error_message=None,
            created_by=request.user,
        )
        return JsonResponse({
            "status": "success",
            "message": "Ripoti imewekwa kwenye queue ya Momo WhatsApp." if provider_status == "queued" else "Ripoti imetumwa kupitia Momo WhatsApp.",
            "provider_message_id": provider_id,
            "provider_status": provider_status or "accepted",
            "report_url": report_link,
        })

    provider_message = str(
        first.get("error_message")
        or (response_data.get("message") if isinstance(response_data, dict) else "")
        or "Momo WhatsApp imekataa ripoti."
    )[:500]
    return JsonResponse({"status": "error", "message": provider_message}, status=400)

